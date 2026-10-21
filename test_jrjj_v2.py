"""Offline regression tests for dual-source scraper."""
import io
import json
import os
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from docx import Document
import jrjj_v2 as j

GOV_HTML='''<html><div class="redesigned-list"><ul><li><a href="/zwzx/jrjj/202610/t20261019_1234567.html">市委召开专项会议</a><em>2026-10-19</em></li><li><a href="/zwzx/jrjj/202610/t20261018_2222222.html">另一条普通新闻</a><time>2026年10月18日</time></li></ul></div></html>'''
JJXW_HTML='''<html><section class="news"><ul><li><h3><a href="/column/668/article/2026/10/19/803640.html">陈云在九江调研</a></h3><a href="/column/668/index.html">九江时政</a><span>2026-10-19</span></li><li><h3><a href="/column/668/article/2026/10/18/803641.html">市委召开专项会议</a></h3><p>2026-10-18</p></li></ul></section></html>'''
ARTICLE_MATCH='''<html><h1>市委召开专项会议</h1><div class="TRS_Editor"><p>会议召开后，各部门进行了深入讨论。黄强在发言中对工作作出安排，要求加强协作推动各项工作落实。</p><p>与会人员表示，将进一步细化任务、落实责任，持续推进项目建设与管理，确保年度各项任务落实到位。</p></div><div class="related"><p>其他领导相关新闻陈永奇</p></div></html>'''
ARTICLE_JJ='''<html><h1>陈云在九江调研</h1><div class="article-content"><p>10月19日上午，陈云率领相关部门负责同志进行实地调研，深入了解产业发展中的困难。</p><p>他要求进一步优化营商环境，加强保障服务，推进区域高质量发展取得扎实成效。</p></div></html>'''
ARTICLE_OTHER='''<html><h1>另一条普通新闻</h1><div id="artibody"><p>今天上午，多部门组织开展环境卫生改善行动，推动辖区生产生活环境进一步改善。</p><p>有关部门研究部署了下一阶段工作重点，确定任务分工和落实计划，现场检查了重点区域。</p></div></html>'''

GOV_XML='''<?xml version="1.0" encoding="UTF-8"?><ROOT>
<ITEM><TITLE><![CDATA[市委召开专项会议]]></TITLE><PUBDATE>2026-10-19 11:30:00</PUBDATE><PUBURL>./202610/t20261019_1234567.html</PUBURL></ITEM>
<ITEM><TITLE>另一条普通新闻</TITLE><PUBURL>./202610/t20261018_2222222.html</PUBURL></ITEM></ROOT>'''
JJXW_JSON=json.dumps({"data":{"list":[{"fileID":803640,"publishTime":"2026-10-19 10:45:07","title":"陈云在九江调研"},
                                      {"fileID":"803641","publishTime":"2026-10-18 09:00:00","title":"市委召开专项会议"}]}},ensure_ascii=False)

class TestJrjj(unittest.TestCase):
    def test_gov_xml_direct_requests_get_prefers_ipv4_and_preserves_bytes(self):
        response = j.requests.Response()
        response.status_code = 200
        response._content = GOV_XML.encode('utf-8')
        original_family = j.urllib3_connection.allowed_gai_family
        with patch.object(j.requests, 'get', return_value=response) as request_get:
            xml = j.fetch_government_xml(
                j.build_session('test-agent'),
                'https://www.jiujiang.gov.cn/zwzx/jrjj/data_2026.xml', 10)
        self.assertEqual(j.urllib3_connection.allowed_gai_family, original_family)
        self.assertEqual(len(j.parse_gov_xml(xml)), 2)
        self.assertEqual(request_get.call_count, 1)
        self.assertEqual(request_get.call_args.kwargs['timeout'], (6.0, 10.0))
        self.assertIn('application/xml', request_get.call_args.kwargs['headers']['Accept'])

    def test_gov_xml_ipv4_failure_retries_default_stack_once(self):
        response = j.requests.Response()
        response.status_code = 200
        response._content = GOV_XML.encode('utf-8')
        original_family = j.urllib3_connection.allowed_gai_family
        with patch.object(j.requests, 'get', side_effect=[
            j.requests.exceptions.ConnectionError('[Errno 101] Network is unreachable'),
            response,
        ]) as request_get:
            xml = j.fetch_government_xml(
                j.build_session('test-agent'),
                'https://www.jiujiang.gov.cn/zwzx/jrjj/data_2026.xml', 8)
        self.assertEqual(j.urllib3_connection.allowed_gai_family, original_family)
        self.assertEqual(request_get.call_count, 2)
        self.assertEqual(len(j.parse_gov_xml(xml)), 2)

    def test_curl_fallback_valid_xml(self):
        from subprocess import CompletedProcess
        result = CompletedProcess(['curl'], 0, stdout=GOV_XML.encode('utf-8'), stderr=b'')
        with patch.object(j.subprocess, 'run', return_value=result) as run:
            rows = j.parse_gov_xml(j.fetch_gov_xml_with_curl(
                'https://www.jiujiang.gov.cn/zwzx/jrjj/data_2026.xml', 8))
        self.assertEqual(len(rows), 2)
        self.assertIn('--compressed', run.call_args.args[0])
        self.assertIn('--fail', run.call_args.args[0])

    def test_curl_fallback_rejects_html_and_http_failure(self):
        from subprocess import CompletedProcess, CalledProcessError
        url = 'https://www.jiujiang.gov.cn/zwzx/jrjj/data_2026.xml'
        with patch.object(j.subprocess, 'run', return_value=CompletedProcess(
                ['curl'], 0, stdout=b'<html>Forbidden</html>', stderr=b'')):
            with self.assertRaises(Exception):
                j.fetch_gov_xml_with_curl(url, 8)
        with patch.object(j.subprocess, 'run', side_effect=CalledProcessError(
                22, ['curl'], stderr=b'HTTP 403')):
            with self.assertRaisesRegex(RuntimeError, '403'):
                j.fetch_gov_xml_with_curl(url, 8)

    def test_curl_used_only_after_both_requests_connection_failures(self):
        err = j.requests.exceptions.ConnectionError('network unreachable')
        with patch.object(j.requests, 'get', side_effect=[err, err]) as get, \
             patch.object(j, 'fetch_gov_xml_with_curl', return_value=GOV_XML.encode('utf-8')) as curl:
            content = j.fetch_government_xml(j.build_session('test-agent'),
                'https://www.jiujiang.gov.cn/zwzx/jrjj/data_2026.xml', 7)
        self.assertEqual(get.call_count, 2)
        curl.assert_called_once()
        self.assertEqual(len(j.parse_gov_xml(content)), 2)

    def test_gov_xml_http_error_is_not_retried(self):
        response = j.requests.Response()
        response.status_code = 403
        response.url = 'https://www.jiujiang.gov.cn/zwzx/jrjj/data_2026.xml'
        response._content = b'Forbidden'
        response.raw = io.BytesIO(b'Forbidden')
        with patch.object(j.requests, 'get', return_value=response) as request_get:
            with self.assertRaises(j.requests.exceptions.HTTPError):
                j.fetch_government_xml(j.build_session('test-agent'), response.url, 8)
        self.assertEqual(request_get.call_count, 1)

    def test_curl_uses_authorized_headers_without_process_arg_leaks(self):
        response = CompletedProcess(['curl'], 0, stdout=GOV_XML.encode('utf-8'), stderr=b'')
        secrets = {'JIUJIANG_API_KEY': 'example-test-key',
                   'JIUJIANG_COOKIE': 'example-test-cookie'}
        with patch.dict(os.environ, secrets), patch.object(j.subprocess, 'run',
                                                            return_value=response) as run:
            payload = j.fetch_gov_xml_with_curl(
                'https://www.jiujiang.gov.cn/zwzx/jrjj/data_2026.xml', 8)
        self.assertEqual(len(j.parse_gov_xml(payload)), 2)
        args = run.call_args.args[0]
        self.assertNotIn('example-test-key', ' '.join(args))
        self.assertNotIn('example-test-cookie', ' '.join(args))
        self.assertIn('--config', args)
        self.assertIn(b'x-api-key: example-test-key', run.call_args.kwargs['input'])
        self.assertIn(b'Cookie: example-test-cookie', run.call_args.kwargs['input'])

    def test_local_gov_xml_import_skips_network(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'data_2026.xml'
            path.write_bytes(GOV_XML.encode('utf-8'))
            config = {**j.config_defaults(), 'gov_xml_file': str(path)}
            with patch.object(j, 'fetch_government_xml',
                              side_effect=AssertionError('should not fetch')):
                rows, ok, _ = j.fetch_list(j.build_session('test-agent'),
                                           'gov', 0, config, False, Path(folder), year=2026)
            self.assertTrue(ok)
            self.assertEqual(len(rows), 2)

    def test_invalid_local_gov_xml_is_rejected(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'data_2026.xml'
            path.write_text('<html>Forbidden</html>', encoding='utf-8')
            config = {**j.config_defaults(), 'gov_xml_file': str(path)}
            rows, ok, _ = j.fetch_list(j.build_session('test-agent'),
                                       'gov', 0, config, False, Path(folder), year=2026)
            self.assertFalse(ok)
            self.assertEqual(rows, [])

    def test_government_annual_xml_feed(self):
        records=j.parse_gov_xml(GOV_XML)
        self.assertEqual(len(records),2)
        self.assertEqual(records[0][0], '市委召开专项会议')
        self.assertEqual(records[0][1], 'https://www.jiujiang.gov.cn/zwzx/jrjj/202610/t20261019_1234567.html')
        self.assertEqual(records[1][2], '2026-10-18')

    def test_jjxw_api_100_per_page_and_links(self):
        records=j.parse_jjxw_json(JJXW_JSON)
        self.assertEqual(len(records),2)
        self.assertEqual(records[0][1], 'https://www.jjxw.cn/article/2026/10/19/803640.html')
        self.assertEqual(records[0][2], '2026-10-19')
        self.assertEqual(records[1][0], '市委召开专项会议')

    def test_structured_fetch_routing_no_html_listing(self):
        with tempfile.TemporaryDirectory() as folder:
            session=j.build_session('test-agent')
            with patch.object(j,'fetch_government_xml',return_value=GOV_XML.encode('utf-8')) as gov_get, patch.object(j,'fetch_html',side_effect=[JJXW_JSON, '{"data":{"list":[]}}']) as fetch:
                gov, ok, _=j.fetch_list(session,'gov',0,j.config_defaults(),False,Path(folder),year=2026)
                self.assertTrue(ok)
                self.assertEqual(len(gov),2)
                self.assertIn('data_2026.xml',gov_get.call_args.args[1])
                jx, ok, _=j.fetch_list(session,'jjxw',0,j.config_defaults(),False,Path(folder),year=2026)
                self.assertTrue(ok)
                self.assertEqual(len(jx),2)
                self.assertIn('pageSize=100&pageNumber=1',fetch.call_args_list[0].args[1])
                empty, ok, _=j.fetch_list(session,'jjxw',1,j.config_defaults(),False,Path(folder),year=2026)
                self.assertTrue(ok)
                self.assertFalse(empty)
    def test_unapproved_organization_terms_never_match(self):
        """No organization name or unrelated leader is approved implicitly."""
        cfg = {**j.config_defaults(), 'names': ['陈云', '黄强']}
        names = j.get_names(cfg)
        title = '市人大常委会党组理论学习中心组举行集体学习，廖奇志主持并讲话'
        body = ('9月2日，市人大常委会党组理论学习中心组举行集体学习。'
                '市人大常委会党组书记、主任廖奇志主持并讲话。'
                '习近平党建思想深化了党对新时代管党治党的规律性认识。')
        self.assertEqual(j.match_names(title, body, names, True), [])
        self.assertEqual(j.match_names('陈云在九江调研', '', names, True), ['陈云'])

    def test_names_only_from_config_by_default(self):
        cfg = j.config_defaults()
        self.assertNotIn('fetch_ifeng_names', cfg)
        self.assertEqual(j.get_names(cfg), sorted(cfg['names']))
        self.assertFalse(hasattr(j, 'extract_ifeng_names'))

    def test_cli_help_with_cp1252_output(self):
        """Chinese argparse help must print even when inherited output is cp1252."""
        env = os.environ.copy()
        env['PYTHONIOENCODING'] = 'cp1252'
        result = subprocess.run(
            [sys.executable, str(Path(j.__file__).resolve()), '--help'],
            capture_output=True, env=env, timeout=30,
        )
        self.assertEqual(result.returncode, 0, result.stderr.decode('utf-8', errors='replace'))
        self.assertIn('抓取', result.stdout.decode('utf-8'))

    def test_sqlite_context_releases_file_handle(self):
        """The database must be closed, not merely committed, on Windows."""
        with tempfile.TemporaryDirectory() as folder:
            db = Path(folder) / 'lock_test.sqlite3'
            with j.connect_db(db) as connection:
                connection.execute(
                    "INSERT INTO metadata(key,value) VALUES(?,?)",
                    ('close_test', 'ok')
                )
            with self.assertRaises(sqlite3.ProgrammingError):
                connection.execute("SELECT 1")
            db.unlink()
            self.assertFalse(db.exists())

    def test_parse_redesigned_gov(self):
        items=j.parse_listing(GOV_HTML,'gov',j.SOURCES['gov']['base'])
        self.assertEqual(len(items),2)
        self.assertEqual(items[0].pub_date,'2026-10-19')

    def test_parse_jjxw_category_link_and_relative_paths(self):
        items=j.parse_listing(JJXW_HTML,'jjxw',j.SOURCES['jjxw']['base'])
        self.assertEqual(len(items),2)
        self.assertIn('/article/2026/10/',items[0].url)
        self.assertNotIn('/column/668/article/',items[0].url)
        self.assertEqual(j.next_page_urls('jjxw',1)[0],'https://www.jjxw.cn/column/668/index2.html')

    def test_article_body_no_related(self):
        title,paras,selector=j.parse_article(ARTICLE_MATCH,'fallback')
        self.assertEqual(title,'市委召开专项会议')
        self.assertEqual(len(paras),2)
        self.assertIn('黄强',''.join(paras))
        self.assertNotIn('陈永奇',''.join(paras))
        self.assertEqual(j.match_names(title,''.join(paras),['黄强','陈永奇'],True),['黄强'])

    def test_names_strictly_from_config_no_external_lookup(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'config.json'
            config = {**j.config_defaults(), 'names': ['陈云', '黄强'],
                      'fetch_ifeng_names': True}
            path.write_text(json.dumps(config, ensure_ascii=False), encoding='utf-8')
            loaded = j.load_config(path)
            self.assertNotIn('fetch_ifeng_names', loaded)
            self.assertEqual(j.get_names(loaded), ['陈云', '黄强'])
            self.assertFalse(hasattr(j, 'extract_ifeng_names'))

    def test_organization_words_cannot_trigger_inclusion(self):
        cfg = {**j.config_defaults(), 'names': ['陈云', '黄强']}
        names = j.get_names(cfg)
        title = '市人大常委会党组理论学习中心组举行集体学习，廖奇志主持并讲话'
        body = ('九江新闻网讯 9月2日，市人大常委会党组理论学习中心组举行集体学习。'
                '人大、省委、省政府、政协、法院和检察院参与相关活动。')
        self.assertEqual(j.match_names(title, body, names, True), [])
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)
            book = j.open_book(path / 'news.xlsx')
            record = j.News('jjxw', title, 'https://www.jjxw.cn/article/2026/09/02/999999.html', '2026-09-02')
            j.write_new_rows(book, {'gov': [], 'jjxw': [record]})
            with j.connect_db(path / 'cache.sqlite3') as con:
                j.upsert_news(con, record)
                con.execute("UPDATE news SET status='成功', body_json=?, article_title=? WHERE url=?",
                            (json.dumps([body], ensure_ascii=False), title, record.url))
                con.execute("INSERT OR REPLACE INTO metadata(key,value) VALUES (?,?)",
                            ('last_names', json.dumps(['人大', '省委', '省政府', '政协'], ensure_ascii=False)))
                selected = j.refresh_book(book, con, j.get_names(cfg), cfg)
                self.assertEqual(selected, [])
                sheet = book['九江新闻网']
                self.assertEqual(sheet['C2'].value, '未命中')
                self.assertFalse(sheet['D2'].value)
        self.assertEqual(j.match_names('陈云在九江调研', '', names, True), ['陈云'])

    def test_upgrade_preserve_yellow_and_word(self):
        with tempfile.TemporaryDirectory() as folder:
            p=Path(folder)
            book=Workbook();ws=book.active;ws.title='九江新闻'
            ws.append(['新闻标题','发布日期'])
            c=ws.cell(2,1,'=HYPERLINK("https://www.jiujiang.gov.cn/zwzx/jrjj/202602/t20260223_8888888.html","旧新闻")')
            c.fill=PatternFill(fill_type='solid',fgColor='FFFF00')
            ws.cell(2,2,'2026-02-23');book.save(p/'jiujiang_news.xlsx')
            b=j.open_book(p/'jiujiang_news.xlsx')
            self.assertEqual(b.sheetnames,['九江市政府','九江新闻网'])
            with j.connect_db(p/'cache.sqlite3') as con:
                seen=j.existing_rows(b,con)
                self.assertEqual(len(seen['gov']),1)
                self.assertEqual(b['九江市政府']['A2'].value,'旧新闻')
                self.assertTrue(b['九江市政府']['A2'].hyperlink.target.endswith('8888888.html'))
                j.write_new_rows(b,{'gov':[j.News('gov','新增新闻','https://www.jiujiang.gov.cn/zwzx/jrjj/202610/t20261019_1234567.html','2026-10-19')],'jjxw':[]})
                self.assertTrue(j.is_yellow(b['九江市政府'].cell(3,1)))
                b.save(p/'jiujiang_news.xlsx')
                articles=[{'source':'gov','url':'x','date':'2026-10-19','title':'新闻标题','body':['这是测试正文。'*15],'status':'成功','matches':['黄强']}]
                dest=j.build_word('2026-10',articles,p)
                self.assertTrue(dest.exists())
                doc=Document(dest)
                self.assertIn('2026年10月',doc.paragraphs[0].text)
                self.assertIn('新闻标题',[x.text for x in doc.paragraphs])
                j.build_word('2026-10',articles,p)
                self.assertEqual(len(list((p/'备份').glob('*.docx'))),1)

    def test_cross_source_dedup_by_title_and_body(self):
        def record(source, title, day, text, status='成功'):
            return {'source': source, 'url': source + title, 'date': day,
                    'title': title, 'body': [text], 'status': status,
                    'matches': ['陈云']}
        body = '陈云对全市重点项目进行实地调研并提出具体要求，推动高质量发展。' * 5
        changed = '陈云主持另外一次会议，研究审议不同政策和工作方案。' * 5
        items = [
            record('jjxw', '陈云赴基层调研', '2026-09-10', body),
            record('gov', '陈云赴基层调研', '2026-09-11', body),
            record('jjxw', '陈云调研重点项目', '2026-09-12', body),
            record('gov', '陈云赴基层调研', '2026-09-10', changed),
        ]
        articles, failed = j.deduplicate_for_month(items, '2026-09')
        self.assertEqual(len(articles), 2)
        self.assertFalse(failed)
        self.assertIn(changed, [a['body'][0] for a in articles])
        self.assertEqual(next(a for a in articles if a['body'][0] == body)['source'], 'gov')

    def test_dedup_punctuation_minor_edits_and_month_filter(self):
        body = ('陈云出席市委专题会议并讲话，各单位认真学习贯彻会议精神。' * 8)
        altered = body.replace('，', '，\n').replace('市委专题会议', '市委重要专题会议', 1)
        def item(source, title, date, text, status='成功'):
            return {'source': source, 'title': title, 'date': date,
                    'body': [text], 'status': status, 'matches': ['陈云'], 'url': source + date}
        stories = [
            item('jjxw', '陈云 召开专题会议', '2026-09-05', body),
            item('gov', '陈云召开专题会议！', '2026-09-06', altered),
            item('jjxw', '陈云召开专题会议', '2026-09-20', body.replace('讲话', '调研')),
            item('gov', '未抓取成功', '2026-09-08', [], '失败'),
            item('gov', '不同月份', '2026-10-02', body)
        ]
        articles, failed = j.deduplicate_for_month(stories, '2026-09')
        self.assertEqual(len(articles), 2)
        self.assertEqual(len(failed), 1)
        self.assertTrue(all(a['date'].startswith('2026-09') for a in articles))

    def test_fake_full_pipeline(self):
        with tempfile.TemporaryDirectory() as folder:
            p=Path(folder)
            config={**j.config_defaults(),'delay_seconds':0,'max_pages':2,'stop_after_known_pages':2,'names':['陈云','黄强']}
            (p/'cfg.json').write_text(json.dumps(config,ensure_ascii=False),encoding='utf-8')
            def fake_fetch(session,url,timeout):
                if 'data_2026.xml' in url:return GOV_XML
                if 'getArticles' in url:
                    return JJXW_JSON if 'pageNumber=1' in url else '{"data":{"list":[]}}'
                if '1234567' in url or '803641' in url:return ARTICLE_MATCH
                if '803640' in url:return ARTICLE_JJ
                if '2222222' in url:return ARTICLE_OTHER
                raise RuntimeError('模拟第二页列表结束')
            with patch.object(j,'fetch_html',side_effect=fake_fetch), patch.object(j,'fetch_government_xml',return_value=GOV_XML.encode('utf-8')):
                code=j.main(['--config',str(p/'cfg.json'),'--excel',str(p/'data.xlsx'),'--database',str(p/'cache.sqlite3'),'--output-dir',str(p),'--debug-dir',str(p/'debug'),'--month','2026-10'])
            self.assertEqual(code,0)
            w=load_workbook(p/'data.xlsx')
            self.assertEqual(w.sheetnames,['九江市政府','九江新闻网'])
            doc=Document(p/'市委大事记（2026年10月）.docx')
            txt='\n'.join(x.text for x in doc.paragraphs)
            self.assertIn('陈云在九江调研',txt)
            self.assertIn('市委召开专项会议',txt)
            self.assertNotIn('另一条普通新闻',txt)
            self.assertEqual(txt.count('市委召开专项会议'),1)

if __name__=='__main__':
    unittest.main()
