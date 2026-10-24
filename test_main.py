"""九江新闻网单数据源离线回归与端到端测试。"""
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from docx import Document
import main as j


def article(title, body):
    return ('<html><h1>' + title + '</h1><div class="article-content">'
            '<p>' + body + '</p><p>' + body + '</p></div></html>')


class TestJjxw(unittest.TestCase):
    def test_only_news_source(self):
        self.assertEqual(list(j.SOURCES), ['jjxw'])
        self.assertNotIn('--gov-xml-file', j.args_parser().format_help())
        self.assertNotIn('--browser', j.args_parser().format_help())

    def test_simplified_cli_and_month_validation(self):
        parser = j.args_parser()
        self.assertIsNone(parser.parse_args([]).month)
        self.assertEqual(parser.parse_args(['2026-09']).month, '2026-09')
        self.assertTrue(parser.parse_args(['--test']).test)
        self.assertTrue(parser.parse_args(['2026-09', '--export']).export)
        self.assertTrue(parser.parse_args(['2026-09', '--verify']).verify)
        with self.assertRaises(SystemExit):
            parser.parse_args(['--export', '--sync'])
        with tempfile.TemporaryDirectory() as folder:
            with self.assertRaises(SystemExit):
                j.main(['2026-13'], base_dir=Path(folder))
            self.assertFalse((Path(folder) / 'results').exists())

    def test_no_args_defaults_to_previous_month_and_results_dir(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)
            with patch.object(j, 'previous_month', return_value='2026-09'), \
                 patch.object(j, 'crawl', return_value={'jjxw': []}) as crawl, \
                 patch.object(j, 'fetch_article_details'):
                status = j.main([], base_dir=p)
            self.assertEqual(status, 1)  # no matching body -> no Word, not success
            self.assertTrue((p / 'results' / 'jiujiang_news.xlsx').exists())
            self.assertTrue((p / 'results' / 'jiujiang_cache.sqlite3').exists())
            self.assertEqual(crawl.call_args.args[-1], '2026-09')

    def test_grace_window_inclusive_and_cross_year(self):
        self.assertEqual(j.publication_window('2026-09', 10),
                         ('2026-09-01', '2026-10-10'))
        self.assertEqual(j.publication_window('2026-09', 0),
                         ('2026-09-01', '2026-09-30'))
        self.assertEqual(j.publication_window('2026-12', 10),
                         ('2026-12-01', '2027-01-10'))

    def test_event_date_from_lead_and_relative_day(self):
        self.assertEqual(j.infer_event_date(
            '陈云赴九江调研',
            ['本报讯（记者某某）9月30日，陈云到九江调研并讲话。'],
            '2026-10-02'), '2026-09-30')
        self.assertEqual(j.infer_event_date(
            '邓永翔主持会议', ['昨日，邓永翔主持召开专题会议。'],
            '2026-10-01'), '2026-09-30')
        self.assertEqual(j.infer_event_date(
            '陈云主持会议', ['12月31日，陈云主持召开专题会议。'],
            '2027-01-04'), '2026-12-31')
        self.assertEqual(j.infer_event_date(
            '某事项继续推进', ['此前9月30日的统计数字显示全市情况。'],
            '2026-10-04'), '')

    def test_report_date_is_not_mistaken_for_prior_month_event(self):
        text = '10月2日，记者从有关部门获悉，9月30日，市委主持召开专题会议。'
        self.assertEqual(j.infer_event_date('重点会议', [text], '2026-10-02'),
                         '2026-09-30')
        text = '截至9月30日，全市召开了工作会议12场。'
        self.assertEqual(j.infer_event_date('情况通报', [text], '2026-10-02'), '')

    def test_reporter_learned_date_not_written_as_event(self):
        self.assertEqual(j.infer_event_date(
            '市委新闻', ['10月2日，记者获悉市委召开了重点工作会议。'],
            '2026-10-02'), '')
        self.assertEqual(j.infer_event_date(
            '市委新闻', ['记者10月2日获悉，9月30日，市委召开了重点工作会议。'],
            '2026-10-02'), '2026-09-30')

    def test_relative_event_dates_and_cross_year(self):
        self.assertEqual(j.infer_event_date(
            '工作动态', ['前天，市领导主持召开会议。'], '2026-10-02'),
            '2026-09-30')
        self.assertEqual(j.infer_event_date(
            '工作动态', ['前日，市委举行座谈会。'], '2027-01-02'),
            '2026-12-31')
        self.assertEqual(j.infer_event_date(
            '工作动态', ['上月30日，市领导主持召开会议。'], '2026-10-03'),
            '2026-09-30')
        self.assertEqual(j.infer_event_date(
            '工作动态', ['本月2日，市领导调研基层情况。'], '2026-10-03'),
            '2026-10-02')
        self.assertEqual(j.infer_event_date(
            '工作动态', ['上月31日，市领导召开会议。'], '2026-05-03'), '')
        self.assertEqual(j.infer_event_date(
            '工作动态', ['下月2日将举行会议。'], '2026-10-03'), '')

    def test_future_schedule_and_october_event_not_backdated(self):
        self.assertEqual(j.infer_event_date(
            '会议计划', ['将于9月30日举行新闻发布会。'],
            '2026-10-02'), '')
        self.assertEqual(j.infer_event_date(
            '会议', ['10月2日，陈云主持召开会议。'],
            '2026-10-03'), '2026-10-02')
        self.assertEqual(j.infer_event_date(
            '工作', ['9月30日，陈云主持召开会议。'],
            '2026-10-11'), '')  # More than the 10-day lag

    def test_list_crawl_includes_october_10_not_11(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)
            source = 'jjxw'
            def record(day, id_):
                return j.News(source, '陈云调研', f'https://www.jjxw.cn/article/{id_}.html', day)
            listing = [record('2026-10-11', 11),
                       record('2026-10-10', 10),
                       record('2026-10-01', 1),
                       record('2026-09-30', 30),
                       record('2026-09-01', 101),
                       record('2026-08-31', 831)]
            cfg = dict(j.config_defaults(), max_pages=3, delay_seconds=0)
            with j.connect_db(p / 'cache.sqlite3') as con:
                with patch.object(j, 'fetch_list', side_effect=[
                    (listing, True), ([], True)
                ]):
                    found = j.crawl(None, con, {'jjxw': set()}, cfg,
                                    2026, p / 'debug', '2026-09')
                    months = {n.pub_date for n in found['jjxw']}
            self.assertEqual(months,
                             {'2026-09-01', '2026-09-30',
                              '2026-10-01', '2026-10-10'})

    def test_next_year_cached_articles_have_body_requested(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)
            cfg = dict(j.config_defaults(), delay_seconds=0)
            with j.connect_db(p / 'db.sqlite3') as con:
                url = 'https://www.jjxw.cn/article/2027/01/04/90001.html'
                con.execute("INSERT INTO news(url,source,title,pub_date) VALUES(?,?,?,?)",
                            (url, 'jjxw', '陈云召开会议', '2027-01-04'))
                book = j.open_book(p / 'news.xlsx')
                article_html = article('陈云召开会议',
                    '12月31日，陈云主持召开市委常委会会议并部署重点工作。' * 6)
                with patch.object(j, 'fetch_html', return_value=article_html) as fetch:
                    j.fetch_article_details(j.build_session('agent'), con, cfg,
                        2026, '2026-12', {'jjxw': []}, book, 'month')
                    fetch.assert_called_once()
                actual = con.execute("SELECT status FROM news WHERE url=?", (url,)).fetchone()
                self.assertEqual(actual['status'], '成功')

    def test_late_news_requires_september_event_date_or_manual_review(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)
            cfg = dict(j.config_defaults(), names=['陈云'], delay_seconds=0)
            url = 'https://www.jjxw.cn/article/2026/10/05/91005.html'
            with j.connect_db(p / 'cache.sqlite3') as con:
                con.execute(
                    "INSERT INTO news(url,source,title,pub_date,article_title,body_json,status) "
                    "VALUES(?,?,?,?,?,?,?)",
                    (url, 'jjxw', '陈云部署工作', '2026-10-05',
                     '陈云部署工作',
                     json.dumps(['近日，陈云研究推进相关工作的情况。' * 8],
                                ensure_ascii=False), '成功'))
                book = j.open_book(p / 'data.xlsx')
                ws = book.active
                ws.cell(2, 1, '陈云部署工作')
                ws.cell(2, 1).hyperlink = url
                ws.cell(2, 2, '2026-10-05')

                chosen = j.refresh_book(book, con, ['陈云'], cfg)
                self.assertEqual(len(chosen), 1)
                self.assertEqual(chosen[0]['date'], '2026-10-05')
                self.assertEqual(j.deduplicate_for_month(chosen, '2026-09')[0], [])
                self.assertFalse(ws.cell(2, 8).value)
                # Manual verification: column H date plus column E include.
                ws.cell(2, 8, '2026-09-29')
                ws.cell(2, 5, '收录')
                chosen = j.refresh_book(book, con, ['陈云'], cfg)
                articles, failures = j.deduplicate_for_month(chosen, '2026-09')
                self.assertEqual((len(articles), len(failures)), (1, 0))
                self.assertEqual(articles[0]['date'], '2026-09-29')

    def test_cross_month_word_dedup_and_excel_report(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)
            results = p / 'results'
            results.mkdir()
            cfg = dict(j.config_defaults(), names=['陈云'], delay_seconds=0)
            body_a = '9月30日，陈云在九江主持召开重点项目建设推进会议。' * 6
            body_b = '10月3日，陈云主持召开新的专题工作会议。' * 6
            body_c = '近期，陈云对重点事项进行了研究部署。' * 7
            body_d = '9月30日，陈云在城区调研基层治理情况。' * 6
            data = [
                ('2026-09-30', '陈云研究项目甲', body_a),
                ('2026-10-02', '陈云推进项目乙', body_a),  # same body, dedup
                ('2026-10-03', '陈云主持其他会议', body_b),
                ('2026-10-04', '陈云部署工作', body_c),  # no event date
                ('2026-10-10', '陈云调研城区', body_d),
            ]
            with j.connect_db(results / 'cache.sqlite3') as con:
                book = j.open_book(results / 'jiujiang_news.xlsx')
                ws = book.active
                for index, (published, title, body) in enumerate(data, 2):
                    ident = 90000 + index
                    yy, mm, dd = published.split('-')
                    url = f'https://www.jjxw.cn/article/{yy}/{mm}/{dd}/{ident}.html'
                    ws.cell(index, 1, title)
                    ws.cell(index, 1).hyperlink = url
                    ws.cell(index, 2, published)

                    con.execute(
                        "INSERT INTO news(url,source,title,pub_date,article_title,body_json,status) "
                        "VALUES(?,?,?,?,?,?,?)",
                        (url, 'jjxw', title, published, title,
                         json.dumps([body], ensure_ascii=False), '成功'))
                chosen = j.refresh_book(book, con, ['陈云'], cfg)
                self.assertEqual(len(chosen), 5)
                self.assertEqual(ws.cell(3, 8).value, '2026-09-30')
                self.assertEqual(ws.cell(4, 8).value, '2026-10-03')
                self.assertFalse(ws.cell(5, 8).value)
                self.assertEqual(ws.cell(6, 8).value, '2026-09-30')
                book.save(results / 'jiujiang_news.xlsx')
                selected, failures = j.deduplicate_for_month(chosen, '2026-09')
                self.assertEqual(len(selected), 2)
                self.assertFalse(failures)
                self.assertTrue(j.build_word('2026-09', selected, results).exists())
            self.assertEqual(j.verify_month_output(
                '2026-09', results / 'jiujiang_news.xlsx', results, 10), 0)
            report = json.loads((results / '抓取检查报告.json').read_text('utf-8'))
            self.assertEqual(report['news'], 5)
            self.assertEqual(report['late_candidates'], 4)
            self.assertEqual(report['late_confirmed'], 2)
            self.assertEqual(report['word_articles'], 2)

    def test_cached_pages_do_not_cut_off_unseen_news(self):
        """Even 3+ known pages must not hide the older new target page."""
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)
            def rec(day, i):
                return j.News('jjxw', f'新闻{i}',
                              f'https://www.jjxw.cn/article/2026/09/20/{i}.html', day)
            old_known = [rec('2026-09-29', 1), rec('2026-09-28', 2),
                         rec('2026-09-27', 3)]
            missing = rec('2026-09-20', 4)
            pages = [
                ([old_known[0]], True), ([old_known[1]], True),
                ([old_known[2]], True), ([missing], True),
                ([rec('2026-08-31', 5)], True)
            ]
            with j.connect_db(p / 'cache.sqlite3') as con:
                cfg = dict(j.config_defaults(), max_pages=7, delay_seconds=0)
                known = {'jjxw': {x.url for x in old_known}}
                state = {}
                with patch.object(j, 'fetch_list', side_effect=pages) as fetch:
                    result = j.crawl(None, con, known, cfg, 2026, p,
                                     '2026-09', status=state)
                self.assertEqual(fetch.call_count, 5)
                self.assertTrue(state['complete'])
                self.assertEqual([x.url for x in result['jjxw']], [missing.url])

    def test_list_page_failure_preserves_excel_and_returns_failure(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)
            listing = j.News(
                'jjxw', '陈云主持会议',
                'https://www.jjxw.cn/article/2026/09/18/99101.html',
                '2026-09-18')
            # Successful first page, second page network failure.
            with patch.object(j, 'fetch_list',
                              side_effect=[([listing], True), ([], False)]), \
                 patch.object(j, 'fetch_html',
                              side_effect=AssertionError('body should not fetch')) as html:
                result = j.main(['2026-09'], base_dir=p)
            self.assertEqual(result, 1)
            html.assert_not_called()
            path = p / 'results' / 'jiujiang_news.xlsx'
            self.assertTrue(path.exists())
            wb = load_workbook(path)
            self.assertEqual(wb.active['A2'].value, '陈云主持会议')
            self.assertFalse(list((p / 'results').glob('*.docx')))

    def test_max_page_limit_is_incomplete_not_success(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)
            cfg = dict(j.config_defaults(), max_pages=2, delay_seconds=0)
            item = j.News('jjxw', '陈云主持会议',
                          'https://www.jjxw.cn/article/2026/09/18/99101.html',
                          '2026-09-18')
            with j.connect_db(p / 'db.sqlite3') as con:
                state = {}
                second = j.News('jjxw', '陈云参加会议',
                                    'https://www.jjxw.cn/article/2026/09/17/99102.html',
                                    '2026-09-17')
                with patch.object(j, 'fetch_list', side_effect=[
                    ([item], True), ([second], True)]):
                    j.crawl(None, con, {'jjxw': set()}, cfg, 2026,
                            p, '2026-09', status=state)
            self.assertFalse(state['complete'])
            self.assertIn('max_pages', state['error'])

    def test_auto_date_refresh_and_human_overrides_persist(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)
            cfg = dict(j.config_defaults(), names=['陈云'], delay_seconds=0)
            url = 'https://www.jjxw.cn/article/2026/10/02/99401.html'
            def body(day):
                return [f'{day}，陈云在九江主持召开会议并部署重点工作。' * 6]
            with j.connect_db(p / 'cache.sqlite3') as con:
                con.execute(
                    "INSERT INTO news(url,source,title,pub_date,article_title,body_json,status) "
                    "VALUES(?,?,?,?,?,?,?)",
                    (url, 'jjxw', '陈云部署工作', '2026-10-02',
                     '陈云部署工作', json.dumps(body('9月30日'), ensure_ascii=False),
                     '成功'))
                book = j.open_book(p / 'news.xlsx')
                ws = book.active
                ws.cell(2, 1, '陈云部署工作').hyperlink = url
                ws.cell(2, 2, '2026-10-02')

                def refresh():
                    return j.refresh_book(book, con, ['陈云'], cfg)
                refresh()
                cell = ws.cell(2, 8)
                self.assertEqual(cell.value, '2026-09-30')
                self.assertTrue(cell.comment.text.startswith('AUTO_EVENT_DATE:'))
                con.execute(
                    "UPDATE news SET body_json=? WHERE url=?",
                    (json.dumps(body('9月29日'), ensure_ascii=False), url))
                refresh()
                self.assertEqual(cell.value, '2026-09-29')
                # Explicit deletion of a machine guess requests human review.
                cell.value = None
                refresh()
                self.assertIsNone(cell.value)
                self.assertIn('MANUAL_EVENT_DATE_REVIEW', cell.comment.text)
                refresh()
                self.assertIsNone(cell.value)
                # A new manually chosen date must never be overwritten.
                cell.value = '2026-09-28'
                refresh()
                self.assertEqual(cell.value, '2026-09-28')
                self.assertIsNone(cell.comment)
                con.execute(
                    "UPDATE news SET body_json=? WHERE url=?",
                    (json.dumps(body('9月27日'), ensure_ascii=False), url))
                refresh()
                self.assertEqual(cell.value, '2026-09-28')

    def test_api_schema_change_does_not_look_like_last_page(self):
        """Valid HTTP 200 with unexpected JSON must fail, never signal EOF."""
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)
            class Response:
                content = b'{"success":true,"data":{"unknown_articles":[1,2]}}'
                def raise_for_status(self):
                    return None
                def __enter__(self):
                    return self
                def __exit__(self, *args):
                    return False

            session = j.build_session('test-agent')
            cfg = dict(j.config_defaults(), delay_seconds=0)
            with patch.object(session, 'get', return_value=Response()):
                articles, ok = j.fetch_list(session, 1, cfg, p)
            self.assertFalse(ok)
            self.assertFalse(articles)
            self.assertIn('无法确认', (p / 'jjxw_feed_page_2.txt').read_text('utf-8'))
            with patch.object(session, 'get', return_value=Response()):
                with j.connect_db(p / 'cache.sqlite3') as con:
                    state = {}
                    j.crawl(session, con, {'jjxw': set()}, cfg, 2026, p,
                            '2026-09', status=state)
                    self.assertFalse(state['complete'])

    def test_real_empty_page_still_stops_pagination(self):
        self.assertTrue(j.is_confirmed_empty_feed(b'{"data":{"list":[]}}'))
        self.assertTrue(j.is_confirmed_empty_feed(b'{"data":{"records":[]}}'))
        self.assertFalse(j.is_confirmed_empty_feed(b'{"data":{"unknown":[]}}'))
        self.assertFalse(j.is_confirmed_empty_feed(b'{"data":{"list":[{"id":4}]}}'))
        class Response:
            content = b'{"data":{"list":[]}}'
            def raise_for_status(self):
                return None
            def __enter__(self):
                return self
            def __exit__(self, *args):
                return False
        with tempfile.TemporaryDirectory() as folder:
            with patch.object(j.requests.Session, 'get', return_value=Response()):
                news, ok = j.fetch_list(
                    j.build_session('test'), 1,
                    j.config_defaults(), Path(folder))
            self.assertTrue(ok)
            self.assertEqual(news, [])

    def test_invalid_manual_event_date_blocks_export(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)
            cfg = dict(j.config_defaults(), names=['陈云'], delay_seconds=0)
            url = 'https://www.jjxw.cn/article/2026/09/18/99001.html'
            with j.connect_db(p / 'cache.sqlite3') as con:
                con.execute(
                    "INSERT INTO news(url,source,title,pub_date,body_json,status) "
                    "VALUES(?,?,?,?,?,?)",
                    (url, 'jjxw', '陈云主持会议', '2026-09-18',
                     json.dumps(['9月17日，陈云主持召开会议并部署重点工作。' * 6],
                                ensure_ascii=False), '成功'))
                book = j.open_book(p / 'news.xlsx')
                ws = book.active
                ws.cell(2, 1, '陈云主持会议').hyperlink = url
                ws.cell(2, 2, '2026-09-18')
                ws.cell(2, 5, '收录')

                ws.cell(2, 8, '2026-09-31')
                chosen = j.refresh_book(book, con, ['陈云'], cfg)
                self.assertEqual(chosen, [])
                self.assertEqual(ws.cell(2, 8).value, '2026-09-31')
                self.assertIn('日期无效', ws.cell(2, 7).value)
                ws.cell(2, 8, '2026-09-17')
                chosen = j.refresh_book(book, con, ['陈云'], cfg)
                self.assertEqual(len(chosen), 1)
                self.assertEqual(chosen[0]['date'], '2026-09-17')

    def test_repeated_api_page_fails_early_and_preserves_data(self):
        """An API that ignores pageNumber must not keep looping and claim success."""
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)
            cfg = dict(j.config_defaults(), max_pages=50, delay_seconds=0)
            entry = j.News('jjxw', '陈云召开会议',
                'https://www.jjxw.cn/article/2026/09/18/918001.html',
                '2026-09-18')
            status = {}
            with j.connect_db(p / 'db.sqlite3') as con:
                with patch.object(j, 'fetch_list',
                                  return_value=([entry], True)) as pages:
                    new = j.crawl(None, con, {'jjxw': set()}, cfg, 2026, p,
                                  '2026-09', status=status)
                stored = con.execute(
                    "SELECT count(*) FROM news WHERE source='jjxw'").fetchone()[0]
            self.assertEqual(pages.call_count, 2)
            self.assertFalse(status['complete'])
            self.assertIn('完全相同', status['error'])
            self.assertEqual(len(new['jjxw']), 1)
            self.assertEqual(stored, 1)

    def test_config_validation_catches_bad_values_before_network(self):
        good = j.config_defaults()
        samples = [
            ('max_pages', 0), ('max_pages', 'ten'),
            ('timeout_seconds', 0), ('delay_seconds', -1),
            ('delay_seconds', 'slow'), ('include_body_matches', 'true'),
            ('names', '陈云'), ('names', [42]),
        ]
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'config.json'
            for key, invalid in samples:
                cfg = dict(good, **{key: invalid})
                path.write_text(json.dumps(cfg, ensure_ascii=False), 'utf-8')
                with self.subTest(key=key, invalid=invalid):
                    with self.assertRaises(ValueError):
                        j.load_config(path)
            path.write_text(json.dumps(good, ensure_ascii=False), 'utf-8')
            self.assertEqual(j.load_config(path)['max_pages'], good['max_pages'])

    def test_empty_export_does_not_report_success(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)
            self.assertEqual(j.main(['2026-09', '--export'], base_dir=p), 1)
            self.assertFalse(list((p / 'results').glob('*.docx')))

    def test_repeated_export_applies_manual_exclusions_and_clears_stale_word(self):
        """Re-export must rebuild from E choices; no stale generated Word."""
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            results = root / 'results'
            results.mkdir()
            excel = results / 'jiujiang_news.xlsx'
            wb = j.open_book(excel)
            ws = wb.active
            with j.connect_db(results / 'jiujiang_cache.sqlite3') as con:
                for index, (title, body) in enumerate([
                    ('会议甲', '9月18日，市委主持召开会议甲并部署相关工作。' * 8),
                    ('会议乙', '9月19日，市委开展项目乙调研并研究相关事项。' * 8)
                ], 2):
                    url = f'https://www.jjxw.cn/article/2026/09/20/{index}.html'
                    ws.cell(index, 1, title).hyperlink = url
                    ws.cell(index, 2, '2026-09-20')
                    ws.cell(index, 5, '收录')
                    con.execute(
                        "INSERT INTO news(url,source,title,pub_date,article_title,body_json,status) "
                        "VALUES(?,?,?,?,?,?,?)",
                        (url, 'jjxw', title, '2026-09-20', title,
                         json.dumps([body], ensure_ascii=False), '成功'))
            j.save_book_atomic(wb, excel)
            word = results / '市委大事记（2026年9月）.docx'
            with patch.object(j, 'fetch_html',
                              side_effect=AssertionError('must use offline cache')):
                self.assertEqual(j.main(['2026-09', '--export'], base_dir=root), 0)
                self.assertTrue(word.exists())
                self.assertIn('会议乙', '\n'.join(p.text for p in Document(word).paragraphs))
                # Exclude only one article, without removing its Excel/cache record.
                wb = load_workbook(excel)
                wb.active['E3'] = '排除'
                wb.save(excel)
                self.assertEqual(j.main(['2026-09', '--export'], base_dir=root), 0)
                content = '\n'.join(p.text for p in Document(word).paragraphs)
                self.assertIn('会议甲', content)
                self.assertNotIn('会议乙', content)
                self.assertEqual(load_workbook(excel).active['E3'].value, '排除')
                # Now exclude the last item; generated old Word must disappear.
                wb = load_workbook(excel)
                wb.active['E2'] = '排除'
                wb.save(excel)
                self.assertEqual(j.main(['2026-09', '--export'], base_dir=root), 1)
            self.assertFalse(word.exists())
            self.assertTrue(list((results / '备份').glob(
                '市委大事记（2026年9月）_*.docx')))
            self.assertEqual(load_workbook(excel).active['E2'].value, '排除')
            self.assertEqual(load_workbook(excel).active['E3'].value, '排除')
            with j.connect_db(results / 'jiujiang_cache.sqlite3') as con:
                self.assertEqual(con.execute(
                    "SELECT count(*) FROM news WHERE source='jjxw'").fetchone()[0], 2)

    def test_empty_export_never_removes_user_authored_word(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            path = root / '市委大事记（2026年9月）.docx'
            doc = Document()
            doc.add_paragraph('人工编写的内容')
            doc.save(path)
            self.assertIsNone(j.build_word('2026-09', [], root))
            self.assertTrue(path.exists())
            self.assertIn('人工编写的内容',
                          '\n'.join(p.text for p in Document(path).paragraphs))

    def test_failed_selected_body_marks_partial_export(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)
            article = dict(source='jjxw', date='2026-09-20',
                           title='陈云在九江调研', url='https://www.jjxw.cn/a',
                           body=['陈云调研并部署相关工作。' * 8], status='成功')
            bad = dict(source='jjxw', date='2026-09-21',
                       title='陈云主持会议', url='https://www.jjxw.cn/b',
                       body=[], status='失败')
            with patch.object(j, 'fetch_article_details'), \
                 patch.object(j, 'refresh_book', return_value=[article, bad]):
                code = j.main(['2026-09', '--export'], base_dir=p)
            self.assertEqual(code, 1)
            self.assertTrue((p / 'results' / '市委大事记（2026年9月）.docx').exists())

    def test_existing_excel_does_not_overwrite_authoritative_cache(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)
            url = 'https://www.jjxw.cn/article/2026/09/19/100001.html'
            second = 'https://www.jjxw.cn/article/2026/09/20/100002.html'
            with j.connect_db(p / 'cache.sqlite3') as con:
                con.execute(
                    "INSERT INTO news(url,source,title,pub_date,article_title,"
                    "body_json,status) VALUES(?,?,?,?,?,?,?)",
                    (url, 'jjxw', '网站修订后的标题', '2026-09-19',
                     '陈云在九江调研', json.dumps(
                         ['9月17日，陈云在九江开展调研并主持会议。' * 8],
                         ensure_ascii=False), '成功'))
                book = j.open_book(p / 'news.xlsx')
                ws = book.active
                ws.cell(2, 1, 'Excel旧标题').hyperlink = url
                ws.cell(2, 2, '2026-09-18')  # old, not authoritative
                ws.cell(2, 5, '排除')

                ws.cell(2, 8, '2026-09-17')  # verified event date
                ws.cell(3, 1, 'Excel补录新文章').hyperlink = second
                ws.cell(3, 2, '2026-09-20')

                urls = j.existing_rows(book, con)
                self.assertEqual(urls['jjxw'], {url, second})
                existing = con.execute(
                    "SELECT * FROM news WHERE url=?", (url,)).fetchone()
                self.assertEqual(existing['title'], '网站修订后的标题')
                self.assertEqual(existing['pub_date'], '2026-09-19')
                self.assertEqual(existing['status'], '成功')
                fresh = con.execute(
                    "SELECT title,pub_date FROM news WHERE url=?",
                    (second,)).fetchone()
                self.assertEqual((fresh['title'], fresh['pub_date']),
                                 ('Excel补录新文章', '2026-09-20'))
                chosen = j.refresh_book(book, con, ['陈云'], j.config_defaults())
                self.assertEqual(ws.cell(2, 2).value, '2026-09-19')
                self.assertEqual(ws.cell(2, 8).value, '2026-09-17')
                self.assertEqual(ws.cell(2, 5).value, '排除')
                self.assertFalse(any(a['url'] == url for a in chosen))

    def test_atomic_excel_save_retains_original_after_failed_serialization(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder) / 'jiujiang_news.xlsx'
            book = j.open_book(p)
            book.active.cell(2, 1, '人工保留的标题')
            book.active.cell(2, 5, '收录')
            j.save_book_atomic(book, p)
            baseline = p.read_bytes()
            book.active.cell(2, 1, '替换途中失败的数据')

            def fail_mid_save(temp_path):
                Path(temp_path).write_bytes(b'partial invalid XLSX')
                raise OSError('模拟保存失败')

            with patch.object(book, 'save', side_effect=fail_mid_save):
                with self.assertRaises(OSError):
                    j.save_book_atomic(book, p)
            self.assertEqual(p.read_bytes(), baseline)
            self.assertEqual(load_workbook(p).active['E2'].value, '收录')
            self.assertFalse(list(Path(folder).glob('.jiujiang_news.*.xlsx')))

    def test_atomic_excel_success_retains_manual_review(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder) / 'jiujiang_news.xlsx'
            book = j.open_book(p)
            book.active.cell(2, 5, '收录')
            book.active.cell(2, 8, '2026-09-30')
            book.active.cell(2, 8).comment = j.Comment(
                'MANUAL_EVENT_DATE_REVIEW', '人工日期标记')
            j.save_book_atomic(book, p)
            reopened = load_workbook(p)
            self.assertEqual(reopened.sheetnames, ['九江新闻网'])
            self.assertEqual(reopened.active['E2'].value, '收录')
            self.assertEqual(reopened.active['H2'].value, '2026-09-30')
            self.assertEqual(reopened.active['H2'].comment.text,
                             'MANUAL_EVENT_DATE_REVIEW')

    def test_batch_names_pasted_from_excel_column(self):
        pasted = '\ufeff姓名\n陈云\n黄强\n陈云\n欧阳修\n# 第一组\n阿布都·拉西提\n'
        names = j.parse_names_text(pasted)
        self.assertEqual(names, ['陈云', '黄强', '欧阳修', '阿布都·拉西提'])
        self.assertEqual(j.get_names({'names': names}), names)
        self.assertEqual(j.match_names('黄强主持会议', '', names, False),
                         ['黄强'])

    def test_batch_names_mixed_separators_and_duplicate_dedup(self):
        pasted = ('姓名\n陈云、邓永翔，黄强;庄兆林；陈永奇\t黄喜忠\n'
                  '陈云,黄强\n＃复制自通讯录\n')
        self.assertEqual(j.parse_names_text(pasted),
                         ['陈云', '邓永翔', '黄强', '庄兆林',
                          '陈永奇', '黄喜忠'])

    def test_names_file_precedes_config_without_mutating_json(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            path = root / 'config.json'
            stored = dict(j.config_defaults(), names=['旧名单'])
            original = json.dumps(stored, ensure_ascii=False)
            path.write_text(original, encoding='utf-8')
            (root / 'names.txt').write_text(
                '姓名\n新人物甲\n新人物乙\n新人物甲\n', encoding='utf-8-sig')
            cfg = j.load_config(path)
            self.assertEqual(j.get_names(cfg), ['新人物甲', '新人物乙'])
            self.assertEqual(path.read_text('utf-8'), original)
            self.assertEqual(j.match_names('新人物乙主持会议', '', j.get_names(cfg), False),
                             ['新人物乙'])
            (root / 'names.txt').unlink()
            self.assertEqual(j.get_names(j.load_config(path)), ['旧名单'])

    def test_empty_or_malformed_names_file_fails_without_network(self):
        examples = ['\n# 暂无名单\n', '张三 李四', '张三,主持会议！',
                    '张三,王1', '1.张三', '姓名\n']
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / 'config.json').write_text(
                json.dumps(j.config_defaults(), ensure_ascii=False), 'utf-8')
            for content in examples:
                (root / 'names.txt').write_text(content, encoding='utf-8')
                with self.subTest(content=content):
                    with patch.object(j, 'fetch_list',
                                      side_effect=AssertionError('should be offline')):
                        with self.assertRaisesRegex(ValueError, 'names.txt'):
                            j.main(['2026-09', '--sync'], base_dir=root)
            self.assertFalse((root / 'results' / 'jiujiang_cache.sqlite3').exists())

    def test_export_rechecks_cached_articles_with_updated_batch_names(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            cfg = dict(j.config_defaults(), delay_seconds=0,
                       names=['没有出现在正文中的人'])
            (root / 'config.json').write_text(
                json.dumps(cfg, ensure_ascii=False), encoding='utf-8')
            news = '黄强'
            url = 'https://www.jjxw.cn/article/2026/09/18/909999.html'
            output = root / 'results'
            output.mkdir()
            with j.connect_db(output / 'jiujiang_cache.sqlite3') as con:
                content = ['9月17日，黄强在九江主持召开工作会议。' * 8]
                con.execute(
                    "INSERT INTO news(url,source,title,pub_date,article_title,"
                    "body_json,status) VALUES(?,?,?,?,?,?,?)",
                    (url, 'jjxw', f'{news}主持工作会议', '2026-09-18',
                     f'{news}主持工作会议',
                     json.dumps(content, ensure_ascii=False), '成功'))
            book = j.open_book(output / 'jiujiang_news.xlsx')
            ws = book.active
            ws.cell(2, 1, f'{news}主持工作会议').hyperlink = url
            ws.cell(2, 2, '2026-09-18')

            j.save_book_atomic(book, output / 'jiujiang_news.xlsx')
            (root / 'names.txt').write_text('黄强\n邓永翔\n', 'utf-8')
            with patch.object(j, 'fetch_html',
                              side_effect=AssertionError('cache should be used')) as fetch:
                code = j.main(['2026-09', '--export'], base_dir=root)
            self.assertEqual(code, 0)
            fetch.assert_not_called()
            wb = load_workbook(output / 'jiujiang_news.xlsx')
            self.assertEqual(wb.active['C2'].value, '命中')
            self.assertEqual(wb.active['D2'].value, '黄强')
            self.assertTrue((output / '市委大事记（2026年9月）.docx').exists())

    def test_names_xlsx_detects_header_in_non_first_column(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)
            wb = Workbook()
            ws = wb.active
            ws.append(['部门', '职务', '姓名'])
            ws.append(['办公室', '主任', '黄强'])
            ws.append(['办公室', '副主任', '陈云'])
            ws.append(['办公室', '副主任', '黄强'])
            ws.append(['办公室', '职工', None])
            wb.save(p / 'names.xlsx')
            self.assertEqual(
                j.read_names_excel(p / 'names.xlsx'), ['黄强', '陈云'])
            (p / 'config.json').write_text(
                json.dumps(dict(j.config_defaults(), names=['老名单']),
                           ensure_ascii=False), encoding='utf-8')
            (p / 'names.txt').write_text('另一个名单\n', encoding='utf-8')
            cfg = j.load_config(p / 'config.json')
            self.assertEqual(cfg['names_source'], 'names.xlsx')
            self.assertEqual(j.get_names(cfg), ['黄强', '陈云'])

    def test_names_xlsx_without_header_reads_column_a(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder) / 'names.xlsx'
            wb = Workbook()
            ws = wb.active
            ws.append(['陈云', '项目科'])
            ws.append(['邓永翔', '会议组'])
            ws.append(['陈云', '重复'])
            wb.save(p)
            self.assertEqual(j.read_names_excel(p), ['陈云', '邓永翔'])

    def test_names_xlsx_header_in_second_row(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder) / 'names.xlsx'
            wb = Workbook()
            ws = wb.active
            ws.append(['名单'])
            ws.append(['序号', '姓名', '部门'])
            ws.append([1, '庄兆林', '市委'])
            ws.append([2, '黄喜忠', '组织'])
            wb.save(p)
            self.assertEqual(j.read_names_excel(p), ['庄兆林', '黄喜忠'])

    def test_names_xlsx_invalid_cells_include_coordinates(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder) / 'names.xlsx'
            wb = Workbook()
            ws = wb.active
            ws.append(['序号', '姓名'])
            ws.append([1, '陈云'])
            ws.append([2, '主任 李四'])
            wb.save(p)
            with self.assertRaisesRegex(ValueError, 'B3'):
                j.read_names_excel(p)
            wb = Workbook()
            ws = wb.active
            ws.append(['姓名', '姓名'])
            ws.append(['陈云', '黄强'])
            wb.save(p)
            with self.assertRaisesRegex(ValueError, '多个'):
                j.read_names_excel(p)
            ws.delete_cols(2)
            ws.delete_rows(2)
            wb.save(p)
            with self.assertRaisesRegex(ValueError, '未找到有效姓名'):
                j.read_names_excel(p)

    def test_check_names_does_not_write_any_files_or_connect(self):
        import io
        from contextlib import redirect_stdout
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)
            wb = Workbook()
            wb.active.append(['姓名'])
            wb.active.append(['陈云'])
            wb.active.append(['黄强'])
            wb.save(p / 'names.xlsx')
            with patch.object(j, 'fetch_list', side_effect=AssertionError('network')), \
                 patch.object(j, 'connect_db', side_effect=AssertionError('database')):
                output = io.StringIO()
                with redirect_stdout(output):
                    rc = j.main(['--check-names'], base_dir=p)
            self.assertEqual(rc, 0)
            self.assertIn('生效姓名：2 人', output.getvalue())
            self.assertIn('姓名来源：names.xlsx', output.getvalue())
            self.assertFalse((p / 'results').exists())
            self.assertFalse((p / 'config.json').exists())

    def test_local_excel_switch_to_txt_without_changing_json(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)
            stored = dict(j.config_defaults(), names=['旧人员'])
            (p / 'config.json').write_text(
                json.dumps(stored, ensure_ascii=False), encoding='utf-8')
            wb = Workbook()
            wb.active.append(['姓名'])
            wb.active.append(['黄强'])
            wb.save(p / 'names.xlsx')
            (p / 'names.txt').write_text('陈云\n', encoding='utf-8')
            self.assertEqual(j.load_config(p / 'config.json')['names'], ['黄强'])
            (p / 'names.xlsx').unlink()
            self.assertEqual(j.load_config(p / 'config.json')['names'], ['陈云'])
            (p / 'names.txt').unlink()
            self.assertEqual(j.load_config(p / 'config.json')['names'], ['旧人员'])

    def test_leaders_table_only_four_sections_and_normalizes_spaces(self):
        markup = """
        <html><body><a href="#">不应收录</a>
        <table width="98%" border="0" align="center" cellpadding="0"
               cellspacing="5" style="line-height:30px;">
          <tr><td>省委</td>
              <td><a>尹　弘</a> <a>叶建春</a> <a>陈永奇</a></td></tr>
          <tr><td><img alt="人大"></td>
              <td><a>尹　弘</a> <a>吴忠琼</a></td></tr>
          <tr><td>政府</td>
              <td><a>叶建春</a><a>孙洪山</a></td></tr>
          <tr><td>政协</td>
              <td><a>宋福龙</a><a>陈俊卿</a></td></tr>
          <tr><td>人事任免</td><td><a>旧新闻标题</a></td></tr>
        </table>
        <table><tr><td><a>另一则新闻</a></td></tr></table>
        </body></html>
        """
        groups = j.extract_leader_groups(markup)
        self.assertEqual(groups['省委'], ['尹弘', '叶建春', '陈永奇'])
        self.assertEqual(groups['人大'], ['尹弘', '吴忠琼'])
        self.assertEqual(groups['政府'], ['叶建春', '孙洪山'])
        self.assertEqual(groups['政协'], ['宋福龙', '陈俊卿'])
        self.assertNotIn('旧新闻标题', j.render_leader_names(groups))
        self.assertEqual(len(set(sum(groups.values(), []))), 7)

    def test_leaders_wrong_table_or_missing_group_rejected(self):
        bad_markup = '<html><table><tr><td>省委<a>张三</a></td></tr></table></html>'
        with self.assertRaisesRegex(ValueError, '指定表格'):
            j.extract_leader_groups(bad_markup)
        incomplete = """
        <table width="98%" border="0" align="center" cellpadding="0"
               cellspacing="5" style="line-height:30px;">
          <tr><td>省委</td><td><a>尹弘</a></td></tr>
          <tr><td>人大</td><td><a>吴忠琼</a></td></tr>
          <tr><td>政府</td><td><a>叶建春</a></td></tr>
        </table>
        """
        with self.assertRaisesRegex(ValueError, '政协'):
            j.extract_leader_groups(incomplete)

    def test_update_leaders_creates_auto_file_and_is_idempotent(self):
        markup = """
        <table width="98%" border="0" align="center" cellpadding="0"
               cellspacing="5" style="line-height: 30px;">
        <tr><td>省委</td><td><a>尹弘</a><a>陈永奇</a></td></tr>
        <tr><td>人大</td><td><a>吴忠琼</a></td></tr>
        <tr><td>政府</td><td><a>叶建春</a><a>陈永奇</a></td></tr>
        <tr><td>政协</td><td><a>宋福龙</a></td></tr>
        </table>"""
        class Response:
            content = markup
            def raise_for_status(self):
                return None
        class Session:
            def get(self, url, **kwargs):
                self.assert_url = url
                return Response()
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)
            groups, changed = j.update_leaders(p, session=Session())
            self.assertTrue(changed)
            auto = p / 'names.auto.txt'
            text = auto.read_text('utf-8')
            self.assertTrue(text.startswith('# 江西省领导'))
            self.assertEqual(text.count('\n陈永奇\n'), 2)  # grouped
            self.assertEqual(len(groups), 4)
            saved = auto.stat().st_mtime_ns
            groups2, changed2 = j.update_leaders(p, session=Session())
            self.assertFalse(changed2)
            self.assertEqual(groups, groups2)
            self.assertEqual(auto.stat().st_mtime_ns, saved)

    def test_failed_leaders_update_keeps_last_good_roster(self):
        class BrokenResponse:
            content = b'<html>upstream site changed</html>'
            def raise_for_status(self):
                return None
        class Session:
            def get(self, *_args, **_kwargs):
                return BrokenResponse()
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)
            old = '# 省委\n尹弘\n# 人大\n吴忠琼\n# 政府\n叶建春\n# 政协\n宋福龙\n'
            auto = p / 'names.auto.txt'
            auto.write_text(old, encoding='utf-8')
            with self.assertRaisesRegex(ValueError, '指定表格'):
                j.update_leaders(p, session=Session())
            self.assertEqual(auto.read_text('utf-8'), old)

    def test_manual_names_and_auto_names_union_without_duplicates(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)
            (p / 'config.json').write_text(json.dumps(
                dict(j.config_defaults(), names=['默认人物']), ensure_ascii=False), 'utf-8')
            auto = '# 省委\n尹弘\n陈永奇\n# 人大\n尹弘\n# 政府\n叶建春\n# 政协\n宋福龙\n'
            (p / 'names.auto.txt').write_text(auto, encoding='utf-8')
            cfg = j.load_config(p / 'config.json')
            self.assertEqual(cfg['names_source'], 'names.auto.txt')
            self.assertEqual(cfg['names'], ['尹弘', '陈永奇', '叶建春', '宋福龙'])
            (p / 'names.txt').write_text('陈永奇\n黄强\n', 'utf-8')
            cfg = j.load_config(p / 'config.json')
            self.assertEqual(cfg['names'], ['陈永奇', '黄强', '尹弘', '叶建春', '宋福龙'])
            self.assertEqual(cfg['names_source'], 'names.txt + names.auto.txt')
            (p / 'names.txt').unlink()
            wb = Workbook()
            wb.active.append(['姓名'])
            wb.active.append(['自定义人员'])
            wb.save(p / 'names.xlsx')
            cfg = j.load_config(p / 'config.json')
            self.assertEqual(cfg['names'],
                             ['自定义人员', '尹弘', '陈永奇', '叶建春', '宋福龙'])
            self.assertEqual(cfg['names_source'], 'names.xlsx + names.auto.txt')

    def test_update_leaders_cli_only_writes_auto_file(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)
            (p / 'names.txt').write_text('人工甲\n', 'utf-8')
            fake = {
                '省委': ['尹弘'], '人大': ['吴忠琼'],
                '政府': ['叶建春'], '政协': ['宋福龙'],
            }
            def stub(root):
                (root / 'names.auto.txt').write_text(
                    j.render_leader_names(fake), encoding='utf-8')
                return fake, True
            with patch.object(j, 'update_leaders', side_effect=stub):
                self.assertEqual(j.main(['--update-leaders'], base_dir=p), 0)
            self.assertEqual((p / 'names.txt').read_text('utf-8'), '人工甲\n')
            self.assertFalse((p / 'results').exists())
            self.assertFalse((p / 'config.json').exists())

    def test_check_names_warns_when_auto_roster_is_unverified_snapshot(self):
        import io
        from contextlib import redirect_stdout
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / 'names.auto.txt').write_text(
                '# 初始网页快照：待复核\n'
                '# 省委\n尹弘\n'
                '# 人大\n吴忠琼\n'
                '# 政府\n叶建春\n'
                '# 政协\n宋福龙\n', encoding='utf-8')
            output = io.StringIO()
            with redirect_stdout(output):
                self.assertEqual(j.main(['--check-names'], base_dir=root), 0)
            self.assertIn('尚未通过实时抓取复核', output.getvalue())

    def test_legacy_excel_migration_preserves_manual_fields_and_links(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder) / 'legacy.xlsx'
            old = Workbook()
            ws = old.active
            ws.title = '九江新闻网'
            ws.append(['新闻标题', '发布日期', '自动匹配', '命中姓名',
                       '人工选择', '正文状态', '原文链接', '说明', '事件日期'])
            url = 'https://www.jjxw.cn/article/2026/09/18/101.html'
            ws['A2'] = '旧标题'
            ws['B2'] = '2026-09-18'
            ws['E2'] = '排除'
            ws['G2'] = url
            ws['H2'] = '已有人工备注'
            ws['H2'] = '2026-09-17'
            ws['H2'].comment = j.Comment('手工核对', '复核人员')
            ws['A2'].fill = PatternFill('solid', fgColor='FFFF00')
            old.save(p)
            book = j.open_book(p)
            current = book.active
            self.assertEqual(current.max_column, 8)
            self.assertEqual([current.cell(1, c).value for c in range(1, 9)],
                             j.HEADERS)
            self.assertEqual(current['A2'].hyperlink.target, url)
            self.assertEqual(current['G2'].value, '已有人工备注')
            self.assertEqual(current['H2'].value, '2026-09-17')
            self.assertEqual(current['H2'].comment.text, '手工核对')
            self.assertEqual(current['E2'].value, '排除')
            self.assertTrue(j.is_yellow(current['A2']))
            j.save_book_atomic(book, p)
            reopened = j.open_book(p)
            self.assertEqual(reopened.active.max_column, 8)
            self.assertEqual(reopened.active['A2'].hyperlink.target, url)

    def test_legacy_conflict_report_survives_column_removal(self):
        import csv
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            book = Workbook()
            ws = book.active
            ws.title = '九江新闻网'
            ws.append(['新闻标题', '发布日期', '自动匹配', '命中姓名',
                       '人工选择', '正文状态', '原文链接', '说明', '事件日期'])
            a_url = 'https://www.jjxw.cn/article/2026/09/18/100.html'
            old_g_url = 'https://www.jjxw.cn/article/2026/09/18/200.html'
            ws['A2'] = '两条链接待核对'
            ws['A2'].hyperlink = a_url
            ws['B2'] = '2026-09-18'
            ws['E2'] = '收录'
            ws['G2'] = old_g_url
            ws['H2'] = '历史说明'
            ws['H2'] = '2026-09-17'
            path = root / 'news.xlsx'
            book.save(path)
            migrated = j.open_book(path)
            ws = migrated.active
            self.assertEqual(ws.max_column, 8)
            self.assertIn(j.LEGACY_LINK_NOTE + old_g_url, ws['G2'].value)
            with j.connect_db(root / 'db.sqlite3') as con:
                issues = []
                known = j.existing_rows(migrated, con, issues)
                self.assertEqual(known['jjxw'], {a_url, old_g_url})
                self.assertEqual(len(issues), 1)
                self.assertEqual(issues[0]['row'], 2)
                j.write_new_rows(migrated, {'jjxw': [
                    j.News('jjxw', '原A', a_url, '2026-09-18'),
                    j.News('jjxw', '原G', old_g_url, '2026-09-18')]})
                self.assertEqual(ws.max_row, 2)
                report = root / '链接冲突待复核.csv'
                j.write_link_review_report(issues, report)
                with report.open(encoding='utf-8-sig', newline='') as stream:
                    rows = list(csv.reader(stream))
                self.assertEqual(rows[1][2:4], [a_url, old_g_url])
                # A becomes the canonical original link once manually checked.
                ws['A2'].hyperlink = old_g_url
                self.assertEqual(j.reconcile_row_link(
                    [ws.cell(2, c) for c in range(1, 9)],
                    j.SOURCES['jjxw']['base'], con), old_g_url)
                self.assertNotIn(j.LEGACY_LINK_NOTE, ws['G2'].value)
                self.assertEqual(ws['H2'].value, '2026-09-17')
                self.assertEqual(ws['E2'].value, '收录')

    def test_old_formula_link_becomes_native_without_losing_yellow(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'formula.xlsx'
            book = j.open_book(path)
            ws = book.active
            url = 'https://www.jjxw.cn/article/2026/09/18/222.html'
            ws['A2'] = f'=HYPERLINK("{url}","市委召开会议")'
            ws['B2'] = '2026-09-18'
            ws['A2'].fill = PatternFill('solid', fgColor='FFFF00')
            with j.connect_db(Path(folder) / 'db.sqlite3') as con:
                j.existing_rows(book, con)
            self.assertEqual(ws['A2'].value, '市委召开会议')
            self.assertEqual(ws['A2'].hyperlink.target, url)
            self.assertTrue(j.is_yellow(ws['A2']))

    def test_incremental_excel_append_has_only_title_links(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder) / 'news.xlsx'
            book = j.open_book(p)
            ws = book.active
            a = 'https://www.jjxw.cn/article/2026/09/18/1001.html'
            b = 'https://www.jjxw.cn/article/2026/10/03/1002.html'
            c = 'https://www.jjxw.cn/article/2026/09/30/1003.html'
            ws['A2'] = '原有记录'
            ws['A2'].hyperlink = a
            ws['B2'] = '2026-09-18'
            ws['E2'] = '收录'
            ws['H2'] = '2026-09-17'
            ws['H2'].comment = j.Comment('手动校对', '核验')
            items = {'jjxw': [
                j.News('jjxw', '迟发报道', b, '2026-10-03'),
                j.News('jjxw', '月底报道', c, '2026-09-30'),
                j.News('jjxw', '原有记录', a, '2026-09-18')]}
            j.write_new_rows(book, items)
            j.write_new_rows(book, items)
            self.assertEqual(ws.max_row, 4)
            self.assertEqual(ws.max_column, 8)
            self.assertEqual(ws['E2'].value, '收录')
            self.assertEqual(ws['H2'].value, '2026-09-17')
            self.assertEqual(ws['H2'].comment.text, '手动校对')
            self.assertEqual([ws.cell(i, 2).value for i in range(2, 5)],
                             ['2026-09-18', '2026-09-30', '2026-10-03'])
            for i in range(2, 5):
                self.assertTrue(ws.cell(i, 1).hyperlink.target)
                self.assertIsNone(ws.cell(i, 7).hyperlink)
            self.assertEqual(ws.auto_filter.ref, 'A1:H4')

    def test_manual_selection_overrides_keyword_auto_match(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            book = j.open_book(root / 'news.xlsx')
            ws = book.active
            with j.connect_db(root / 'db.sqlite3') as con:
                data = [
                    ('普通市政报道', '收录', '无领导名字但由人工决定收录。'),
                    ('陈云出席活动', '排除', '虽然命中姓名但人工决定排除。'),
                    ('陈云主持会议', None, '自动命中领导姓名。')]
                for i, (title, decision, body) in enumerate(data, 2):
                    url = f'https://www.jjxw.cn/article/2026/09/18/{i}.html'
                    ws.cell(i, 1, title).hyperlink = url
                    ws.cell(i, 2, '2026-09-18')
                    ws.cell(i, 5, decision)
                    con.execute(
                        "INSERT INTO news(url,source,title,pub_date,article_title,body_json,status) "
                        "VALUES(?,?,?,?,?,?,?)",
                        (url, 'jjxw', title, '2026-09-18', title,
                         json.dumps([body], ensure_ascii=False), '成功'))
                chosen = j.refresh_book(book, con, ['陈云'],
                                         j.config_defaults())
                self.assertEqual(len(chosen), 2)
                self.assertEqual({c['title'] for c in chosen},
                                 {'普通市政报道', '陈云主持会议'})
                self.assertEqual(ws['E2'].value, '收录')
                self.assertEqual(ws['E3'].value, '排除')

    def test_month_subtitle_outline_level_two(self):
        with tempfile.TemporaryDirectory() as folder:
            path = j.build_word('2026-09', [{
                'title': '市委专题会议', 'body': ['会议部署了工作。']}],
                Path(folder))
            doc = Document(path)
            headings = [(p.text, [x.get(j.qn('w:val')) for x in
                                  p._p.xpath('./w:pPr/w:outlineLvl')])
                        for p in doc.paragraphs
                        if p._p.xpath('./w:pPr/w:outlineLvl')]
            self.assertEqual(headings[0][1], ['0'])
            self.assertEqual(headings[1], ('（2026年9月）', ['1']))
            self.assertEqual(headings[2], ('市委专题会议', ['2']))

    def test_csv_link_report_escapes_formula(self):
        import csv
        with tempfile.TemporaryDirectory() as folder:
            report = Path(folder) / '链接冲突待复核.csv'
            j.write_link_review_report([{
                'row': 2, 'title': '=HYPERLINK("url","test")',
                'a_url': 'https://www.jjxw.cn/article/a.html',
                'g_url': 'https://www.jjxw.cn/article/b.html',
                'reason': 'legacy mismatch'}], report)
            with report.open(encoding='utf-8-sig', newline='') as stream:
                data = list(csv.reader(stream))
            self.assertTrue(data[1][1].startswith("'=HYPERLINK("))

    def test_verify_detects_pending_link_review_report(self):
        with tempfile.TemporaryDirectory() as folder:
            out = Path(folder)
            wb = j.open_book(out / 'jiujiang_news.xlsx')
            ws = wb.active
            url = 'https://www.jjxw.cn/article/2026/09/18/101.html'
            ws['A2'] = '新闻'
            ws['A2'].hyperlink = url
            ws['B2'] = '2026-09-18'
            wb.save(out / 'jiujiang_news.xlsx')
            j.write_link_review_report([{
                'row': 3, 'title': '另一篇待审报道',
                'a_url': url, 'g_url': url, 'reason': '待审'}],
                out / '链接冲突待复核.csv')
            self.assertEqual(j.verify_month_output(
                '2026-09', out / 'jiujiang_news.xlsx', out), 1)
            report = json.loads((out / '抓取检查报告.json').read_text('utf-8'))
            self.assertTrue(any('链接冲突' in message
                                for message in report['errors']))

    def test_daily_backup_caps_seven_and_keeps_manually_saved_files(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder) / 'jiujiang_news.xlsx'
            j.save_book_atomic(j.open_book(p), p)
            dest = p.parent / '备份'
            dest.mkdir()
            for day in range(1, 13):
                (dest / f'jiujiang_news_运行前_202509{day:02d}_121200_000001.xlsx').write_bytes(b'old')
            manual = dest / '手动整理的重要存档.xlsx'
            manual.write_bytes(b'never delete')
            j.backup_spreadsheet(p)
            copies = sorted(dest.glob('jiujiang_news_运行前_*.xlsx'))
            self.assertEqual(len(copies), 7)
            self.assertTrue(manual.exists())
            first_run = {item.name for item in copies}
            j.backup_spreadsheet(p)  # same calendar day
            self.assertEqual(first_run,
                             {item.name for item in dest.glob(
                                 'jiujiang_news_运行前_*.xlsx')})
            self.assertTrue(all(item.exists() for item in copies))

    def test_excel_audit_counts_months_without_modifying_file(self):
        import io
        from contextlib import redirect_stdout
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            output = root / 'results'
            output.mkdir()
            path = output / 'jiujiang_news.xlsx'
            book = j.open_book(path)
            ws = book.active
            for i, day in enumerate(['2026-09-18', '2026-09-30',
                                     '2026-10-02'], 2):
                url = f'https://www.jjxw.cn/article/{i}.html'
                ws.cell(i, 1, f'新闻{i}').hyperlink = url
                ws.cell(i, 2, day)
            j.save_book_atomic(book, path)
            original = path.read_bytes()
            report = j.audit_excel(path)
            self.assertEqual(report['total'], 3)
            self.assertEqual(report['months'], {'2026-09': 2, '2026-10': 1})
            self.assertEqual(report['issue_count'], 0)
            with patch.object(j, 'fetch_html',
                              side_effect=AssertionError('offline')), \
                 patch.object(j, 'connect_db',
                              side_effect=AssertionError('no db')):
                output_text = io.StringIO()
                with redirect_stdout(output_text):
                    self.assertEqual(j.main(['--check-excel'], base_dir=root), 0)
            self.assertIn('2026-09：2 条', output_text.getvalue())
            self.assertEqual(path.read_bytes(), original)

    def test_excel_audit_checks_title_links_and_h_event_dates(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'news.xlsx'
            book = j.open_book(path)
            ws = book.active
            url = 'https://www.jjxw.cn/article/2026/09/18/100.html'
            ws['A2'] = '没有链接的报道'
            ws['B2'] = '2026-09-18'
            ws['A3'] = '另一个报道'
            ws['A3'].hyperlink = url
            ws['B3'] = '2026-09-19'
            ws['A4'] = '又一篇报道'
            ws['A4'].hyperlink = url
            ws['B4'] = '日期有误'
            ws['H4'] = '2026-09-31'
            book.save(path)
            report = j.audit_excel(path)
            details = '\n'.join(report['issues'])
            self.assertIn('A列没有有效', details)
            self.assertIn('第3行', details)
            self.assertIn('发布日期缺失或无效', details)
            self.assertIn('H列事件日期无效', details)
            self.assertGreaterEqual(report['issue_count'], 4)

    def test_check_excel_missing_workbook_is_read_only(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            self.assertEqual(j.main(['--check-excel'], base_dir=root), 1)
            self.assertFalse((root / 'results').exists())

    def test_generated_word_backup_caps_seven_and_preserves_manual_files(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            path = root / '市委大事记（2026年9月）.docx'
            Document().save(path)
            backup_dir = root / '备份'
            backup_dir.mkdir()
            prefix = path.stem + '_'
            for day in range(1, 13):
                (backup_dir / f'{prefix}202509{day:02d}_121200.docx').write_bytes(b'old')
            manual = backup_dir / '我的手工修订终稿.docx'
            manual.write_bytes(b'keep forever')
            j.backup_generated_word(path)
            generated = list(backup_dir.glob(prefix + '*.docx'))
            self.assertEqual(len(generated), 7)
            names = {f.name for f in generated}
            j.backup_generated_word(path)
            self.assertEqual(names, {f.name for f in backup_dir.glob(prefix + '*.docx')})
            self.assertEqual(manual.read_bytes(), b'keep forever')

    def test_news_json_api(self):
        obj = {'data': {'list': [
            {'fileID': 190001, 'publishTime': '2026-09-18 10:00:00',
             'title': '陈云调研'},
            {'fileID': 190002, 'publishTime': '2026-09-17 10:00:00',
             'title': '普通报道'}
        ]}}
        rows = j.parse_jjxw_json(obj)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0][2], '2026-09-18')
        self.assertEqual(rows[0][1],
                         'https://www.jjxw.cn/article/2026/09/18/190001.html')

    def test_verify_missing_artifacts_does_not_create_database(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)
            with patch.object(j, 'fetch_html', side_effect=AssertionError('offline')):
                result = j.main(['2026-09', '--verify'], base_dir=p)
            self.assertEqual(result, 1)
            self.assertTrue((p / 'results' / '抓取检查报告.json').exists())
            self.assertFalse((p / 'results' / 'jiujiang_cache.sqlite3').exists())

    def test_short_name_is_not_matched_inside_a_configured_longer_name(self):
        names = ['陈云', '陈云飞', '黄强']
        self.assertEqual(
            j.match_names('陈云飞主持会议', '', names, True), ['陈云飞'])
        # A standalone occurrence of the short name elsewhere still counts.
        self.assertEqual(
            j.match_names('陈云飞主持会议', '陈云随后发言', names, True),
            ['陈云', '陈云飞'])
        # Preserve the configured order in the reported matches.
        self.assertEqual(
            j.match_names('陈云飞参加会议', '', ['陈云飞', '陈云'], True),
            ['陈云飞'])

    def test_short_name_remains_matchable_when_no_longer_name_is_configured(self):
        self.assertEqual(
            j.match_names('陈云飞参加会议', '', ['陈云'], True), ['陈云'])
        self.assertEqual(
            j.match_names('没有提及名单人物', '陈云飞参加会议',
                          ['陈云'], False), [])

    def test_check_names_warns_about_contained_name_pairs(self):
        import io
        from contextlib import redirect_stdout
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)
            (p / 'config.json').write_text(
                json.dumps(dict(j.config_defaults(), names=['陈云', '陈云飞']),
                           ensure_ascii=False), encoding='utf-8')
            out = io.StringIO()
            with redirect_stdout(out):
                self.assertEqual(j.main(['--check-names'], base_dir=p), 0)
            self.assertIn('陈云 ⊂ 陈云飞', out.getvalue())

    def test_organization_terms_not_personal_names(self):
        names = j.get_names(j.config_defaults())
        title = '市人大常委会党组理论学习中心组举行集体学习，廖奇志主持并讲话'
        self.assertEqual(j.match_names(title, '省委政协人大', names, True), [])
        self.assertEqual(j.match_names('陈云调研', '', names, False), ['陈云'])

    def test_duplicate_body_different_titles(self):
        body = '陈云部署推进重大项目建设质量效率。' * 9
        items = [
            dict(source='jjxw', title='报道甲', url='a', date='2026-09-18',
                 body=[body], status='成功'),
            dict(source='jjxw', title='报道乙', url='b', date='2026-09-18',
                 body=[body], status='成功')
        ]
        retained, failures = j.deduplicate_for_month(items, '2026-09')
        self.assertEqual((len(retained), len(failures)), (1, 0))

    def test_same_title_different_body_kept(self):
        a = '陈云安排产业建设发展合作工作。' * 10
        b = '黄强主持财政会议研究财政预算计划。' * 10
        rows = [
            dict(source='jjxw', title='市委召开会议', url='a', date='2026-09-18',
                 body=[a], status='成功'),
            dict(source='jjxw', title='市委召开会议', url='b', date='2026-09-18',
                 body=[b], status='成功')
        ]
        retained, failures = j.deduplicate_for_month(rows, '2026-09')
        self.assertEqual((len(retained), len(failures)), (2, 0))

    def test_migrate_old_excel_with_backup(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder) / 'old.xlsx'
            old = Workbook()
            govt = old.active
            govt.title = '九江市政府'
            govt.append(['新闻标题', '发布日期'])
            govt.append(['旧新闻', '2026-09-18'])
            news = old.create_sheet('九江新闻网')
            news.append(['新闻标题', '发布日期', '自动匹配', '命中姓名',
                         '人工选择', '正文状态', '原文链接', '说明', '事件日期'])
            news.cell(2, 1, '陈云在九江调研')
            news.cell(2, 2, '2026-09-18')
            news.cell(2, 5, '收录')
            news.cell(2, 7,
                      'https://www.jjxw.cn/article/2026/09/18/190001.html')
            news.cell(2, 1).fill = PatternFill('solid', fgColor='FFFF00')
            old.save(p)
            j.backup_spreadsheet(p)
            converted = j.open_book(p)
            self.assertEqual(converted.sheetnames, ['九江新闻网'])
            self.assertEqual(converted.active['E2'].value, '收录')
            self.assertTrue(j.is_yellow(converted.active['A2']))
            converted.save(p)
            copies = list((Path(folder) / '备份').glob('*.xlsx'))
            self.assertEqual(len(copies), 1)
            self.assertIn('九江市政府', load_workbook(copies[0]).sheetnames)

    def test_old_government_cache_is_ignored(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)
            with j.connect_db(p / 'cache.sqlite3') as con:
                con.execute("INSERT INTO news(url,source,title,pub_date) VALUES(?,?,?,?)",
                            ('https://www.jiujiang.gov.cn/old.html',
                             'gov', '陈云旧报道', '2026-09-18'))
                book = j.open_book(p / 'news.xlsx')
                cfg = dict(j.config_defaults(), delay_seconds=0)
                with patch.object(j, 'fetch_html',
                                  side_effect=AssertionError('should not fetch')) as fetch:
                    j.fetch_article_details(j.build_session('test'), con, cfg,
                                            2026, '2026-09', {'jjxw': []},
                                            book, 'month')
                    fetch.assert_not_called()

    def test_mock_month_export_and_document(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder)
            cfg = dict(j.config_defaults(), names=['陈云'],
                       max_pages=3, delay_seconds=0, stop_after_known_pages=3)
            (p / 'config.json').write_text(
                json.dumps(cfg, ensure_ascii=False), encoding='utf-8')
            listing = {'data': {'list': [
                {'fileID': 190001, 'publishTime': '2026-09-18 10:00:00',
                 'title': '陈云调研甲'},
                {'fileID': 190002, 'publishTime': '2026-09-18 09:00:00',
                 'title': '陈云调研乙'},
                {'fileID': 190003, 'publishTime': '2026-09-17 09:00:00',
                 'title': '普通报道'}
            ]}}
            shared = '陈云部署产业发展与高质量建设项目相关工作。' * 7
            ordinary = '全市落实交通建设和公共服务改善各项任务。' * 7

            def fake_html(session, url, timeout):
                if url.endswith('190001.html'):
                    return article('陈云调研甲', shared)
                if url.endswith('190002.html'):
                    return article('陈云调研乙', shared)
                if url.endswith('190003.html'):
                    return article('普通报道', ordinary)
                raise AssertionError('Unexpected article URL: ' + url)

            class Response:
                def __init__(self, content):
                    self.content = content
                def raise_for_status(self):
                    return None
                def __enter__(self):
                    return self
                def __exit__(self, *args):
                    return False

            def fake_get(self, url, **kwargs):
                if 'getArticles?' not in url:
                    raise AssertionError('Unexpected feed URL: ' + url)
                page = listing if 'pageNumber=1' in url else {'data': {'list': []}}
                return Response(json.dumps(page, ensure_ascii=False).encode('utf-8'))

            with patch.object(j.requests.Session, 'get', fake_get), \
                 patch.object(j, 'fetch_html', fake_html):
                code = j.main(['2026-09'], base_dir=p)
            self.assertEqual(code, 0)
            wb = load_workbook(p / 'results' / 'jiujiang_news.xlsx')
            self.assertEqual(wb.sheetnames, ['九江新闻网'])
            self.assertEqual(wb.active.max_row, 4)
            self.assertEqual(sum(wb.active.cell(i, 3).value == '命中'
                                 for i in range(2, 5)), 2)
            doc = Document(p / 'results' / '市委大事记（2026年9月）.docx')
            titles = [paragraph.text for paragraph in doc.paragraphs
                      if any(node.get(j.qn('w:val')) == '2'
                             for node in paragraph._p.xpath('./w:pPr/w:outlineLvl'))]
            self.assertEqual(len(titles), 1)
            self.assertIn('陈云', titles[0])
            verified = j.main(['2026-09', '--verify'], base_dir=p)
            self.assertEqual(verified, 0)
            report = json.loads((p / 'results' / '抓取检查报告.json').read_text('utf-8'))
            self.assertEqual(report['news'], 3)
            self.assertEqual(report['matched'], 2)
            self.assertEqual(report['word_articles'], 1)
            self.assertFalse(report['errors'])


if __name__ == '__main__':
    unittest.main()
