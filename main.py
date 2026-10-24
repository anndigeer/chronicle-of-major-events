#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""九江新闻网单来源市委大事记采集器。"""
from __future__ import annotations

import argparse
import csv
import json
from contextlib import contextmanager
import logging
import os
import re
import shutil
import sqlite3
import sys
import tempfile
import time
import html
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from urllib.parse import urljoin, urlsplit, urlunsplit, unquote

import requests
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

LOG = logging.getLogger('jrjj')
SOURCES = {'jjxw': {'sheet': '九江新闻网', 'base': 'https://www.jjxw.cn/column/668/'}}
LEADERS_URL = 'https://jiangxi.jxnews.com.cn/leader.shtml'
LEADER_SECTIONS = ('省委', '人大', '政府', '政协')
AUTO_NAMES_FILE = 'names.auto.txt'
DEFAULT_NAMES = ['陈云', '邓永翔', '黄强', '陈永奇', '庄兆林', '黄喜忠']
HEADERS = ['新闻标题', '发布日期', '自动匹配', '命中姓名', '人工选择', '正文状态', '说明', '事件日期']
LEGACY_LINK_NOTE = '【原文链接待核】'
RE_DATE = re.compile(r'(20\d{2})[-年/.](1[0-2]|0?[1-9])[-月/.](3[01]|[12]\d|0?[1-9])(?:日)?(?!\d)')
RE_ISO = re.compile(r'20\d{2}-\d{2}-\d{2}')
INVALID_URL = re.compile(r'^(?:#|javascript:|mailto:|tel:)', re.I)
CONTENT_SELECTORS = [
    '#artibody > div.trs_editor_view', '#artibody', '#zoom',
    '.TRS_Editor', '.trs_editor_view', '.article-content', '.article_content',
    '.article-body', '.article_body', '.article-text', '.article_text',
    '.news-content', '.news_content', '.content-article', '.detail-content',
    '.detail_content', '.text-content', '.wzcon', '#content', '.content',
]
BLACKLIST_CLASS = re.compile(r'(related|recommend|copyright|share|comment|footer|nav|ad[-_]|pagination|source)', re.I)


@dataclass
class News:
    source: str
    title: str
    url: str
    pub_date: str


def norm_space(s: object) -> str:
    return re.sub(r'\s+', ' ', str(s or '').replace('\u200b', '')).strip()


def clean_title(s: object) -> str:
    return norm_space(s).strip(' -|_·\u3000')


def normalize_url(base: str, href: str) -> str:
    href = (href or '').strip()
    if not href or INVALID_URL.match(href):
        return ''
    # Some sites write a root-relative path without '/', e.g. column/668/article/...
    if re.match(r'^column/\d+/', href, re.I):
        parts = urlsplit(base)
        href = f'{parts.scheme}://{parts.netloc}/' + href
    url = urljoin(base, href)
    p = urlsplit(url)
    if p.scheme not in {'https', 'http'}:
        return ''
    path = p.path
    # The 668 column has published relative links that resolve to a 404
    # (/column/668/article/...), while articles live at /article/... .
    if 'jjxw.cn' in p.netloc.lower():
        path = re.sub(r'^/(?:column/\d+/)+article/', '/article/', path)
    return urlunsplit((p.scheme, p.netloc.lower(), path, p.query, ''))


def date_from_text(raw: str) -> str:
    m = RE_DATE.search(raw or '')
    if not m:
        return ''
    try:
        return date(*map(int, m.groups())).isoformat()
    except ValueError:
        return ''


NAME_TOKEN = re.compile(r'[^\W\d_]+(?:[·•-][^\W\d_]+)*', re.UNICODE)
NAME_SEPARATORS = re.compile(r'[,，、;；\t]+')


def parse_names_text(content: str) -> list[str]:
    """Read pasted Excel columns or separated names, without guessing from prose.

    Names are separated by newlines, tabs, Chinese/ASCII commas, semicolons
    or ideographic enumeration marks. A middle dot is part of a person's
    name, not a separator. Reject malformed entries instead of silently
    changing the names used for matching.
    """
    found: list[str] = []
    seen: set[str] = set()
    for line_no, line in enumerate(content.lstrip('\ufeff').splitlines(), 1):
        if line.strip().startswith(('#', '＃')):
            continue
        for token in NAME_SEPARATORS.split(line):
            name = norm_space(token)
            if not name or name in {'姓名', '人名', '关键词'}:
                continue
            if not 2 <= len(name) <= 30 or not NAME_TOKEN.fullmatch(name):
                raise ValueError(
                    f'names.txt 第{line_no}行姓名格式不正确：{name!r}；'
                    '请只复制人名列，用换行、逗号、顿号或分号分隔，不要包含职务、序号或空格')
            if name not in seen:
                seen.add(name)
                found.append(name)
    if not found:
        raise ValueError(
            'names.txt 没有有效姓名；请每行填写一个人名（可直接粘贴Excel的一列）')
    return found


def read_names_excel(path: Path) -> list[str]:
    """Read the names column in the first worksheet of a local .xlsx file.

    Use a column headed 姓名 within the first 10 rows if present; otherwise
    use column A. Do not read unrelated department/position columns. Reject
    invalid cells with their Excel coordinates rather than guessing.
    """
    try:
        book = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(
            f'无法读取姓名Excel：{path.name}；请确认是有效.xlsx文件且未损坏：{exc}'
        ) from exc
    try:
        sheet = book.worksheets[0]
        name_columns = []
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 10),
                                   max_col=min(sheet.max_column, 30)):
            for cell in row:
                if norm_space(cell.value) in ('姓名', '人名'):
                    name_columns.append((cell.row, cell.column))
        if len(name_columns) > 1:
            raise ValueError(
                'names.xlsx 前10行出现多个“姓名”表头，'
                '请保留一个姓名列，或将所需姓名复制到A列')
        header_row, column = name_columns[0] if name_columns else (0, 1)
        found: list[str] = []
        seen: set[str] = set()
        for row in sheet.iter_rows(min_row=header_row + 1, min_col=column,
                                   max_col=column):
            cell = row[0]
            if cell.value is None:
                continue
            raw = norm_space(cell.value)
            if not raw or raw.startswith(('#', '＃')) or raw in ('姓名', '人名', '关键词'):
                continue
            # One cell = one personal name, never silently split a text cell
            # that might be an unrelated description or two separate people.
            if not 2 <= len(raw) <= 30 or not NAME_TOKEN.fullmatch(raw):
                raise ValueError(
                    f'names.xlsx {cell.coordinate} 单元格姓名格式不正确：{raw!r}；'
                    '请只保留人名（不要包含序号、职务或多个姓名）')
            if raw not in seen:
                seen.add(raw)
                found.append(raw)
        if not found:
            raise ValueError(
                'names.xlsx 未找到有效姓名；请在第一个工作表填写姓名列')
        return found
    finally:
        book.close()


def extract_leader_groups(markup: str | bytes) -> dict[str, list[str]]:
    """Extract only names in the four leadership groups in the specified table.

    Never scrape the general portrait list or adjacent 人事任免 headlines.
    If the page changes shape, fail instead of silently replacing the saved
    roster with incomplete or unrelated content.
    """
    soup = BeautifulSoup(markup, 'html.parser')
    groups: dict[str, list[str]] = {section: [] for section in LEADER_SECTIONS}
    tables = [
        table for table in soup.find_all('table')
        if (str(table.get('width', '')).strip() == '98%'
            and str(table.get('border', '')).strip() == '0'
            and str(table.get('align', '')).lower() == 'center'
            and str(table.get('cellpadding', '')).strip() == '0'
            and str(table.get('cellspacing', '')).strip() == '5'
            and 'line-height:30px' in re.sub(
                r'\s+', '', str(table.get('style', '')).lower()))
    ]
    if not tables:
        raise ValueError('没有找到江西领导名单指定表格，网页结构可能已变化')

    def heading(value: str) -> str:
        value = re.sub(r'[\s\u3000:：]+', '', value or '')
        return value if value in groups else ''

    # The page can use text or image alt/title for section headings.
    for table in tables:
        active = ''
        for node in table.descendants:
            if isinstance(node, NavigableString):
                if node.parent and node.parent.name in ('a', 'script', 'style'):
                    continue
                text = str(node).strip()
                if '人事任免' in text:
                    active = ''
                    continue
                active = heading(text) or active
            elif isinstance(node, Tag) and node.name == 'img':
                for attr in ('alt', 'title'):
                    label = heading(str(node.get(attr, '')))
                    if label:
                        active = label
                        break
            elif isinstance(node, Tag) and node.name == 'a' and active:
                raw = node.get_text(' ', strip=True)
                name = re.sub(r'[\s\u3000]+', '', raw)
                if 2 <= len(name) <= 12 and NAME_TOKEN.fullmatch(name):
                    if name not in groups[active]:
                        groups[active].append(name)

    missing = [section for section in LEADER_SECTIONS if not groups[section]]
    if missing:
        raise ValueError(
            '江西领导名单页面缺少板块或人名：' + '、'.join(missing)
            + '；保留现有名单不覆盖')
    if len({name for values in groups.values() for name in values}) > 100:
        raise ValueError('江西领导名单提取异常：人名数量超过100，拒绝覆盖')
    return groups


def render_leader_names(groups: dict[str, list[str]]) -> str:
    """Deterministic grouped plaintext: unchanged page means no git diff."""
    lines = [
        '# 江西省领导自动名单（由程序生成，勿在此文件手动维护）',
        '# 来源: ' + LEADERS_URL,
        '# 四个板块人名分别保存；跨板块同名在匹配时自动去重',
    ]
    for section in LEADER_SECTIONS:
        lines.extend(['', f'# {section}', *groups[section]])
    return '\n'.join(lines) + '\n'


def update_leaders(root: Path, session: requests.Session | None = None,
                   timeout: int = 20) -> tuple[dict[str, list[str]], bool]:
    """Fetch and atomically save the auto roster, preserving the old on error."""
    session = session or build_session('Mozilla/5.0 (compatible; news-leader-sync/1.0)')
    response = session.get(LEADERS_URL, timeout=timeout)
    response.raise_for_status()
    groups = extract_leader_groups(response.content)
    output = root / AUTO_NAMES_FILE
    contents = render_leader_names(groups)
    if output.is_file() and output.read_text(encoding='utf-8-sig') == contents:
        return groups, False
    root.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
            mode='w', encoding='utf-8', newline='\n',
            prefix='.names-auto-', suffix='.tmp', dir=root,
            delete=False) as handle:
        temporary = Path(handle.name)
        try:
            handle.write(contents)
        except BaseException:
            temporary.unlink(missing_ok=True)
            raise
    try:
        os.replace(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)
    return groups, True


def resolve_names_file(path: Path, fallback: list[str]) -> tuple[list[str], str]:
    """Combine manually managed names with the latest public auto roster.

    Manual priority is Excel > TXT > JSON. Auto names are always appended
    and de-duplicated; updating them cannot erase a user's manual entries.
    Without a manual file, the generated roster supersedes sample defaults.
    """
    excel = path.with_name('names.xlsx')
    text_file = path.with_name('names.txt')
    auto_file = path.with_name(AUTO_NAMES_FILE)
    if excel.is_file():
        manual, source = read_names_excel(excel), excel.name
    elif text_file.is_file():
        manual, source = parse_names_text(
            text_file.read_text(encoding='utf-8-sig')), text_file.name
    elif auto_file.is_file():
        manual, source = [], ''
    else:
        manual, source = fallback, path.name if path.is_file() else '默认配置'
    if auto_file.is_file():
        automatic = parse_names_text(auto_file.read_text(encoding='utf-8-sig'))
        return list(dict.fromkeys([*manual, *automatic])), (
            f'{source} + {AUTO_NAMES_FILE}' if source else AUTO_NAMES_FILE)
    return manual, source


def config_defaults() -> dict:
    return {
        'names': DEFAULT_NAMES,
        'include_body_matches': True,
        'max_pages': 170,
        'late_publication_days': 10,
        'delay_seconds': 0.4,
        'timeout_seconds': 20,
        'user_agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/126.0 Safari/537.36',
    }


def load_config(path: Path, *, create_missing: bool = True) -> dict:
    if not path.exists() and create_missing:
        path.write_text(json.dumps(config_defaults(), ensure_ascii=False, indent=2), encoding='utf-8')
        LOG.info('已创建配置文件：%s', path)
    data = json.loads(path.read_text(encoding='utf-8')) if path.is_file() else {}
    if not isinstance(data, dict):
        raise ValueError('config.json 须为JSON对象')
    cfg = {**config_defaults(), **data}
    if (not isinstance(cfg['names'], list)
            or not all(isinstance(name, str) for name in cfg['names'])):
        raise ValueError('names 必须是字符串数组，例如 ["陈云", "黄强"]')
    cfg['names'] = list(dict.fromkeys(
        norm_space(x) for x in cfg['names'] if norm_space(x) and len(norm_space(x)) >= 2))
    # Preserve the manual source priority, and append safely fetched
    # leadership names from the separately tracked names.auto.txt.
    cfg['names'], cfg['names_source'] = resolve_names_file(path, cfg['names'])
    if cfg['names_source'] != path.name:
        LOG.info('已从 %s 读取 %d 个人名', cfg['names_source'],
                 len(cfg['names']))
    grace = cfg['late_publication_days']
    if type(grace) is not int or not 0 <= grace <= 31:
        raise ValueError('late_publication_days 必须为 0～31 的整数')
    for key in ('max_pages', 'timeout_seconds'):
        value = cfg[key]
        if type(value) is not int or not 1 <= value <= 1000:
            raise ValueError(f'{key} 必须为 1～1000 的整数')
    delay = cfg['delay_seconds']
    if type(delay) not in (int, float) or not 0 <= delay <= 60:
        raise ValueError('delay_seconds 必须为 0～60 的非负数字')
    if not isinstance(cfg['include_body_matches'], bool):
        raise ValueError('include_body_matches 必须为 true 或 false')
    return cfg


def build_session(user_agent: str) -> requests.Session:
    session = requests.Session()
    session.headers.update({'User-Agent': user_agent, 'Accept-Language': 'zh-CN,zh;q=0.9', 'Accept': 'text/html,application/xhtml+xml'})
    retry = Retry(total=3, connect=2, read=2, backoff_factor=0.8,
                  status_forcelist=(429, 500, 502, 503, 504), respect_retry_after_header=True)
    adapter = HTTPAdapter(max_retries=retry, pool_connections=8, pool_maxsize=8)
    session.mount('https://', adapter)
    session.mount('http://', adapter)
    return session


def fetch_html(session: requests.Session, url: str, timeout: int) -> str:
    r = session.get(url, timeout=timeout)
    r.raise_for_status()
    if not r.encoding or r.encoding.lower() in {'iso-8859-1', 'ascii'}:
        r.encoding = r.apparent_encoding or 'utf-8'
    return r.text


FEED_DATE = re.compile(r'(20\d{2})[-年/.](1[0-2]|0?[1-9])[-月/.](3[01]|[12]\d|0?[1-9])')


def day_from_text(text: object) -> str:
    m = FEED_DATE.search(str(text or ''))
    if not m:
        return ''
    try:
        return date(*map(int, m.groups())).isoformat()
    except ValueError:
        return ''


def parse_jjxw_json(payload: str | bytes | dict) -> list[tuple[str, str, str]]:
    """Return (title, URL, day) from column 668's paginated API."""
    data = json.loads(payload) if isinstance(payload, (str, bytes)) else payload
    arrays = []

    def walk(x, depth=0):
        if depth > 8:
            return
        if isinstance(x, list):
            if any(isinstance(v, dict) and ('fileID' in v or 'fileId' in v) for v in x):
                arrays.append(x)
            else:
                for v in x[:5]:
                    walk(v, depth + 1)
        elif isinstance(x, dict):
            for v in x.values():
                if isinstance(v, (dict, list)):
                    walk(v, depth + 1)

    walk(data)
    found = []
    seen = set()
    for arr in arrays:
        for obj in arr:
            if not isinstance(obj, dict):
                continue
            ident = str(obj.get('fileID') or obj.get('fileId') or '').strip()
            day = day_from_text(obj.get('publishTime') or obj.get('publishedTime'))
            title = html.unescape(str(obj.get('title') or obj.get('articleTitle') or
                                      obj.get('fileTitle') or obj.get('name') or '')).strip()
            if not re.fullmatch(r'\d{1,15}', ident) or not day or len(title) < 2:
                continue
            yy, mm, dd = day.split('-')
            url = f'https://www.jjxw.cn/article/{yy}/{mm}/{dd}/{ident}.html'
            if url not in seen:
                found.append((title, url, day))
                seen.add(url)
    return found


def is_confirmed_empty_feed(payload: bytes | str | dict) -> bool:
    """Only an explicitly empty API article array proves the end of pagination.

    A 200 response with a changed JSON schema must be reported as an error,
    not confused with an empty page; otherwise the exported month may be partial.
    """
    data = json.loads(payload) if isinstance(payload, (bytes, str)) else payload
    if isinstance(data, list):
        return not data
    if not isinstance(data, dict):
        return False
    for container in (data, data.get('data'), data.get('result')):
        if isinstance(container, list) and not container:
            return True
        if isinstance(container, dict):
            for key in ('list', 'records', 'rows', 'items'):
                if key in container and isinstance(container[key], list):
                    if not container[key]:
                        return True
    return False


def fetch_list(session: requests.Session, page: int, cfg: dict,
               debug_dir: Path) -> tuple[list[News], bool]:
    """Only 九江新闻网 JSON listing, 100 articles per page."""
    url = (f'https://api.jjxw.cn/api/getArticles?cid=668&pageSize=100'
           f'&pageNumber={page + 1}&containSubCols=1')
    try:
        with session.get(url, timeout=cfg['timeout_seconds'],
                         headers={'Accept': 'application/json,text/plain,*/*',
                                  'Referer': 'https://www.jjxw.cn/'}) as response:
            response.raise_for_status()
            content = response.content
            rows = parse_jjxw_json(content)
        if not rows:
            if page and is_confirmed_empty_feed(content):
                return [], True
            raise ValueError('新闻网接口未返回可识别的新闻；无法确认是否已翻页结束')
        LOG.info('新闻网 API 第%d页：%d条', page + 1, len(rows))
        return [News('jjxw', title, link, published)
                for title, link, published in rows], True
    except Exception as exc:
        debug_dir.mkdir(parents=True, exist_ok=True)
        (debug_dir / f'jjxw_feed_page_{page+1}.txt').write_text(
            f'URL: {url}\nERROR: {type(exc).__name__}: {exc}\n', encoding='utf-8')
        LOG.error('九江新闻网接口失败：%s — %s', url, exc)
        return [], False


def get_names(cfg: dict) -> list[str]:
    """Use the stable union of manually entered names and published leaders."""
    return list(dict.fromkeys(cfg['names']))


def clean_body(soup: BeautifulSoup) -> tuple[str, list[str]]:
    """Choose the best article container and omit nav/recommendation/ads before matching."""
    for tag in soup.select('script, style, noscript, iframe, form, svg, nav, footer, aside'):
        tag.decompose()
    for tag in list(soup.find_all(True)):
        if tag.attrs is None:  # A parent may have been decomposed already.
            continue
        attrs = f"{tag.get('id', '')} {' '.join(tag.get('class', []))}"
        if BLACKLIST_CLASS.search(attrs) and tag.name in {'div', 'section', 'ul'}:
            tag.decompose()
    candidates: list[tuple[float, list[str], str]] = []
    for priority, selector in enumerate(CONTENT_SELECTORS):
        for node in soup.select(selector)[:5]:
            paragraphs = [norm_space(p.get_text(' ', strip=True)) for p in node.select('p')]
            paragraphs = [p for p in paragraphs if len(p) >= 8]
            if not paragraphs:
                paragraphs = [norm_space(t) for t in node.get_text('\n', strip=True).split('\n')]
                paragraphs = [p for p in paragraphs if len(p) >= 12]
            text = '\n'.join(paragraphs)
            if len(text) < 65:
                continue
            all_a = sum(len(a.get_text(' ', strip=True)) for a in node.select('a'))
            ratio = all_a / max(len(text), 1)
            if ratio > .65:
                continue
            # More real text is useful; prefer article-specific containers for close ties.
            score = min(len(text), 20000) * (1 - ratio) + max(0, 20 - priority) * 18
            candidates.append((score, paragraphs, selector))
    if not candidates:
        # An unknown redesign should not force code changes solely to add a CSS selector.
        # Find the densest prose block with a few substantial paragraphs.
        for node in soup.select('article, main, section, div'):
            paragraphs = [norm_space(x.get_text(' ', strip=True)) for x in node.select('p')]
            paragraphs = [x for x in paragraphs if len(x) >= 15]
            length = sum(map(len, paragraphs))
            if length < 65 or len(paragraphs) < 2:
                continue
            link_chars = sum(len(a.get_text(' ', strip=True)) for a in node.select('a'))
            if link_chars / length > .22:
                continue
            # Penalize giant wrappers and prefer coherent, well-contained paragraphs.
            score = min(length, 12000) - link_chars * 4 - max(0, len(paragraphs) - 35) * 30
            candidates.append((score, paragraphs, '[自动发现的正文容器]'))
    if not candidates:
        return '', []
    score, paragraphs, selector = max(candidates, key=lambda x: x[0])
    cleaned = []
    for p in paragraphs:
        if p in cleaned[-1:]:
            continue
        if re.match(r'^(?:责任编辑|编辑：|来源：|分享到|打印本页|关闭窗口|上一篇|下一篇)', p) and len(p) < 90:
            continue
        cleaned.append(p)
    return selector, cleaned


def parse_article(html: str, fallback_title: str) -> tuple[str, list[str], str]:
    soup = BeautifulSoup(html, 'html.parser')
    title = ''
    for selector in ['#sybt .mainTitle', '#sybt', 'h1.article-title', '.article-title h1', '.article_title h1', 'h1.title', 'article h1', 'h1']:
        node = soup.select_one(selector)
        if node:
            candidate = clean_title(node.get_text(' ', strip=True))
            if len(candidate) >= 4 and len(candidate) <= 250:
                title = candidate
                break
    if not title:
        meta = soup.select_one('meta[property="og:title"], meta[name="title"]')
        title = clean_title(meta.get('content', '')) if meta else ''
    selector, paragraphs = clean_body(soup)
    return title or fallback_title, paragraphs, selector


@contextmanager
def connect_db(path: Path):
    """Manage a SQLite transaction AND close its file handle on all platforms.

    sqlite3.Connection.__exit__ commits/rolls back, but does not close the
    connection. Without an explicit close, Windows retains a lock on the
    database file and temporary-directory cleanup fails with WinError 32.
    """
    con = sqlite3.connect(path)
    try:
        con.row_factory = sqlite3.Row
        con.executescript('''
            CREATE TABLE IF NOT EXISTS news (
                url TEXT PRIMARY KEY,
                source TEXT NOT NULL,
                title TEXT NOT NULL,
                pub_date TEXT NOT NULL,
                article_title TEXT DEFAULT '',
                body_json TEXT DEFAULT '',
                status TEXT DEFAULT '未抓取',
                message TEXT DEFAULT '',
                fetched_at TEXT DEFAULT '',
                created_at TEXT DEFAULT CURRENT_TIMESTAMP
            );
            CREATE INDEX IF NOT EXISTS idx_news_month ON news(pub_date, source);
            CREATE TABLE IF NOT EXISTS metadata (key TEXT PRIMARY KEY, value TEXT NOT NULL);
        ''')
        try:
            yield con
            con.commit()
        except BaseException:
            con.rollback()
            raise
    finally:
        con.close()


def upsert_news(con: sqlite3.Connection, record: News) -> bool:
    found = con.execute('SELECT url FROM news WHERE url=?', (record.url,)).fetchone()
    if found:
        con.execute('UPDATE news SET title=?,pub_date=?,source=? WHERE url=?',
                    (record.title, record.pub_date, record.source, record.url))
        return False
    con.execute('INSERT INTO news(url,source,title,pub_date) VALUES(?,?,?,?)',
                (record.url, record.source, record.title, record.pub_date))
    return True


def get_cell_url(cell, explicit='') -> str:
    if explicit and str(explicit).startswith(('https://', 'http://')):
        return str(explicit)
    if cell.hyperlink and cell.hyperlink.target:
        return cell.hyperlink.target
    if cell.data_type == 'f' and isinstance(cell.value, str):
        m = re.search(r'^=HYPERLINK\(\s*"((?:[^"]|"")*)"', cell.value, re.I)
        return m.group(1).replace('""', '"') if m else ''
    if isinstance(cell.value, str) and cell.value.startswith('http'):
        return cell.value
    return ''


def get_cell_title(cell) -> str:
    v = cell.value
    if isinstance(v, str) and v.upper().startswith('=HYPERLINK('):
        m = re.search(r'^=HYPERLINK\(\s*"(?:[^"]|"")*"\s*[,;]\s*"((?:[^"]|"")*)"', v, re.I)
        return m.group(1).replace('""', '"') if m else ''
    return clean_title(v)


def is_yellow(cell) -> bool:
    if cell.fill is None or cell.fill.patternType != 'solid':
        return False
    color = cell.fill.fgColor
    if color.type == 'rgb' and color.rgb:
        return color.rgb.upper()[-6:] in {'FFFF00', 'FFEB3B', 'FFF200', 'FFEE00', 'FFF2CC'}
    return False


def selection(row) -> str:
    value = norm_space(row[4].value if len(row) > 4 else '').lower()
    if value in {'排除', '否', 'no', 'n', '0', '忽略'}:
        return 'exclude'
    if value in {'收录', '是', 'yes', 'y', '1', '人工收录'}:
        return 'include'
    if is_yellow(row[0]) or (len(row) > 1 and is_yellow(row[1])):
        return 'include'
    return 'auto'


def canonical_news_url(base: str, href: object) -> str:
    """Normalize article URLs and reject other domains and non-news links."""
    raw = str(href or '').strip()
    if not raw:
        return ''
    url = normalize_url(base, raw)
    parts = urlsplit(url)
    if parts.hostname not in {'www.jjxw.cn', 'jjxw.cn'}:
        return ''
    if not parts.path.lower().endswith(('.html', '.htm')):
        return ''
    return url


def legacy_link_candidate(note: object) -> str:
    """Recover a legacy G-column conflict saved in the new G note cell."""
    match = re.search(re.escape(LEGACY_LINK_NOTE) + r'(https?://[^\s]+)',
                      str(note or ''))
    return match.group(1) if match else ''


def migrate_legacy_columns(ws, base: str) -> None:
    """Convert old A:I to A:H exactly once, keeping links and manual data.

    The former G URL is redundant. If it disagrees with the actual A
    hyperlink, retain it inside a review marker in the former H note;
    after the physical column deletion, that note becomes G.
    """
    if norm_space(ws.cell(1, 7).value) != '原文链接':
        return
    for idx in range(2, ws.max_row + 1):
        a = ws.cell(idx, 1)
        old_g = ws.cell(idx, 7)
        old_note = ws.cell(idx, 8)
        a_url = canonical_news_url(base, get_cell_url(a))
        old_url = canonical_news_url(base, get_cell_url(old_g) or old_g.value)
        title = get_cell_title(a)
        if not a_url and old_url:
            if a.data_type == 'f' and title:
                a.value = title
            a.hyperlink = old_url
        elif a_url and old_url and a_url != old_url:
            # Do not discard mismatched source links when removing column G.
            marker = LEGACY_LINK_NOTE + old_url
            old_note.value = (str(old_note.value or '').rstrip() + '\n' + marker).strip()
    ws.delete_cols(7, 1)
    LOG.info('已将旧版Excel从A:I迁移为A:H；原G列原文链接已取消')


def reconcile_row_link(row, base: str, con: sqlite3.Connection | None = None,
                       source: str = 'jjxw',
                       conflicts: list[dict] | None = None) -> str:
    """A-column hyperlink is the only authoritative URL after migration."""
    a = row[0]
    url = canonical_news_url(base, get_cell_url(a))
    title = get_cell_title(a)
    if a.data_type == 'f' and title:
        a.value = title
        if url:
            a.hyperlink = url
    if a.hyperlink and a.hyperlink.target and url and a.hyperlink.target != url:
        a.hyperlink = url
    note = row[6] if len(row) > 6 else None
    old_url = canonical_news_url(base, legacy_link_candidate(note.value)) if note else ''
    if old_url:
        if url == old_url:
            # Editing A to the saved original G candidate resolves the
            # discrepancy; remove just the review marker, not other notes.
            note.value = re.sub(re.escape(LEGACY_LINK_NOTE) + r'https?://[^\s]+',
                                '', str(note.value)).strip() or None
        else:
            if conflicts is not None:
                conflicts.append({
                    'row': a.row, 'title': title, 'a_url': url,
                    'g_url': old_url, 'reason': '旧版A/G原文链接不一致；请核实后清除G列待核标记'
                })
            LOG.error('第%s行旧版链接存在冲突，A=%s；历史G=%s', a.row, url, old_url)
            return ''
    return url


def open_book(path: Path):
    """One worksheet, A:H; migrate historical G links before rewriting headers."""
    if path.exists():
        book = load_workbook(path)
    else:
        book = Workbook()
        book.active.title = '九江新闻网'
    name = SOURCES['jjxw']['sheet']
    if name not in book.sheetnames:
        book.create_sheet(name)
    for old in list(book.sheetnames):
        if old != name:
            LOG.warning('旧工作表“%s”不再处理，已保留在运行前Excel备份中', old)
            del book[old]
    ws = book[name]
    migrate_legacy_columns(ws, SOURCES['jjxw']['base'])
    for i, header in enumerate(HEADERS, 1):
        cell = ws.cell(1, i)
        cell.value = header
        cell.fill = PatternFill('solid', fgColor='203A58')
        cell.font = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center')
    for col, width in {'A': 73, 'B': 16, 'C': 14, 'D': 25,
                       'E': 14, 'F': 16, 'G': 35, 'H': 18}.items():
        ws.column_dimensions[col].width = width
    # Remove the former I-column width left behind by openpyxl after a shift.
    if 'I' in ws.column_dimensions:
        del ws.column_dimensions['I']
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:H{max(ws.max_row, 2)}'
    ws.row_dimensions[1].height = 27
    if not ws.data_validations.dataValidation:
        validation = DataValidation(type='list', formula1='"收录,排除"', allow_blank=True)
        validation.error = '只能填写“收录”或“排除”'
        validation.errorTitle = '人工选择'
        validation.showErrorMessage = True
        ws.add_data_validation(validation)
        validation.add('E2:E1048576')
    return book


def existing_rows(book, con: sqlite3.Connection,
                  conflicts: list[dict] | None = None) -> dict[str, set[str]]:
    urls = {source: set() for source in SOURCES}
    for source, info in SOURCES.items():
        ws = book[info['sheet']]
        for idx in range(2, ws.max_row + 1):
            row = [ws.cell(idx, j) for j in range(1, 9)]
            url = reconcile_row_link(row, info['base'], con, source, conflicts)
            if not url:
                # Hold both historical candidates to avoid duplicate
                # insertion until a legacy mismatch is reviewed.
                candidate_a = canonical_news_url(info['base'], get_cell_url(row[0]))
                candidate_old = canonical_news_url(
                    info['base'], legacy_link_candidate(row[6].value))
                urls[source].update(filter(None, (candidate_a, candidate_old)))
                continue
            if url in urls[source]:
                continue
            urls[source].add(url)
            d = date_from_text(str(row[1].value))
            if isinstance(row[1].value, (datetime, date)):
                d = row[1].value.strftime('%Y-%m-%d')
            title = get_cell_title(row[0])
            if row[0].data_type == 'f' and title:
                # Native hyperlinks survive editing and do not depend on WPS/Excel recalculation.
                # Changing only the value leaves the manual yellow fill intact.
                row[0].value = title
                row[0].hyperlink = url
            if d and title:
                # Excel is a manual-review surface and may be stale.
                # Use it to populate missing cache records only; the website
                # listing remains authoritative for known title/date metadata.
                con.execute(
                    "INSERT OR IGNORE INTO news(url,source,title,pub_date) "
                    "VALUES(?,?,?,?)", (url, source, title, d))
    con.commit()
    return urls


def save_book_atomic(book, path: Path) -> None:
    """Save beside the target, then atomically replace it when complete.

    An interrupted or failed Excel write leaves the original intact.
    Works on Windows and Linux without adding external dependencies.
    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
            prefix=f'.{path.stem}.', suffix=path.suffix,
            dir=path.parent, delete=False) as handle:
        temporary = Path(handle.name)
    try:
        book.save(temporary)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def write_new_rows(book, records: dict[str, list[News]]):
    """Append newly published articles to one lifelong worksheet.

    Insertion into row 2 used to shift existing hyperlinks/comments/styles.
    Appending keeps all existing row positions stable for manual review.
    Column B's date filter can be used to sort/view individual months.
    """
    for source, news in records.items():
        if not news:
            continue
        ws = book[SOURCES[source]['sheet']]
        base = SOURCES[source]['base']
        known = set()
        for idx in range(2, ws.max_row + 1):
            a_link = canonical_news_url(base, get_cell_url(ws.cell(idx, 1)))
            old_link = canonical_news_url(
                base, legacy_link_candidate(ws.cell(idx, 7).value))
            known.update(filter(None, (a_link, old_link)))
        added = 0
        for item in sorted(news, key=lambda n: (n.pub_date, n.title)):
            url = canonical_news_url(base, item.url)
            if not url or url in known:
                continue
            known.add(url)
            added += 1
            idx = ws.max_row + 1
            a = ws.cell(idx, 1, item.title)
            a.hyperlink = url
            a.font = Font(name='楷体_GB2312', size=12, color='1565C0', underline='single')
            a.alignment = Alignment(vertical='center', wrap_text=True)
            b = ws.cell(idx, 2, item.pub_date)
            b.font = Font(name='微软雅黑', size=11)
            b.alignment = Alignment(horizontal='center', vertical='center')
            ws.row_dimensions[idx].height = 28
        ws.auto_filter.ref = f'A1:H{ws.max_row}'
        LOG.info('%s：按发布日期增量累计新增 %d 条（保留所有原有行）',
                 SOURCES[source]['sheet'], added)



def publication_window(month: str, grace_days: int = 10) -> tuple[str, str]:
    """Inclusive listing window, e.g. 2026-09-01 through 2026-10-10."""
    y, m = map(int, month.split('-'))
    first = date(y, m, 1)
    next_month = (first.replace(day=28) + timedelta(days=4)).replace(day=1)
    last = next_month + timedelta(days=grace_days - 1)
    return first.isoformat(), last.isoformat()


EVENT_ACTION = re.compile(
    r'召开|举行|主持|出席|会见|调研|视察|走访|督导|考察|部署|'
    r'参加|启动|开幕|闭幕|签约|举行座谈|听取汇报|检查|慰问|'
    r'宣讲|在.{0,12}讲话|开展'
)
EVENT_DATE = re.compile(
    r'(?:(?P<year>20\d{2})年)?(?P<month>1[0-2]|0?[1-9])月'
    r'(?P<day>3[01]|[12]\d|0?[1-9])日|'
    r'(?P<relative_month>本月|上月)(?P<relative_day>3[01]|[12]\d|0?[1-9])日|'
    r'(?P<relative>昨日|昨天|前日|前天)'
)


def infer_event_date(title: str, paragraphs: list[str], publication: str,
                     max_lag: int = 10) -> str:
    """Conservatively infer event dates without mistaking a report date for one."""
    published = date.fromisoformat(publication)
    for text in [title, *paragraphs[:2]]:
        lead = norm_space(text)[:280]
        for match in EVENT_DATE.finditer(lead):
            before = lead[max(0, match.start() - 8):match.start()]
            if re.search(r'将于|计划于|预计于|拟于', before):
                continue
            # A statistics cutoff is not an event's date, even when a
            # subsequent clause describes meetings or inspections.
            if re.search(r'(?:截至|截止(?:到)?|统计至)\s*$', before):
                continue
            context = lead[match.end():match.end() + 70].split('。', 1)[0]
            # Dates attached to a journalist receiving or publishing a
            # report are not actual event dates and never belong in H.
            reporter_clause = re.search(
                r'记者.{0,28}(?:获悉|获知|了解到|得知)|'
                r'(?:记者|通讯员).{0,14}(?:报道|发稿)|(?:消息|通报)称',
                context[:45])
            if reporter_clause or (
                re.search(r'记者\s*
            # not October 2. Each candidate can only use words occurring
            # before the *next* full event-date expression.
            next_date = EVENT_DATE.search(context)
            if next_date:
                context = context[:next_date.start()]
            action = EVENT_ACTION.search(context)
            if not action:
                continue
            if re.search(r'预计|计划|拟于|将于|即将|将|拟',
                         context[:action.end()]):
                continue
            if match.group('relative'):
                lag = 2 if match.group('relative') in {'前日', '前天'} else 1
                happened = published - timedelta(days=lag)
            else:
                year = int(match.group('year') or published.year)
                if match.group('relative_month'):
                    month_start = published.replace(day=1)
                    if match.group('relative_month') == '上月':
                        month_start -= timedelta(days=1)
                    year, mm = month_start.year, month_start.month
                    dd = int(match.group('relative_day'))
                else:
                    mm, dd = int(match.group('month')), int(match.group('day'))
                    if not match.group('year') and published.month == 1 and mm == 12:
                        year -= 1
                try:
                    happened = date(year, mm, dd)
                except ValueError:
                    continue
            lag = (published - happened).days
            if 0 <= lag <= max_lag:
                return happened.isoformat()
    return ''


def crawl(session: requests.Session, con: sqlite3.Connection,
          known: dict[str, set[str]], cfg: dict, year: int, debug: Path,
          target_month: str | None = None, *, status: dict | None = None
          ) -> dict[str, list[News]]:
    """Scan through the full date window, even when intermediate pages are cached.

    Record completion status instead of silently accepting partial pagination.
    Every successfully fetched page is committed to SQLite; the caller can
    persist new Excel rows before reporting a failure.
    """
    source = 'jjxw'
    new = {source: []}
    window_start, window_end = publication_window(
        target_month or f'{year}-01', cfg['late_publication_days'])
    complete = False
    error = ''
    last_page_signature: tuple[str, ...] | None = None
    for page in range(cfg['max_pages']):
        news, ok = fetch_list(session, page, cfg, debug)
        if not ok:
            error = f'新闻网第{page + 1}页请求失败，采集不完整'
            break
        if not news:
            complete = True  # API reached the end of the listing
            break
        page_signature = tuple(n.url for n in news)
        if page_signature == last_page_signature:
            error = (f'新闻网第{page + 1}页与上一页内容完全相同，'
                     '分页接口可能未正确翻页，采集不完整')
            break
        last_page_signature = page_signature
        current = [n for n in news if window_start <= n.pub_date <= window_end]
        count = 0
        for n in current:
            if n.url not in known[source]:
                known[source].add(n.url)
                new[source].append(n)
                count += 1
            upsert_news(con, n)
        con.commit()
        LOG.info('第%d页：解析%d条，时间范围内%d条，新%d条',
                 page + 1, len(news), len(current), count)
        if all(n.pub_date < window_start for n in news):
            complete = True  # All records on this page predate the window.
            break
        # Never stop simply because pages are already cached: that can skip
        # unseen articles further down in the target month.
        if cfg['delay_seconds']:
            time.sleep(cfg['delay_seconds'])
    else:
        error = (f'已达到最大分页数 max_pages={cfg["max_pages"]}，'
                 '尚未扫描到目标月份之前的记录，结果可能不完整')
    if status is not None:
        status['complete'] = complete
        status['error'] = error
    if error:
        LOG.error('%s；已获取的新闻仍可保留在Excel和SQLite中', error)
    return new


def fetch_article_details(session: requests.Session, con: sqlite3.Connection,
                          cfg: dict, year: int, month: str,
                          new: dict[str, list[News]], book, scope: str):
    source = 'jjxw'
    fresh_urls = {x.url for x in new[source]}
    manual_urls = set()
    for row in book[SOURCES[source]['sheet']].iter_rows(min_row=2, max_col=8):
        url = reconcile_row_link(row, SOURCES[source]['base'], con, source)
        if selection(row) == 'include' and url:
            manual_urls.add(url)
    # Query by full date range so December news published in January is included.
    start, end = publication_window(month, cfg['late_publication_days'])
    rows = con.execute(
        "SELECT * FROM news WHERE source=? AND pub_date>=? AND pub_date<=? "
        "ORDER BY pub_date DESC", (source, start, end)).fetchall()
    to_fetch = [
        n for n in rows if not (n['status'] == '成功' and n['body_json'])
        and ((scope == 'all' and n['url'] in fresh_urls)
             or start <= n['pub_date'] <= end
             or n['url'] in manual_urls)
    ]
    LOG.info('需要抓取正文%d篇；成功缓存直接复用', len(to_fetch))
    for idx, n in enumerate(to_fetch, 1):
        try:
            html = fetch_html(session, n['url'], cfg['timeout_seconds'])
            title, paragraphs, selector = parse_article(html, n['title'])
            if sum(map(len, paragraphs)) < 65:
                raise ValueError('正文少于65字，可能是访问限制或页面改版')
            con.execute(
                "UPDATE news SET article_title=?,body_json=?,status='成功',message=?,fetched_at=? WHERE url=?",
                (title, json.dumps(paragraphs, ensure_ascii=False), selector,
                 datetime.now().isoformat(timespec='seconds'), n['url']))
            LOG.info('[%d/%d] 正文成功：%s', idx, len(to_fetch), n['title'][:35])
        except Exception as exc:
            con.execute("UPDATE news SET status='失败',message=?,fetched_at=? WHERE url=?",
                        (str(exc)[:300], datetime.now().isoformat(timespec='seconds'), n['url']))
            LOG.warning('[%d/%d] 正文失败：%s — %s',
                        idx, len(to_fetch), n['title'][:35], exc)
        con.commit()
        if cfg['delay_seconds']:
            time.sleep(cfg['delay_seconds'])


def match_names(title: str, body: str, names: list[str], include_body: bool) -> list[str]:
    """Match configured names, avoiding substring-only hits inside longer listed names.

    If both 陈云 and 陈云飞 are configured, an occurrence of 陈云 inside
    陈云飞 does not count as a separate 陈云 mention. A separate occurrence
    of 陈云 elsewhere in the same article still counts.
    """
    target = title + ('\n' + body if include_body else '')
    unique_names = list(dict.fromkeys(name for name in names if name))
    occurrences: dict[str, list[tuple[int, int]]] = {}
    for name in unique_names:
        spans = []
        start = 0
        while True:
            start = target.find(name, start)
            if start < 0:
                break
            spans.append((start, start + len(name)))
            start += 1
        occurrences[name] = spans

    matched = []
    for name in unique_names:
        for start, end in occurrences[name]:
            covered_by_longer_name = any(
                len(other) > len(name)
                and other_start <= start and end <= other_end
                for other, spans in occurrences.items()
                if len(other) > len(name)
                for other_start, other_end in spans
            )
            if not covered_by_longer_name:
                matched.append(name)
                break
    return matched


def refresh_book(book, con: sqlite3.Connection, names: list[str], cfg: dict) -> list[dict]:
    chosen = []
    for source in SOURCES:
        ws = book[SOURCES[source]['sheet']]
        for i in range(2, ws.max_row + 1):
            row = [ws.cell(i, k) for k in range(1, 9)]
            url = reconcile_row_link(row, SOURCES[source]['base'], con, source)
            if not url:
                continue
            n = con.execute('SELECT * FROM news WHERE url=? AND source=?', (url, source)).fetchone()
            if n is None:
                continue
            paras = json.loads(n['body_json'] or '[]')
            matches = match_names(n['article_title'] or n['title'], '\n'.join(paras), names, cfg['include_body_matches'])
            choose = selection(row)
            # Column H: preserve human edits, recalculate machine guesses.
            # The comment stores our last guess. If the value differs, it
            # was edited manually and should never be overwritten.
            event_cell = row[7]
            raw_event = event_cell.value
            event_day = (raw_event.strftime('%Y-%m-%d')
                         if isinstance(raw_event, (date, datetime))
                         else date_from_text(str(raw_event or '')))
            auto_marker = 'AUTO_EVENT_DATE:'
            prev_guess = ''
            if event_cell.comment:
                text = event_cell.comment.text
                if text.startswith(auto_marker):
                    prev_guess = text[len(auto_marker):].splitlines()[0].strip()
            if prev_guess:
                if event_day == prev_guess:
                    # Still machine-owned: refresh the guess from latest body.
                    event_day = ''
                    event_cell.value = None
                    event_cell.comment = None
                else:
                    # User has changed or cleared our suggested date.
                    event_cell.comment = None
                    if not event_day:
                        # Clearing a suggestion explicitly asks for review,
                        # not automatic regeneration on the same refresh.
                        event_cell.comment = Comment(
                            'MANUAL_EVENT_DATE_REVIEW', '人工日期标记')
            needs_manual_review = bool(
                event_cell.comment and
                event_cell.comment.text.startswith('MANUAL_EVENT_DATE_REVIEW'))
            if needs_manual_review and event_day:
                # A newly typed valid event date completes the manual review.
                event_cell.comment = None
                needs_manual_review = False
            invalid_event = bool(norm_space(event_cell.value)) and not event_day
            if not event_day and not needs_manual_review and not invalid_event and n['status'] == '成功':
                inferred = infer_event_date(
                    n['article_title'] or n['title'], paras, n['pub_date'],
                    cfg['late_publication_days'])
                if inferred:
                    event_day = inferred
                    event_cell.value = inferred
                    event_cell.comment = Comment(
                        auto_marker + inferred + '\n自动推断，请核实或直接修改单元格。',
                        '采集程序')
            effective_day = event_day or n['pub_date']
            selected = not invalid_event and (choose == 'include' or (choose == 'auto' and bool(matches)))
            row[2].value = '命中' if matches else ('待正文' if n['status'] != '成功' else '未命中')
            row[3].value = '、'.join(matches)
            row[5].value = n['status']
            # B column is the site's publication date; manually reviewed
            # event dates belong in column H and must remain untouched.
            if row[1].value != n['pub_date']:
                row[1].value = n['pub_date']
            if n['status'] != '成功':
                status_note = n['message']
            elif event_day:
                status_note = '事件日期已填写（H列可复核）'
            else:
                status_note = '事件日期未确认；跨月新闻请核对后在H列填写'
            if invalid_event:
                status_note = 'H列事件日期无效：请填写真实日期 YYYY-MM-DD（已暂停收录）'
            elif choose == 'exclude':
                status_note = '人工排除；' + status_note
            prior_note = norm_space(row[6].value)
            generated_note = prior_note.startswith((
                '事件日期已填写', '事件日期未确认', 'H列事件日期无效',
                'I列事件日期无效', '人工排除；'))
            # Historical H-column personal remarks have moved to G. Avoid
            # destroying them on every refresh; column F already shows
            # the body status, and H holds the reviewable event date.
            if not prior_note or generated_note or prior_note == n['message']:
                row[6].value = status_note
            elif invalid_event and status_note not in prior_note:
                row[6].value = prior_note + '\n' + status_note
            row[2].alignment = Alignment(horizontal='center')
            row[4].alignment = Alignment(horizontal='center')
            row[5].alignment = Alignment(horizontal='center')
            if selected:
                chosen.append({'source': source, 'url': url, 'date': effective_day,
                               'published': n['pub_date'],
                               'title': n['article_title'] or n['title'], 'body': paras,
                               'status': n['status'], 'matches': matches})
    return chosen


def canonical_title(title: str) -> str:
    return re.sub(r'[\W_]+', '', (title or '').lower())


def canonical_body(body: list[str]) -> str:
    """Ignore layout whitespace and punctuation when comparing full articles."""
    return re.sub(r'[\W_]+', '', '\n'.join(body).lower())


def _similar_article_body(a: str, b: str) -> bool:
    """Require substantial content overlap to deduplicate equal headlines.

    This prevents two different meetings with the same generic title from
    being collapsed merely because they were published close together.
    """
    if not a or not b:
        return False
    if a == b and len(a) >= 65:
        return True
    if min(len(a), len(b)) < 100:
        return False
    from collections import Counter
    a_pairs = Counter(a[i:i + 2] for i in range(len(a) - 1))
    b_pairs = Counter(b[i:i + 2] for i in range(len(b) - 1))
    if not a_pairs or not b_pairs:
        return False
    # Multiset counts keep long repeated phrases weighted correctly.
    shared = sum((a_pairs & b_pairs).values())
    total = sum((a_pairs | b_pairs).values())
    return shared / total >= 0.90


def _same_news_story(a: dict, b: dict) -> bool:
    a_body, b_body = canonical_body(a['body']), canonical_body(b['body'])
    # Reposting identical substantial body text under a different headline is
    # still one story, even if publication dates or titles differ.
    if a_body == b_body and len(a_body) >= 65:
        return True
    if canonical_title(a['title']) != canonical_title(b['title']):
        return False
    if abs((date.fromisoformat(a['date']) - date.fromisoformat(b['date'])).days) > 2:
        return False
    return _similar_article_body(a_body, b_body)


def deduplicate_for_month(selected: list[dict], month: str) -> tuple[list[dict], list[dict]]:
    """Return distinct articles for Word, leaving all source rows in Excel."""
    articles: list[dict] = []
    failed: list[dict] = []
    duplicates = 0
    eligible = sorted((x for x in selected if x['date'].startswith(month)),
                      key=lambda x: (x['date'], x['title']))
    for item in eligible:
        if item['status'] != '成功' or not item['body']:
            failed.append(item)
            continue
        match_idx = next((i for i, seen in enumerate(articles)
                          if _same_news_story(seen, item)), None)
        if match_idx is None:
            articles.append(item)
            continue
        duplicates += 1
        seen = articles[match_idx]
        # Keep the longer article among same-source duplicates.
        if len(canonical_body(item['body'])) > len(canonical_body(seen['body'])):
            articles[match_idx] = item
    articles.sort(key=lambda x: (x['date'], x['title']))
    LOG.info('Word 正文去重：候选正文 %d 篇，合并重复 %d 篇，最终保留 %d 篇',
             len(eligible) - len(failed), duplicates, len(articles))
    return articles, failed


def set_page_layout(doc):
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin, section.bottom_margin = Cm(2.54), Cm(2.54)
        section.left_margin, section.right_margin = Cm(2.7), Cm(2.7)
    normal = doc.styles['Normal']
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(28)
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)


def add_run(p, text: str, font: str):
    r = p.add_run(text)
    r.font.name = font
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), font)
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return r


def backup_generated_word(path: Path) -> None:
    """At most one generated Word backup per day; retain last seven.

    Matches only timestamped backups created by the program and never
    deletes a manually named document.
    """
    directory = path.parent / '备份'
    directory.mkdir(parents=True, exist_ok=True)
    prefix = f'{path.stem}_'
    generated = re.compile(
        re.escape(prefix) + r'\d{8}_\d{6}(?:_\d{6})?\.docx')
    previous = sorted(
        (item for item in directory.glob(f'{prefix}*.docx')
         if item.is_file() and generated.fullmatch(item.name)),
        key=lambda x: x.name)
    today = prefix + datetime.now().strftime('%Y%m%d') + '_'
    if not any(item.name.startswith(today) for item in previous):
        stamp = datetime.now().strftime('%Y%m%d_%H%M%S_%f')
        snapshot = directory / f'{prefix}{stamp}.docx'
        shutil.copy2(path, snapshot)
        previous.append(snapshot)
    for stale in sorted(previous, key=lambda x: x.name)[:-7]:
        stale.unlink()


def build_word(month: str, articles: list[dict], out_dir: Path) -> Path | None:
    if not articles:
        # Re-exporting a month after all its candidates have been excluded
        # must not leave last week's generated Word looking current.
        # Preserve any user-created Word (no generator signature).
        y, m = map(int, month.split('-'))
        for suffix in ('', '_自动整理'):
            candidate = out_dir / f'市委大事记（{y}年{m}月）{suffix}.docx'
            if not candidate.exists():
                continue
            old_document = Document(candidate)
            if 'jrjj-v2-generated' not in (
                    old_document.core_properties.keywords or ''):
                LOG.info('保留人工创建的Word文档：%s', candidate)
                continue
            backup_generated_word(candidate)
            candidate.unlink()
            LOG.warning('全部候选已排除或不可导出，备份并清理过期自动Word：%s', candidate)
        return None
    y, m = map(int, month.split('-'))
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f'市委大事记（{y}年{m}月）.docx'
    if path.exists():
        existing = Document(path)
        if 'jrjj-v2-generated' not in (existing.core_properties.keywords or ''):
            path = out_dir / f'市委大事记（{y}年{m}月）_自动整理.docx'
        if path.exists():
            backup_generated_word(path)
    doc = Document()
    set_page_layout(doc)
    doc.core_properties.keywords = 'jrjj-v2-generated'
    doc.core_properties.title = f'市委大事记（{y}年{m}月）'
    header = doc.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.paragraph_format.space_after = Pt(0)
    header.paragraph_format.keep_with_next = True
    h_lvl = OxmlElement('w:outlineLvl')
    h_lvl.set(qn('w:val'), '0')  # Word outline level 1 (zero-based)
    header._p.get_or_add_pPr().append(h_lvl)
    hrun = header.add_run('市 委 大 事 记')
    hrun.font.name = '方正小标宋简体'
    hrun._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '方正小标宋简体')
    hrun.font.size = Pt(22)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(0)
    s_lvl = OxmlElement('w:outlineLvl')
    s_lvl.set(qn('w:val'), '1')  # Word outline level 2 (zero-based)
    subtitle._p.get_or_add_pPr().append(s_lvl)
    date_run = subtitle.add_run(f'（{y}年{m}月）')
    date_run.font.name = '楷体_GB2312'
    date_run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '楷体_GB2312')
    date_run.font.size = Pt(18)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pnum = footer.add_run()
    pnum.font.name = '仿宋_GB2312'
    pnum.font.size = Pt(12)
    field_start = OxmlElement('w:fldChar')
    field_start.set(qn('w:fldCharType'), 'begin')
    field_instr = OxmlElement('w:instrText')
    field_instr.set(qn('xml:space'), 'preserve')
    field_instr.text = ' PAGE '
    field_end = OxmlElement('w:fldChar')
    field_end.set(qn('w:fldCharType'), 'end')
    pnum._r.append(field_start)
    pnum._r.append(field_instr)
    pnum._r.append(field_end)
    for article in articles:
        doc.add_paragraph('')
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = Pt(28)
        p.paragraph_format.keep_with_next = True
        add_run(p, article['title'], '黑体')
        ppr = p._p.get_or_add_pPr()
        lvl = OxmlElement('w:outlineLvl')
        lvl.set(qn('w:val'), '2')
        ppr.append(lvl)
        for para in article['body']:
            if not para:
                continue
            body = doc.add_paragraph()
            body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            body.paragraph_format.first_line_indent = Cm(1.12)
            body.paragraph_format.line_spacing = Pt(28)
            add_run(body, para, '仿宋_GB2312')
    temp = path.with_suffix('.tmp.docx')
    doc.save(temp)
    temp.replace(path)
    return path


def audit_excel(path: Path) -> dict:
    """Read-only validation of the sole A hyperlink and A:H schema."""
    result = {
        'file': str(path), 'sheet': '九江新闻网', 'total': 0,
        'months': {}, 'issues': [], 'issue_count': 0,
    }
    if not path.is_file():
        result['issues'].append('找不到Excel：' + str(path))
        result['issue_count'] += 1
        return result
    try:
        book = load_workbook(path, read_only=False, data_only=False)
    except Exception as exc:
        result['issues'].append(f'Excel无法读取：{exc}')
        result['issue_count'] += 1
        return result

    def problem(message):
        result['issue_count'] += 1
        if len(result['issues']) < 30:
            result['issues'].append(message)

    try:
        if book.sheetnames != ['九江新闻网']:
            problem('工作表不是唯一的“九江新闻网”：' + '、'.join(book.sheetnames))
        if '九江新闻网' not in book.sheetnames:
            return result
        ws = book['九江新闻网']
        if norm_space(ws.cell(1, 7).value) == '原文链接':
            problem('检测到旧版A:I工作表，需运行一次同步或导出迁移为A:H')
            event_col = 9
            legacy = True
        else:
            event_col = 8
            legacy = False
            if [ws.cell(1, i).value for i in range(1, 9)] != HEADERS:
                problem('表头不是新版A:H结构')
        seen = {}
        base = SOURCES['jjxw']['base']
        for rowno in range(2, ws.max_row + 1):
            a, b, event = (ws.cell(rowno, n) for n in (1, 2, event_col))
            if not (a.value or b.value):
                continue
            result['total'] += 1
            if not get_cell_title(a):
                problem(f'第{rowno}行：A列缺少新闻标题')
            url = canonical_news_url(base, get_cell_url(a))
            if not url:
                problem(f'第{rowno}行：A列没有有效的新闻原文超链接')
            elif url in seen:
                problem(f'第{rowno}行：超链接与第{seen[url]}行重复：{url}')
            else:
                seen[url] = rowno
            if legacy:
                old = canonical_news_url(base, ws.cell(rowno, 7).value)
                if old and url and old != url:
                    problem(f'第{rowno}行：旧版A/G超链接冲突；A={url}；旧G={old}')
            elif legacy_link_candidate(ws.cell(rowno, 7).value):
                problem(f'第{rowno}行：G列说明中存在历史链接待复核标记')
            if isinstance(b.value, (date, datetime)):
                published = b.value.date().isoformat() if isinstance(b.value, datetime) else b.value.isoformat()
            else:
                published = date_from_text(str(b.value or ''))
            if not published:
                problem(f'第{rowno}行：B列发布日期缺失或无效')
            else:
                month = published[:7]
                result['months'][month] = result['months'].get(month, 0) + 1
            if event.value not in (None, ''):
                if isinstance(event.value, (date, datetime)):
                    event_day = event.value.date().isoformat() if isinstance(event.value, datetime) else event.value.isoformat()
                else:
                    event_day = date_from_text(str(event.value))
                if not event_day:
                    problem(f'第{rowno}行：{("I" if legacy else "H")}列事件日期无效：{event.value}')
    finally:
        book.close()
    result['months'] = dict(sorted(result['months'].items()))
    return result


def show_excel_audit(path: Path) -> int:
    """Display cumulative monthly counts and anomalies without changing files."""
    report = audit_excel(path)
    print(f'检查文件：{path}')
    print(f'累计新闻：{report["total"]} 条')
    print('按发布日期统计：')
    for month, count in report['months'].items():
        print(f'  {month}：{count} 条')
    if not report['months']:
        print('  暂无有效发布日期记录')
    print(f'异常：{report["issue_count"]} 项')
    for problem in report['issues']:
        print('  - ' + problem)
    if report['issue_count'] > len(report['issues']):
        print(f'  ...还有{report["issue_count"] - len(report["issues"])}项未显示')
    return 1 if report['issue_count'] else 0


def write_link_review_report(conflicts: list[dict], path: Path) -> None:
    """Write one replaceable UTF-8 CSV without changing any Excel cells."""
    if not conflicts:
        path.unlink(missing_ok=True)
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
            mode='w', prefix='.link-review-', suffix='.csv',
            encoding='utf-8-sig', newline='', dir=path.parent,
            delete=False) as handle:
        temporary = Path(handle.name)
        try:
            writer = csv.writer(handle)
            writer.writerow(['Excel行号', '新闻标题', 'A列实际超链接',
                             'G列原文链接', '需人工复核的原因'])
            for issue in conflicts:
                # Titles are untrusted web text; do not let Excel interpret
                # a leading formula prefix when opening this CSV.
                title = str(issue['title'])
                if title.lstrip().startswith(('=', '+', '-', '@')):
                    title = "'" + title
                writer.writerow([issue['row'], title,
                                 issue['a_url'], issue['g_url'], issue['reason']])
        except BaseException:
            temporary.unlink(missing_ok=True)
            raise
    try:
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)
    LOG.warning('有%d行原文链接冲突未能安全修复，详见：%s', len(conflicts), path)


def backup_spreadsheet(path: Path) -> None:
    """Keep one pre-run snapshot per day and only seven automatic Excel backups.

    Never delete files that do not match this program's autogenerated filename.
    """
    if not path.exists():
        return
    backups = path.parent / '备份'
    backups.mkdir(parents=True, exist_ok=True)
    prefix = f'{path.stem}_运行前_'
    pattern = re.compile(
        re.escape(prefix) + r'\d{8}_\d{6}_\d{6}' + re.escape(path.suffix))
    snapshots = sorted(
        (item for item in backups.glob(f'{prefix}*{path.suffix}')
         if item.is_file() and pattern.fullmatch(item.name)),
        key=lambda item: item.name)
    today = prefix + datetime.now().strftime('%Y%m%d') + '_'
    if any(item.name.startswith(today) for item in snapshots):
        LOG.info('今天已有Excel运行前备份，跳过重复备份')
    else:
        stamp = datetime.now().strftime('%Y%m%d_%H%M%S_%f')
        dest = backups / f'{prefix}{stamp}{path.suffix}'
        shutil.copy2(path, dest)
        snapshots.append(dest)
        LOG.info('运行前Excel备份：%s', dest)
    for old in sorted(snapshots, key=lambda item: item.name)[:-7]:
        old.unlink()
        LOG.info('清理超出最近7份的自动Excel备份：%s', old)


def previous_month(today: date | None = None) -> str:
    today = today or date.today()
    return (today.replace(day=1) - timedelta(days=1)).strftime('%Y-%m')


def verify_month_output(month: str, excel: Path, output_dir: Path, grace_days: int = 10) -> int:
    """Verify monthly output without network or database writes."""
    start, end = publication_window(month, grace_days)
    report = {
        'month': month, 'source': '九江新闻网', 'news': 0,
        'matched': 0, 'late_candidates': 0, 'late_confirmed': 0,
        'word_articles': 0, 'errors': [],
    }
    if not excel.exists():
        report['errors'].append('缺少Excel文件')
    else:
        try:
            book = load_workbook(excel, read_only=True)
            try:
                if book.sheetnames != ['九江新闻网']:
                    report['errors'].append('Excel必须只有九江新闻网工作表')
                elif book.active.cell(1, 7).value == '原文链接':
                    report['errors'].append('Excel还是旧版A:I九列结构，请先运行同步或导出迁移')
                else:
                    for row in book.active.iter_rows(min_row=2, max_col=8,
                                                     values_only=True):
                        if not row[0]:
                            continue
                        pub = str(row[1] or '')
                        if not start <= pub <= end:
                            continue  # Existing workbook can also contain other months.
                        report['news'] += 1
                        if not pub.startswith(month):
                            report['late_candidates'] += 1
                            if str(row[7] or '').startswith(month):
                                report['late_confirmed'] += 1
                        if row[2] == '命中':
                            report['matched'] += 1
            finally:
                book.close()
        except Exception as exc:
            report['errors'].append(f'Excel读取失败：{exc}')
    if not report['news']:
        report['errors'].append('没有新闻网新闻')
    if (output_dir / '链接冲突待复核.csv').is_file():
        report['errors'].append('仍有链接冲突待人工复核，请检查results/链接冲突待复核.csv')
    year, month_number = map(int, month.split('-'))
    words = list(output_dir.glob(
        f'市委大事记（{year}年{month_number}月）*.docx')) if output_dir.exists() else []
    if not words:
        report['errors'].append('没有生成指定月份Word')
    for path in words:
        try:
            doc = Document(path)
            # The main title and month subtitle now have outline levels 1
            # and 2. Only level 3 headings count as article entries.
            report['word_articles'] += sum(
                any(node.get(qn('w:val')) == '2'
                    for node in p._p.xpath('./w:pPr/w:outlineLvl'))
                for p in doc.paragraphs)
        except Exception as exc:
            report['errors'].append(f'Word读取失败：{path.name} — {exc}')
    if not report['word_articles']:
        report['errors'].append('Word没有新闻正文')
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / '抓取检查报告.json').write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    LOG.info('导出验证：%s', json.dumps(report, ensure_ascii=False))
    return 1 if report['errors'] else 0


def args_parser():
    """Keep the user-facing CLI small; tuning lives in config.json."""
    parser = argparse.ArgumentParser(
        description='九江新闻网市委大事记：默认采集上个月，Excel/Word保存到results目录',
        epilog=('示例：python main.py | python main.py 2026-09 | '
                'python main.py --test | python main.py 2026-09 --export'))
    parser.add_argument(
        'month', nargs='?', metavar='YYYY-MM',
        help='目标月份，默认上一个自然月，如 2026-09')
    operation = parser.add_mutually_exclusive_group()
    operation.add_argument('--test', action='store_true',
                           help='运行离线测试，不联网也不创建采集文件')
    operation.add_argument('--check-names', action='store_true',
                           help='离线检查生效的人名数量与完整名单，不访问新闻网站')
    operation.add_argument('--check-excel', action='store_true',
                           help='离线检查累计Excel各月数量、A/G超链接、重复新闻及日期')
    operation.add_argument('--update-leaders', action='store_true',
                           help='联网更新省委、人大、政府、政协领导名单至names.auto.txt')
    operation.add_argument('--sync', action='store_true',
                           help='只同步新闻到Excel和正文缓存，不生成Word')
    operation.add_argument('--export', action='store_true',
                           help='根据现有Excel与缓存重新生成Word，可补抓缺失正文')
    operation.add_argument('--verify', action='store_true',
                           help='离线检查指定月份的Excel和Word是否完整')
    return parser


def main(argv=None, *, base_dir: Path | None = None) -> int:
    args = args_parser().parse_args(argv)
    # Resolve paths relative to this script, not the terminal's current directory.
    # base_dir is an internal test seam, not an extra command-line option.
    root = Path(base_dir) if base_dir is not None else Path(__file__).resolve().parent
    if args.test:
        import subprocess
        return subprocess.call(
            [sys.executable, '-m', 'unittest', '-v', 'test_main'], cwd=root)
    if args.update_leaders:
        # Separate workflow from news collection; errors retain the old file.
        try:
            groups, updated = update_leaders(root)
        except Exception as exc:
            print(f'领导名单更新失败：{exc}', file=sys.stderr)
            return 1
        for section, people in groups.items():
            print(f'{section}：{len(people)} 人 — ' + '、'.join(people))
        print(f'已保存至 {root / AUTO_NAMES_FILE}' if updated
              else '名单未变化，无需覆盖已有文件')
        print('合计去重人数：' +
              str(len({name for vals in groups.values() for name in vals})))
        return 0
    if args.check_excel:
        return show_excel_audit(root / 'results' / 'jiujiang_news.xlsx')
    if args.check_names:
        # Read-only: no network, no SQLite/Excel outputs or config creation.
        cfg = load_config(root / 'config.json', create_missing=False)
        names = get_names(cfg)
        print(f'姓名来源：{cfg["names_source"]}')
        print(f'生效姓名：{len(names)} 人')
        snapshot = root / AUTO_NAMES_FILE
        if (snapshot.is_file() and
                '初始网页快照' in snapshot.read_text(encoding='utf-8-sig').splitlines()[0]):
            print('提醒：当前 names.auto.txt 为网页缓存快照，尚未通过实时抓取复核。')
        for i, name in enumerate(names, 1):
            print(f'{i:>2}. {name}')
        overlaps = [(short, long) for short in names for long in names
                    if short != long and short in long]
        if overlaps:
            print('注意：名单中存在姓名包含关系；若短姓名只出现在长姓名内部，'
                  '不会单独计为命中：')
            for short, long in overlaps:
                print(f'  {short} ⊂ {long}')
        return 0

    logging.basicConfig(level=logging.INFO, format='%(asctime)s %(levelname)s %(message)s',
                        datefmt='%H:%M:%S')
    month = args.month or previous_month()
    if not re.fullmatch(r'20\d{2}-(0[1-9]|1[0-2])', month):
        args_parser().error('月份格式错误，请输入 YYYY-MM，例如 2026-09')

    results = root / 'results'
    excel = results / 'jiujiang_news.xlsx'
    database = results / 'jiujiang_cache.sqlite3'
    debug_dir = results / 'debug_html'

    if args.verify:
        return verify_month_output(month, excel, results, load_config(root / 'config.json')['late_publication_days'])

    mode = 'sync' if args.sync else 'export' if args.export else 'all'
    cfg = load_config(root / 'config.json')
    year = int(month[:4])
    results.mkdir(parents=True, exist_ok=True)

    LOG.info('目标月份：%s；延后%d天检索；发布日期范围：%s 至 %s；模式：%s；输出目录：%s',
             month, cfg['late_publication_days'],
             *publication_window(month, cfg['late_publication_days']), mode, results)
    backup_spreadsheet(excel)
    book = open_book(excel)
    with connect_db(database) as con:
        link_conflicts = []
        existing = existing_rows(book, con, link_conflicts)
        write_link_review_report(link_conflicts, results / '链接冲突待复核.csv')
        session = build_session(cfg['user_agent'])
        new = {source: [] for source in SOURCES}
        if mode != 'export':
            # Scan the complete month + grace window, then preserve pages even
            # if a later API request fails. Do not issue a misleading Word.
            crawl_status = {}
            new = crawl(session, con, existing, cfg, year, debug_dir, month,
                        status=crawl_status)
            write_new_rows(book, new)
            save_book_atomic(book, excel)
            if crawl_status.get('complete') is False:
                LOG.error('本次采集不完整：%s。已保存Excel和缓存，请稍后重试。',
                          crawl_status.get('error', '未知分页错误'))
                return 1
            fetch_article_details(session, con, cfg, year, month, new, book, 'month')
        else:
            # Reuse cached news and honor manual selection in column E.
            fetch_article_details(session, con, cfg, year, month, new, book, 'month')

        names = get_names(cfg)
        selected = refresh_book(book, con, names, cfg)
        save_book_atomic(book, excel)
        LOG.info('Excel已更新：%s；配置姓名 %d 人', excel, len(names))
        if mode != 'sync':
            chosen, failed = deduplicate_for_month(selected, month)
            word = build_word(month, chosen, results)
            if word:
                LOG.info('Word已生成：%s（%d篇，正文去重）', word, len(chosen))
            else:
                LOG.error('该月份没有成功提取并满足筛选条件的正文，未生成Word；请检查日期、名单和采集日志')
            if failed:
                LOG.error('有%d篇已选新闻正文未抓取成功，Word可能缺少文章；下次将重试', len(failed))
            if not word or failed:
                return 1  # Do not report a complete monthly export.
    if link_conflicts:
        LOG.error('仍有%d行A/G链接待人工复核，详见results/链接冲突待复核.csv；'
                  '本次结果不标记为完整成功', len(link_conflicts))
        return 1
    LOG.info('完成。输出文件位于：%s', results)
    return 0


def configure_cli_encoding() -> None:
    """Make Chinese terminal output safe in Windows CI and redirected consoles.

    Some Windows environments default to cp1252. Setting the encoding before
    argparse handles --help prevents UnicodeEncodeError without requiring the
    caller to manually change the system code page.
    """
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, 'reconfigure'):
            try:
                stream.reconfigure(encoding='utf-8', errors='backslashreplace')
            except (OSError, ValueError):
                pass


if __name__ == '__main__':
    configure_cli_encoding()
    try:
        sys.exit(main())
    except KeyboardInterrupt:
        LOG.error('用户中断；已落盘的缓存、Excel仍被保留')
        sys.exit(130)
    except PermissionError as exc:
        LOG.error('文件被占用，请先关闭 Excel/Word 再运行：%s', exc)
        sys.exit(2)
