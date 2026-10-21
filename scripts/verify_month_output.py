"""Validate the single-source monthly Excel and Word artifacts."""
import argparse
import json
from pathlib import Path
from openpyxl import load_workbook
from docx import Document

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--month', required=True)
    parser.add_argument('--folder', default='results')
    args = parser.parse_args()
    folder = Path(args.folder)
    report = dict(month=args.month, source='九江新闻网',
                  news=0, matched=0, word_articles=0, errors=[])
    excel = folder / 'jiujiang_news.xlsx'
    if not excel.exists():
        report['errors'].append('缺少Excel文件')
    else:
        wb = load_workbook(excel, read_only=True)
        if wb.sheetnames != ['九江新闻网']:
            report['errors'].append('Excel必须只有九江新闻网工作表')
        else:
            for row in wb.active.iter_rows(min_row=2, max_col=8, values_only=True):
                if not row[0]:
                    continue
                report['news'] += 1
                if not str(row[1] or '').startswith(args.month):
                    report['errors'].append('包含非目标月份的记录')
                if row[2] == '命中':
                    report['matched'] += 1
        wb.close()
    if not report['news']:
        report['errors'].append('没有新闻网新闻')
    for path in folder.glob('*.docx'):
        doc = Document(path)
        report['word_articles'] += sum(
            bool(p._p.xpath('./w:pPr/w:outlineLvl')) for p in doc.paragraphs)
    if not report['word_articles']:
        report['errors'].append('Word没有新闻正文')
    folder.mkdir(parents=True, exist_ok=True)
    (folder / '抓取检查报告.json').write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 1 if report['errors'] else 0

if __name__ == '__main__':
    raise SystemExit(main())
