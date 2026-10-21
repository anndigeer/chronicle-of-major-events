from pathlib import Path
from openpyxl import load_workbook
from docx import Document
root=Path('results')
file=root/'jiujiang_news.xlsx'
assert file.exists(), 'Missing Excel workbook'
wb=load_workbook(file,read_only=True)
for name in ('九江市政府','九江新闻网'):
    assert name in wb.sheetnames, f'Missing worksheet: {name}'
    count=max(0,wb[name].max_row-1)
    print(f'{name}: {count} news records')
    assert count > 0, f'No news records: {name}'
words=list(root.glob('*.docx'))
assert words, 'No Word document was produced'
for path in words:
    paragraphs=[p.text for p in Document(path).paragraphs if p.text.strip()]
    print(path.name, len(paragraphs), 'nonempty paragraphs')
    assert len(paragraphs)>2, f'Empty Word document: {path}'
