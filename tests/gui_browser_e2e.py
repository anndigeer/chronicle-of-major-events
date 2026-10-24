#!/usr/bin/env python3
"""Real Chromium browser smoke + core CLI workflow, with an offline news fixture.

Run in CI:
    python -m pip install playwright
    python -m playwright install --with-deps chromium
    python tests/gui_browser_e2e.py
"""
from __future__ import annotations

import json
import os
from pathlib import Path
import sys
import tempfile
import threading
from http.server import ThreadingHTTPServer

from docx import Document
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))
import webui  # noqa: E402

ARTIFACTS = REPO / 'test-artifacts' / 'gui'
ARTIFACTS.mkdir(parents=True, exist_ok=True)

# Runs production main.py via importlib; only the external news endpoints
# are replaced. No fixture changes are committed or written into user data.
FIXTURE_RUNNER = r'''
from __future__ import annotations
import importlib.util
import os
from pathlib import Path
import sys

path = Path(os.environ['CHRONICLE_REAL_CORE'])
spec = importlib.util.spec_from_file_location('chronicle_real_core', path)
core = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = core
spec.loader.exec_module(core)
root = Path(__file__).resolve().parent

def fake_fetch_list(session, page, cfg, debug_dir):
    base = 'https://www.jjxw.cn/article/2026/'
    if page == 0:
        rows = [
            core.News('jjxw', '陈云调研工程丙', base+'10/02/1004.html', '2026-10-02'),
            core.News('jjxw', '普通政策解读', base+'09/20/1003.html', '2026-09-20'),
            core.News('jjxw', '陈云开展调研乙', base+'09/19/1002.html', '2026-09-19'),
            core.News('jjxw', '陈云主持会议甲', base+'09/18/1001.html', '2026-09-18'),
        ]
        if (root/'add_new_article.flag').exists():
            rows.insert(1, core.News(
                'jjxw', '陈云参加会议丁', base+'09/21/1005.html', '2026-09-21'))
        return rows, True
    if page == 1:
        return [core.News(
            'jjxw', '早期报道', base+'08/31/9900.html', '2026-08-31')], True
    return [], True

def fake_fetch_html(session, url, timeout):
    cases = {
      '1001.html': ('陈云主持会议甲', '9月18日，陈云主持召开会议甲，研究推进重点项目与有关工作。'),
      '1002.html': ('陈云开展调研乙', '9月19日，陈云深入基层调研乙，检查重点项目建设情况。'),
      '1003.html': ('普通政策解读', '9月20日，有关部门开展政策解读工作并组织业务交流活动。'),
      '1004.html': ('陈云调研工程丙', '9月30日，陈云赴工程丙调研并部署后续建设重点工作。'),
      '1005.html': ('陈云参加会议丁', '9月21日，陈云参加会议丁并就落实重点工作提出要求。'),
    }
    for ending, (title, body) in cases.items():
        if url.endswith(ending):
            return ('<html><h1>'+title+'</h1><div class="article-content">'
                    +'<p>'+body*4+'</p><p>'+body*4+'</p></div></html>')
    raise AssertionError('unexpected URL in offline CI: '+url)

core.fetch_list = fake_fetch_list
core.fetch_html = fake_fetch_html
sys.exit(core.main(sys.argv[1:], base_dir=root))
'''


def article_headings(doc_file: Path) -> list[str]:
    doc = Document(doc_file)
    return [
        p.text for p in doc.paragraphs
        if any(node.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val') == '2'
               for node in p._p.xpath('./w:pPr/w:outlineLvl'))
    ]


def by_url(ws, suffix: str) -> int:
    for n in range(2, ws.max_row + 1):
        link = ws.cell(n, 1).hyperlink
        if link and link.target and link.target.endswith(suffix):
            return n
    raise AssertionError('Missing fixture URL suffix: '+suffix)


def run_browser_test(page, jobs, folder: Path):
    results = folder/'results'
    workbook = results/'jiujiang_news.xlsx'
    word = results/'市委大事记（2026年9月）.docx'

    page.goto('http://127.0.0.1:'+str(page.context._ci_server_port)+'/', wait_until='domcontentloaded')
    page.locator('#month').fill('2026-09')
    assert page.get_by_role('button', name='增量采集 + 更新 Word').count() == 1
    assert page.get_by_role('button', name='打开 Excel').count() == 1
    page.get_by_role('button', name='增量采集 + 更新 Word').click()
    page.wait_for_function("""() => document.getElementById('status').textContent.includes('最近一次成功')""",
                           timeout=60000)
    assert workbook.is_file(), 'Full GUI run did not create Excel'
    assert word.is_file(), 'Full GUI run did not create Word'
    wb = load_workbook(workbook)
    ws = wb.active
    assert ws.max_row == 5, f'Unexpected article count: {ws.max_row}'
    assert [ws.cell(1, i).value for i in range(1, 9)] == [
        '新闻标题','发布日期','自动匹配','命中姓名',
        '人工选择','正文状态','说明','事件日期']
    assert ws.cell(by_url(ws,'1004.html'),8).value == '2026-09-30'
    headings = article_headings(word)
    assert '陈云主持会议甲' in headings and '陈云开展调研乙' in headings
    assert '普通政策解读' not in headings
    assert '陈云调研工程丙' in headings

    # User edits existing Excel manually before clicking "rebuild Word".
    excluded_row = by_url(ws, '1002.html')
    included_row = by_url(ws, '1003.html')
    ws.cell(excluded_row,5).value = '排除'
    ws.cell(included_row,5).value = '收录'
    ws.cell(included_row,8).value = '2026-09-20'
    wb.save(workbook)
    page.get_by_role('button', name='按人工收录/排除重建 Word').click()
    page.wait_for_function("""() => document.getElementById('status').textContent.includes('最近一次成功')""",
                           timeout=60000)
    headings = article_headings(word)
    assert '陈云开展调研乙' not in headings, 'Exclusion still appeared in Word'
    assert '普通政策解读' in headings, 'Manual include not applied'
    assert load_workbook(workbook).active.cell(excluded_row,5).value == '排除'

    # A subsequent sync must only append new rows and leave the Word alone.
    word_bytes = word.read_bytes()
    (folder/'add_new_article.flag').write_text('yes', encoding='ascii')
    page.get_by_role('button', name='仅增量同步 Excel').click()
    page.wait_for_function("""() => document.getElementById('status').textContent.includes('最近一次成功')""",
                           timeout=60000)
    updated = load_workbook(workbook)
    assert updated.active.max_row == 6
    assert updated.active.cell(by_url(updated.active,'1002.html'),5).value == '排除'
    assert word.read_bytes() == word_bytes, '--sync unexpectedly rewrote Word'

    page.get_by_role('button', name='按人工收录/排除重建 Word').click()
    page.wait_for_function("""() => document.getElementById('status').textContent.includes('最近一次成功')""",
                           timeout=60000)
    assert '陈云参加会议丁' in article_headings(word)
    assert '陈云开展调研乙' not in article_headings(word)

    page.get_by_role('button', name='核验成果').click()
    page.wait_for_function("""() => document.getElementById('status').textContent.includes('最近一次成功')""",
                           timeout=60000)
    report = json.loads((results/'抓取检查报告.json').read_text(encoding='utf-8'))
    assert not report['errors'], report
    assert report['word_articles'] == len(article_headings(word))
    page.get_by_role('button', name='检查 Excel').click()
    page.wait_for_function("""() => document.getElementById('status').textContent.includes('最近一次成功')""",
                           timeout=60000)
    page.get_by_role('button', name='检查名单').click()
    page.wait_for_function("""() => document.getElementById('status').textContent.includes('最近一次成功')""",
                           timeout=60000)
    assert '陈云' in page.locator('#log').inner_text()
    page.screenshot(path=str(ARTIFACTS/'gui-success.png'),full_page=True)
    (ARTIFACTS/'summary.json').write_text(json.dumps({
       'result':'passed','excel_rows':updated.active.max_row,
       'word_articles':report['word_articles'],
       'manual_exclude':True,'manual_include':True,
       'sync_did_not_change_word':True},ensure_ascii=False,indent=2),encoding='utf-8')


def main():
    import webui
    with tempfile.TemporaryDirectory(prefix='chronicle-gui-ci-') as d:
        root=Path(d)
        (root/'main.py').write_text(FIXTURE_RUNNER,encoding='utf-8')
        (root/'names.txt').write_text('陈云\n',encoding='utf-8')
        (root/'config.json').write_text(json.dumps({
            'names':['陈云'],'include_body_matches':True,'delay_seconds':0,
            'max_pages':10,'timeout_seconds':10,'late_publication_days':10
        },ensure_ascii=False),encoding='utf-8')
        os.environ['CHRONICLE_REAL_CORE']=str(REPO/'main.py')
        jobs=webui.Jobs(root,sys.executable)
        token='ci-test-session-token'
        handler=type('CIBrowserHandler',(webui.LocalHandler,),{
            'jobs':jobs,'token':token,
            'page':webui.PAGE.replace('__TOKEN__',token)})
        server=ThreadingHTTPServer(('127.0.0.1',0),handler)
        thread=threading.Thread(target=server.serve_forever,daemon=True)
        thread.start()
        try:
            with sync_playwright() as playwright:
                browser=playwright.chromium.launch(headless=True)
                context=browser.new_context(viewport={'width':1320,'height':900},
                                            locale='zh-CN')
                # Test metadata used by helper without exposing public socket.
                context._ci_server_port=server.server_port
                page=context.new_page()
                try:
                    run_browser_test(page,jobs,root)
                    print('PASS: Chromium GUI + offline collection + manual edits '
                          '+ increment + export + verification')
                except Exception:
                    page.screenshot(path=str(ARTIFACTS/'gui-failure.png'),
                                    full_page=True)
                    (ARTIFACTS/'last-job-log.txt').write_text(
                        jobs.status('2026-09')['log'],encoding='utf-8')
                    raise
                finally:
                    browser.close()
        finally:
            server.shutdown()
            server.server_close()
            thread.join(timeout=10)


if __name__=='__main__':
    main()
