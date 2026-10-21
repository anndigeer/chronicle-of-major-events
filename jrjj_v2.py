#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""双来源市委大事记采集器。详见 README_使用说明.md。"""
from __future__ import annotations

import argparse
import json
from contextlib import contextmanager
import logging
import os
import re
import socket
import shutil
import sqlite3
import sys
import threading
import time
from source_feeds import parse_gov_xml, parse_jjxw_json
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from urllib.parse import urljoin, urlsplit, urlunsplit, unquote

import requests
from bs4 import BeautifulSoup, Tag
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
from urllib3.util import connection as urllib3_connection

LOG = logging.getLogger('jrjj')
SOURCES = {
    'gov': {'sheet': '九江市政府', 'base': 'https://www.jiujiang.gov.cn/zwzx/jrjj/', 'index': 'index_{n}.html'},
    'jjxw': {'sheet': '九江新闻网', 'base': 'https://www.jjxw.cn/column/668/', 'index': 'index{n}.html'},
}
DEFAULT_NAMES = ['陈云', '邓永翔', '黄强', '陈永奇', '庄兆林', '黄喜忠']
HEADERS = ['新闻标题', '发布日期', '自动匹配', '命中姓名', '人工选择', '正文状态', '原文链接', '说明']
RE_DATE = re.compile(r'(20\d{2})[-年/.](1[0-2]|0?[1-9])[-月/.](3[01]|[12]\d|0?[1-9])(?:日)?(?!\d)')
RE_ISO = re.compile(r'20\d{2}-\d{2}-\d{2}')
INVALID_URL = re.compile(r'^(?:#|javascript:|mailto:|tel:)', re.I)
CONTENT_SELECTORS = [
    '#artibody > div.trs_editor_view', '#artibody', '#zoom',
    '.TRS_Editor', '.trs_editor_view', '.article-content', '.article_content',
    '.article-body', '.article_body', '.article-text', '.article_text',
    '.news-content', '.news_content', '.content-article', '.detail-content',
    '.detail_content', '.text-content', '.wzcon', '#content', '.content',
]
BLACKLIST_CLASS = re.compile(r'(related|recommend|copyright|share|comment|footer|nav|ad[-_]|pagination|source)', re.I)


@dataclass
class News:
    source: str
    title: str
    url: str
    pub_date: str


def norm_space(s: object) -> str:
    return re.sub(r'\s+', ' ', str(s or '').replace('\u200b', '')).strip()


def clean_title(s: object) -> str:
    return norm_space(s).strip(' -|_·\u3000')


def normalize_url(base: str, href: str) -> str:
    href = (href or '').strip()
    if not href or INVALID_URL.match(href):
        return ''
    # Some sites write a root-relative path without '/', e.g. column/668/article/...
    if re.match(r'^(?:column/\d+/|zwzx/jrjj/)', href, re.I):
        parts = urlsplit(base)
        href = f'{parts.scheme}://{parts.netloc}/' + href
    url = urljoin(base, href)
    p = urlsplit(url)
    if p.scheme not in {'https', 'http'}:
        return ''
    path = p.path
    # The 668 column has published relative links that resolve to a 404
    # (/column/668/article/...), while articles live at /article/... .
    if 'jjxw.cn' in p.netloc.lower():
        path = re.sub(r'^/(?:column/\d+/)+article/', '/article/', path)
    return urlunsplit((p.scheme, p.netloc.lower(), path, p.query, ''))


def date_from_text(raw: str) -> str:
    m = RE_DATE.search(raw or '')
    if not m:
        return ''
    try:
        return date(*map(int, m.groups())).isoformat()
    except ValueError:
        return ''


def date_from_url(url: str) -> str:
    m = re.search(r'/((?:19|20)\d{2})/(0[1-9]|1[0-2])/(0[1-9]|[12]\d|3[01])/', url)
    if m:
        return date_from_text('-'.join(m.groups()))
    m = re.search(r'/((?:19|20)\d{2})(0[1-9]|1[0-2])/(?:t|\d)', url)
    if m:
        # Publication date may differ from the URL date; this fallback is approximate.
        d = re.search(r't((?:19|20)\d{2})(0[1-9]|1[0-2])(0[1-9]|[12]\d|3[01])', url)
        if d:
            return date_from_text('-'.join(d.groups()))
    return ''


def article_like(url: str, source: str) -> bool:
    p = urlsplit(url)
    if not p.netloc:
        return False
    lower = p.path.lower()
    if lower.endswith(('.png', '.jpg', '.jpeg', '.pdf', '.doc', '.docx', '.xls', '.xlsx', '.js', '.css')):
        return False
    if source == 'gov':
        if p.netloc.endswith('jiujiang.gov.cn'):
            return bool(re.search(r'/(?:19|20)\d{2}(?:\d{2})?/|/t20\d{6}_', lower)) and lower.endswith(('.html', '.htm', '.shtml'))
        return bool(re.search(r'/(?:19|20)\d{2}/|t20\d{6}', lower)) and lower.endswith(('.html', '.htm', '.shtml'))
    if source == 'jjxw':
        return 'jjxw.cn' in p.netloc and (bool(re.search(r'/article/', lower)) or
                 bool(re.search(r'/(?:19|20)\d{2}/\d{2}/\d{2}/', lower))) and lower.endswith(('.html', '.htm', '.shtml'))
    return False


def detect_title_and_date(node: Tag, anchor: Tag, source: str, base: str) -> News | None:
    href = anchor.get('href', '')
    url = normalize_url(base, href)
    if not url or not article_like(url, source):
        return None
    title = clean_title(anchor.get('title') or anchor.get_text(' ', strip=True))
    if len(title) < 4 or len(title) > 240:
        return None
    # Use node-local date, not date from a neighboring or unrelated article.
    raw = node.get_text(' ', strip=True)
    if len(raw) > 1800:
        raw = ''
    day = date_from_text(raw) or date_from_text(anchor.get('data-date', '') or '') or date_from_url(url)
    if not day:
        return None
    return News(source, title, url, day)


def parse_listing(html: str, source: str, base: str) -> list[News]:
    """Parse legacy and new list layouts. Never scan undated navigation links as news."""
    soup = BeautifulSoup(html, 'html.parser')
    # A list item normally contains its title and own date, regardless of CSS class.
    scoped = soup.select('li, article, tr, .news-item, .list-item, .news_item, .list_item')
    items: list[News] = []
    seen: set[str] = set()
    for node in scoped:
        anchors = [a for a in node.select('a[href]') if article_like(normalize_url(base, a.get('href', '')), source)]
        if len(anchors) != 1:  # outer wrapper of several articles isn't an individual item
            continue
        rec = detect_title_and_date(node, anchors[0], source, base)
        if rec and rec.url not in seen:
            seen.add(rec.url)
            items.append(rec)
    if items:
        return items
    # Some CMSes render an <h3><a>title</a></h3><p>date</p> pair inside a DIV.
    for a in soup.select('a[href]'):
        if not article_like(normalize_url(base, a.get('href', '')), source):
            continue
        for parent in list(a.parents)[:4]:
            if not isinstance(parent, Tag) or parent.name in {'body', 'html', 'main'}:
                break
            if len(parent.select('a[href]')) > 3:
                continue
            rec = detect_title_and_date(parent, a, source, base)
            if rec:
                if rec.url not in seen:
                    seen.add(rec.url)
                    items.append(rec)
                break
    return items


def next_page_urls(source: str, n: int) -> list[str]:
    info = SOURCES[source]
    if n == 0:
        return [info['base'] + 'index.html', info['base']]
    first = info['base'] + info['index'].format(n=n if source == 'gov' else n + 1)
    # Different CMS installations number the second page as index_1/index2.
    alt = (info['base'] + f'index_{n + 1}.html') if source == 'gov' else (info['base'] + f'index{n}.html')
    return list(dict.fromkeys([first, alt]))


def config_defaults() -> dict:
    return {
        'names': DEFAULT_NAMES,
        'include_body_matches': True,
        'max_pages': 170,
        'stop_after_known_pages': 3,
        'delay_seconds': 0.4,
        'timeout_seconds': 20,
        'user_agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/126.0 Safari/537.36',
    }


def load_config(path: Path) -> dict:
    if not path.exists():
        path.write_text(json.dumps(config_defaults(), ensure_ascii=False, indent=2), encoding='utf-8')
        LOG.info('已创建配置文件：%s', path)
    data = json.loads(path.read_text(encoding='utf-8'))
    # Ignore the obsolete option in existing configuration files.
    data.pop('fetch_ifeng_names', None)
    cfg = {**config_defaults(), **data}
    cfg['names'] = sorted({norm_space(x) for x in cfg['names'] if norm_space(x) and len(norm_space(x)) >= 2})
    return cfg


def build_session(user_agent: str) -> requests.Session:
    session = requests.Session()
    session.headers.update({'User-Agent': user_agent, 'Accept-Language': 'zh-CN,zh;q=0.9', 'Accept': 'text/html,application/xhtml+xml'})
    retry = Retry(total=3, connect=2, read=2, backoff_factor=0.8,
                  status_forcelist=(429, 500, 502, 503, 504), respect_retry_after_header=True)
    adapter = HTTPAdapter(max_retries=retry, pool_connections=8, pool_maxsize=8)
    session.mount('https://', adapter)
    session.mount('http://', adapter)
    return session


def fetch_html(session: requests.Session, url: str, timeout: int) -> str:
    r = session.get(url, timeout=timeout)
    r.raise_for_status()
    if not r.encoding or r.encoding.lower() in {'iso-8859-1', 'ascii'}:
        r.encoding = r.apparent_encoding or 'utf-8'
    return r.text


_GOV_IPV4_LOCK = threading.RLock()


def fetch_government_xml(session: requests.Session, url: str, timeout: int) -> bytes:
    """Fetch annual government XML with requests.get; prefer IPv4.

    The government site may have a broken IPv6 route from a cloud runner.
    urllib3's address-family selector is changed only during the IPv4
    attempt and always restored afterward. This command-line program uses
    sequential networking; callers must not run parallel requests during
    this scoped override. A second normal-stack attempt is used only when
    the IPv4 connection itself fails.

    Keeping response.content as bytes lets ElementTree honor the XML
    declaration's encoding instead of guessing through response.text.
    """
    headers = {
        'User-Agent': session.headers.get('User-Agent', 'Python requests'),
        'Accept': 'application/xml, text/xml;q=0.9, */*;q=0.8',
        'Accept-Language': session.headers.get('Accept-Language', 'zh-CN,zh;q=0.9'),
        'Referer': f'https://www.jiujiang.gov.cn/zwzx/jrjj/?year={url.rsplit("data_", 1)[-1].split(".", 1)[0]}&page=1',
    }
    # Optional official-access credential; never store it in the repository.
    api_key = os.environ.get('JIUJIANG_API_KEY', '').strip()
    if api_key:
        headers['x-api-key'] = api_key
    connect_timeout = min(max(float(timeout), 1.0), 6.0)
    read_timeout = max(float(timeout), 2.0)

    def get_once() -> bytes:
        # Direct requests.get, no disabled TLS verification or tokenized URL.
        with requests.get(
            url,
            headers=headers,
            timeout=(connect_timeout, read_timeout),
            allow_redirects=True,
        ) as response:
            response.raise_for_status()
            payload = response.content
            if not payload.strip():
                raise ValueError('政府年度 XML 请求成功，但响应内容为空')
            return payload

    try:
        with _GOV_IPV4_LOCK:
            old_family = urllib3_connection.allowed_gai_family
            urllib3_connection.allowed_gai_family = lambda: socket.AF_INET
            try:
                return get_once()
            finally:
                urllib3_connection.allowed_gai_family = old_family
    except requests.exceptions.ConnectionError as first_error:
        if isinstance(first_error, (requests.exceptions.SSLError, requests.exceptions.ProxyError)):
            raise
        LOG.warning('政府 XML IPv4 连接失败，尝试系统默认 IP 协议栈：%s', first_error)
        try:
            return get_once()
        except requests.exceptions.ConnectionError as second_error:
            raise requests.exceptions.ConnectionError(
                f'政府 XML 两次 GET 均无法建立连接（IPv4 优先、默认协议栈回退）。'
                f'网站或云服务器网络可能限制访问；首次：{first_error}；'
                f'回退：{second_error}'
            ) from second_error


def fetch_with_browser(url: str, timeout: int) -> str:
    """Optional Chromium fallback. Runs only with --browser and installed playwright."""
    try:
        from playwright.sync_api import sync_playwright
    except ImportError as exc:
        raise RuntimeError('浏览器回退需要 pip install playwright && python -m playwright install chromium') from exc
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        try:
            page = browser.new_page(locale='zh-CN')
            page.goto(url, wait_until='domcontentloaded', timeout=timeout * 1000)
            page.wait_for_timeout(1400)
            return page.content()
        finally:
            browser.close()


def discover_next_url(html: str, current_url: str, source: str) -> str:
    soup = BeautifulSoup(html, 'html.parser')
    for a in soup.select('a[href]'):
        label = norm_space(a.get_text(' ', strip=True))
        if label not in {'下一页', '下页', '后页', '下一页>', '下一页 »'} and a.get('rel') != ['next']:
            continue
        href = normalize_url(current_url, a.get('href', ''))
        if href and href != current_url and urlsplit(href).netloc == urlsplit(current_url).netloc:
            return href
    return ''


def fetch_list(session: requests.Session, source: str, page: int, cfg: dict,
               use_browser: bool, debug_dir: Path, preferred_url: str = '',
               year: int | None = None) -> tuple[list[News], bool, str]:
    """Read XML / JSON feeds instead of crawling changing HTML listing pages.

    Government XML is one file per year; Jiujiang News uses pageNumber
    starting from 1 and caps each response at 100 articles.
    """
    if source == 'gov':
        if page:
            return [], True, ''
        url = f"https://www.jiujiang.gov.cn/zwzx/jrjj/data_{year or date.today().year}.xml"
    elif source == 'jjxw':
        url = (f'https://api.jjxw.cn/api/getArticles?cid=668&pageSize=100'
               f'&pageNumber={page + 1}&containSubCols=1')
    else:
        raise ValueError(f'Unsupported news source: {source}')

    original_accept = session.headers.get('Accept')
    original_referer = session.headers.get('Referer')
    try:
        if source == 'jjxw':
            session.headers['Accept'] = 'application/json, text/plain, */*'
            session.headers['Referer'] = 'https://www.jjxw.cn/'
        if source == 'gov':
            response_data = fetch_government_xml(session, url, cfg['timeout_seconds'])
        else:
            response_data = fetch_html(session, url, cfg['timeout_seconds'])
        rows = parse_gov_xml(response_data) if source == 'gov' else parse_jjxw_json(response_data)
        news = [News(source, title, link, published) for title, link, published in rows]
        if not news:
            if source == 'jjxw' and page > 0:
                LOG.info('九江新闻网 API 第%d页没有更多新闻，分页结束', page + 1)
                return [], True, ''
            raise ValueError('接口可访问，但未识别到新闻记录，请核对 XML/JSON 结构')
        LOG.info('%s 数据接口第%d页获取 %d 条新闻', source, page + 1, len(news))
        return news, True, ''
    except Exception as exc:
        debug_dir.mkdir(parents=True, exist_ok=True)
        (debug_dir / f'{source}_feed_page_{page + 1}.txt').write_text(
            f'URL: {url}\nERROR: {type(exc).__name__}: {exc}\n', encoding='utf-8')
        LOG.error('%s 数据接口请求或解析失败：%s — %s', source, url, exc)
        return [], False, ''
    finally:
        if original_accept is None:
            session.headers.pop('Accept', None)
        else:
            session.headers['Accept'] = original_accept
        if original_referer is None:
            session.headers.pop('Referer', None)
        else:
            session.headers['Referer'] = original_referer


def get_names(cfg: dict) -> list[str]:
    """Use only personal names explicitly supplied in the configuration."""
    return sorted(set(cfg['names']))


def clean_body(soup: BeautifulSoup) -> tuple[str, list[str]]:
    """Choose the best article container and omit nav/recommendation/ads before matching."""
    for tag in soup.select('script, style, noscript, iframe, form, svg, nav, footer, aside'):
        tag.decompose()
    for tag in list(soup.find_all(True)):
        if tag.attrs is None:  # A parent may have been decomposed already.
            continue
        attrs = f"{tag.get('id', '')} {' '.join(tag.get('class', []))}"
        if BLACKLIST_CLASS.search(attrs) and tag.name in {'div', 'section', 'ul'}:
            tag.decompose()
    candidates: list[tuple[float, list[str], str]] = []
    for priority, selector in enumerate(CONTENT_SELECTORS):
        for node in soup.select(selector)[:5]:
            paragraphs = [norm_space(p.get_text(' ', strip=True)) for p in node.select('p')]
            paragraphs = [p for p in paragraphs if len(p) >= 8]
            if not paragraphs:
                paragraphs = [norm_space(t) for t in node.get_text('\n', strip=True).split('\n')]
                paragraphs = [p for p in paragraphs if len(p) >= 12]
            text = '\n'.join(paragraphs)
            if len(text) < 65:
                continue
            all_a = sum(len(a.get_text(' ', strip=True)) for a in node.select('a'))
            ratio = all_a / max(len(text), 1)
            if ratio > .65:
                continue
            # More real text is useful; prefer article-specific containers for close ties.
            score = min(len(text), 20000) * (1 - ratio) + max(0, 20 - priority) * 18
            candidates.append((score, paragraphs, selector))
    if not candidates:
        # An unknown redesign should not force code changes solely to add a CSS selector.
        # Find the densest prose block with a few substantial paragraphs.
        for node in soup.select('article, main, section, div'):
            paragraphs = [norm_space(x.get_text(' ', strip=True)) for x in node.select('p')]
            paragraphs = [x for x in paragraphs if len(x) >= 15]
            length = sum(map(len, paragraphs))
            if length < 65 or len(paragraphs) < 2:
                continue
            link_chars = sum(len(a.get_text(' ', strip=True)) for a in node.select('a'))
            if link_chars / length > .22:
                continue
            # Penalize giant wrappers and prefer coherent, well-contained paragraphs.
            score = min(length, 12000) - link_chars * 4 - max(0, len(paragraphs) - 35) * 30
            candidates.append((score, paragraphs, '[自动发现的正文容器]'))
    if not candidates:
        return '', []
    score, paragraphs, selector = max(candidates, key=lambda x: x[0])
    cleaned = []
    for p in paragraphs:
        if p in cleaned[-1:]:
            continue
        if re.match(r'^(?:责任编辑|编辑：|来源：|分享到|打印本页|关闭窗口|上一篇|下一篇)', p) and len(p) < 90:
            continue
        cleaned.append(p)
    return selector, cleaned


def parse_article(html: str, fallback_title: str) -> tuple[str, list[str], str]:
    soup = BeautifulSoup(html, 'html.parser')
    title = ''
    for selector in ['#sybt .mainTitle', '#sybt', 'h1.article-title', '.article-title h1', '.article_title h1', 'h1.title', 'article h1', 'h1']:
        node = soup.select_one(selector)
        if node:
            candidate = clean_title(node.get_text(' ', strip=True))
            if len(candidate) >= 4 and len(candidate) <= 250:
                title = candidate
                break
    if not title:
        meta = soup.select_one('meta[property="og:title"], meta[name="title"]')
        title = clean_title(meta.get('content', '')) if meta else ''
    selector, paragraphs = clean_body(soup)
    return title or fallback_title, paragraphs, selector


@contextmanager
def connect_db(path: Path):
    """Manage a SQLite transaction AND close its file handle on all platforms.

    sqlite3.Connection.__exit__ commits/rolls back, but does not close the
    connection. Without an explicit close, Windows retains a lock on the
    database file and temporary-directory cleanup fails with WinError 32.
    """
    con = sqlite3.connect(path)
    try:
        con.row_factory = sqlite3.Row
        con.executescript('''
            CREATE TABLE IF NOT EXISTS news (
                url TEXT PRIMARY KEY,
                source TEXT NOT NULL,
                title TEXT NOT NULL,
                pub_date TEXT NOT NULL,
                article_title TEXT DEFAULT '',
                body_json TEXT DEFAULT '',
                status TEXT DEFAULT '未抓取',
                message TEXT DEFAULT '',
                fetched_at TEXT DEFAULT '',
                created_at TEXT DEFAULT CURRENT_TIMESTAMP
            );
            CREATE INDEX IF NOT EXISTS idx_news_month ON news(pub_date, source);
            CREATE TABLE IF NOT EXISTS metadata (key TEXT PRIMARY KEY, value TEXT NOT NULL);
        ''')
        try:
            yield con
            con.commit()
        except BaseException:
            con.rollback()
            raise
    finally:
        con.close()


def upsert_news(con: sqlite3.Connection, record: News) -> bool:
    found = con.execute('SELECT url FROM news WHERE url=?', (record.url,)).fetchone()
    if found:
        con.execute('UPDATE news SET title=?,pub_date=?,source=? WHERE url=?',
                    (record.title, record.pub_date, record.source, record.url))
        return False
    con.execute('INSERT INTO news(url,source,title,pub_date) VALUES(?,?,?,?)',
                (record.url, record.source, record.title, record.pub_date))
    return True


def get_cell_url(cell, explicit='') -> str:
    if explicit and str(explicit).startswith(('https://', 'http://')):
        return str(explicit)
    if cell.hyperlink and cell.hyperlink.target:
        return cell.hyperlink.target
    if cell.data_type == 'f' and isinstance(cell.value, str):
        m = re.search(r'^=HYPERLINK\(\s*"((?:[^"]|"")*)"', cell.value, re.I)
        return m.group(1).replace('""', '"') if m else ''
    if isinstance(cell.value, str) and cell.value.startswith('http'):
        return cell.value
    return ''


def get_cell_title(cell) -> str:
    v = cell.value
    if isinstance(v, str) and v.upper().startswith('=HYPERLINK('):
        m = re.search(r'^=HYPERLINK\(\s*"(?:[^"]|"")*"\s*[,;]\s*"((?:[^"]|"")*)"', v, re.I)
        return m.group(1).replace('""', '"') if m else ''
    return clean_title(v)


def is_yellow(cell) -> bool:
    if cell.fill is None or cell.fill.patternType != 'solid':
        return False
    color = cell.fill.fgColor
    if color.type == 'rgb' and color.rgb:
        return color.rgb.upper()[-6:] in {'FFFF00', 'FFEB3B', 'FFF200', 'FFEE00', 'FFF2CC'}
    return False


def selection(row) -> str:
    value = norm_space(row[4].value if len(row) > 4 else '').lower()
    if value in {'排除', '否', 'no', 'n', '0', '忽略'}:
        return 'exclude'
    if value in {'收录', '是', 'yes', 'y', '1', '人工收录'}:
        return 'include'
    if is_yellow(row[0]) or (len(row) > 1 and is_yellow(row[1])):
        return 'include'
    return 'auto'


def open_book(path: Path):
    if path.exists():
        book = load_workbook(path)
        # Upgrade old workbook in place while retaining all colored rows, formulas and hyperlinks.
        if '九江新闻' in book.sheetnames and '九江市政府' not in book.sheetnames:
            book['九江新闻'].title = '九江市政府'
    else:
        book = Workbook()
        book.active.title = '九江市政府'
    for key, info in SOURCES.items():
        if info['sheet'] not in book.sheetnames:
            book.create_sheet(info['sheet'])
        ws = book[info['sheet']]
        for i, header in enumerate(HEADERS, 1):
            c = ws.cell(1, i)
            c.value = header
            c.fill = PatternFill('solid', fgColor='203A58')
            c.font = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
            c.alignment = Alignment(horizontal='center', vertical='center')
        for col, width in {'A': 73, 'B': 16, 'C': 14, 'D': 25, 'E': 14, 'F': 16, 'G': 68, 'H': 35}.items():
            ws.column_dimensions[col].width = width
        ws.freeze_panes = 'C2'
        ws.auto_filter.ref = f'A1:H{max(ws.max_row, 2)}'
        ws.row_dimensions[1].height = 27
        if not ws.data_validations.dataValidation:
            validation = DataValidation(type='list', formula1='"收录,排除"', allow_blank=True)
            validation.error = '只能填写“收录”或“排除”'
            validation.errorTitle = '人工选择'
            validation.showErrorMessage = True
            ws.add_data_validation(validation)
            validation.add('E2:E1048576')
    return book


def existing_rows(book, con: sqlite3.Connection) -> dict[str, set[str]]:
    urls = {source: set() for source in SOURCES}
    for source, info in SOURCES.items():
        ws = book[info['sheet']]
        for idx in range(2, ws.max_row + 1):
            row = [ws.cell(idx, j) for j in range(1, 9)]
            url = get_cell_url(row[0], row[6].value)
            if not url:
                continue
            url = normalize_url(info['base'], url)
            if url in urls[source]:
                continue
            urls[source].add(url)
            d = date_from_text(str(row[1].value))
            if isinstance(row[1].value, (datetime, date)):
                d = row[1].value.strftime('%Y-%m-%d')
            title = get_cell_title(row[0])
            if row[0].data_type == 'f' and title:
                # Native hyperlinks survive editing and do not depend on WPS/Excel recalculation.
                # Changing only the value leaves the manual yellow fill intact.
                row[0].value = title
                row[0].hyperlink = url
            if d and title:
                upsert_news(con, News(source, title, url, d))
    con.commit()
    return urls


def write_new_rows(book, records: dict[str, list[News]]):
    for source, news in records.items():
        if not news:
            continue
        ws = book[SOURCES[source]['sheet']]
        new_items = sorted(news, key=lambda x: (x.pub_date, x.title), reverse=True)
        ws.insert_rows(2, len(new_items))
        for ix, item in enumerate(new_items, 2):
            a = ws.cell(ix, 1, item.title)
            a.hyperlink = item.url
            a.font = Font(name='楷体_GB2312', size=12, color='1565C0', underline='single')
            a.alignment = Alignment(vertical='center', wrap_text=True)
            b = ws.cell(ix, 2, item.pub_date)
            b.font = Font(name='微软雅黑', size=11)
            b.alignment = Alignment(horizontal='center', vertical='center')
            ws.cell(ix, 7, item.url)
            ws.row_dimensions[ix].height = 28
        ws.auto_filter.ref = f'A1:H{ws.max_row}'
        LOG.info('%s：新增 %d 条', SOURCES[source]['sheet'], len(new_items))


def crawl(session: requests.Session, con: sqlite3.Connection, known: dict[str, set[str]],
          cfg: dict, year: int, debug: Path, browser: bool,
          target_month: str | None = None) -> dict[str, list[News]]:
    new = {source: [] for source in SOURCES}
    for source in SOURCES:
        known_pages = 0
        next_url = ''
        for page in range(cfg['max_pages']):
            news, ok, discovered_next = fetch_list(session, source, page, cfg, browser, debug, next_url, year)
            next_url = discovered_next
            if not ok or not news:
                break
            in_year = [x for x in news if x.pub_date.startswith(target_month if target_month else f'{year}-')]
            old = len(in_year) and all(x.url in known[source] for x in in_year)
            cutoff = f'{target_month}-01' if target_month else f'{year}-01-01'
            older = bool(news) and all(x.pub_date < cutoff for x in news)
            count = 0
            for n in in_year:
                if n.url not in known[source]:
                    known[source].add(n.url)
                    new[source].append(n)
                    count += 1
                upsert_news(con, n)
            con.commit()
            LOG.info('%s 第%d页：解析 %d 条，目标范围 %d 条，新记录 %d 条', source, page + 1, len(news), len(in_year), count)
            if source == 'gov':
                LOG.info('政府年度 XML 已一次获取完毕，无需分页')
                break
            if older:
                LOG.info('%s 已到 %d 年之前，停止', source, year)
                break
            known_pages = known_pages + 1 if old else 0
            if known_pages >= cfg['stop_after_known_pages']:
                LOG.info('%s 连续 %d 页都是已知记录，停止翻页', source, known_pages)
                break
            if cfg['delay_seconds']:
                time.sleep(cfg['delay_seconds'])
        else:
            LOG.warning('%s 达到 max_pages=%d，可能尚有更早的记录未抓到', source, cfg['max_pages'])
    return new


def fetch_article_details(session: requests.Session, con: sqlite3.Connection, cfg: dict,
                          year: int, month: str, new: dict[str, list[News]], book,
                          browser: bool, scope: str):
    fresh_urls = {n.url for group in new.values() for n in group}
    manual_urls: set[str] = set()
    for source in SOURCES:
        ws = book[SOURCES[source]['sheet']]
        for row in ws.iter_rows(min_row=2, max_col=8):
            url = get_cell_url(row[0], row[6].value)
            if selection(row) == 'include' and url:
                manual_urls.add(normalize_url(SOURCES[source]['base'], url))
    all_rows = con.execute('SELECT * FROM news WHERE pub_date LIKE ? ORDER BY pub_date DESC', (str(year) + '-%',)).fetchall()
    to_fetch = []
    for n in all_rows:
        if n['status'] == '成功' and n['body_json']:
            continue
        if scope == 'all' and n['url'] in fresh_urls:
            to_fetch.append(n)
        elif n['pub_date'].startswith(month) or n['url'] in manual_urls:
            to_fetch.append(n)
    LOG.info('本次需要抓取正文：%d 篇；历史成功正文直接使用缓存', len(to_fetch))
    for idx, n in enumerate(to_fetch, 1):
        try:
            try:
                html = fetch_html(session, n['url'], cfg['timeout_seconds'])
            except Exception:
                if not browser:
                    raise
                html = fetch_with_browser(n['url'], cfg['timeout_seconds'])
            title, paragraphs, selector = parse_article(html, n['title'])
            if sum(map(len, paragraphs)) < 65 and browser:
                html = fetch_with_browser(n['url'], cfg['timeout_seconds'])
                title, paragraphs, selector = parse_article(html, n['title'])
            if sum(map(len, paragraphs)) < 65:
                raise ValueError('未找到足够长的正文；可能是反爬、JS渲染或CSS结构变化')
            con.execute('''UPDATE news SET article_title=?, body_json=?,status='成功',message=?,fetched_at=? WHERE url=?''',
                        (title, json.dumps(paragraphs, ensure_ascii=False), selector, datetime.now().isoformat(timespec='seconds'), n['url']))
            LOG.info('[%d/%d] 正文成功：%s (%d字)', idx, len(to_fetch), n['title'][:33], sum(map(len, paragraphs)))
        except Exception as exc:
            con.execute("UPDATE news SET status='失败',message=?,fetched_at=? WHERE url=?",
                        (str(exc)[:300], datetime.now().isoformat(timespec='seconds'), n['url']))
            LOG.warning('[%d/%d] 正文失败：%s — %s', idx, len(to_fetch), n['title'][:35], exc)
        con.commit()
        if cfg['delay_seconds']:
            time.sleep(cfg['delay_seconds'])


def match_names(title: str, body: str, names: list[str], include_body: bool) -> list[str]:
    target = title + ('\n' + body if include_body else '')
    return [name for name in names if name in target]


def refresh_book(book, con: sqlite3.Connection, names: list[str], cfg: dict) -> list[dict]:
    chosen = []
    for source in SOURCES:
        ws = book[SOURCES[source]['sheet']]
        for i in range(2, ws.max_row + 1):
            row = [ws.cell(i, k) for k in range(1, 9)]
            url = get_cell_url(row[0], row[6].value)
            if not url:
                continue
            url = normalize_url(SOURCES[source]['base'], url)
            n = con.execute('SELECT * FROM news WHERE url=?', (url,)).fetchone()
            if n is None:
                continue
            paras = json.loads(n['body_json'] or '[]')
            matches = match_names(n['article_title'] or n['title'], '\n'.join(paras), names, cfg['include_body_matches'])
            choose = selection(row)
            selected = choose == 'include' or (choose == 'auto' and bool(matches))
            row[2].value = '命中' if matches else ('待正文' if n['status'] != '成功' else '未命中')
            row[3].value = '、'.join(matches)
            row[5].value = n['status']
            row[6].value = url
            row[7].value = n['message'] if n['status'] != '成功' else ('人工收录' if choose == 'include' else '人工排除' if choose == 'exclude' else '')
            row[2].alignment = Alignment(horizontal='center')
            row[4].alignment = Alignment(horizontal='center')
            row[5].alignment = Alignment(horizontal='center')
            if selected:
                chosen.append({'source': source, 'url': url, 'date': n['pub_date'],
                               'title': n['article_title'] or n['title'], 'body': paras,
                               'status': n['status'], 'matches': matches})
    return chosen


def canonical_title(title: str) -> str:
    return re.sub(r'[\W_]+', '', (title or '').lower())


def deduplicate_for_month(selected: list[dict], month: str) -> tuple[list[dict], list[dict]]:
    articles = []
    failed = []
    for n in sorted((x for x in selected if x['date'].startswith(month)), key=lambda x: (x['date'], x['title'])):
        if n['status'] != '成功' or not n['body']:
            failed.append(n)
            continue
        key = canonical_title(n['title'])
        match_idx = None
        for idx, seen in enumerate(articles):
            if canonical_title(seen['title']) != key:
                continue
            if abs((date.fromisoformat(seen['date']) - date.fromisoformat(n['date'])).days) <= 2:
                match_idx = idx
                break
        if match_idx is None:
            articles.append(n)
        elif (len(''.join(n['body'])), n['source'] == 'gov') > (len(''.join(articles[match_idx]['body'])), articles[match_idx]['source'] == 'gov'):
            articles[match_idx] = n
    articles.sort(key=lambda x: (x['date'], x['title']))
    return articles, failed


def set_page_layout(doc):
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin, section.bottom_margin = Cm(2.54), Cm(2.54)
        section.left_margin, section.right_margin = Cm(2.7), Cm(2.7)
    normal = doc.styles['Normal']
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(28)
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)


def add_run(p, text: str, font: str):
    r = p.add_run(text)
    r.font.name = font
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), font)
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return r


def build_word(month: str, articles: list[dict], out_dir: Path) -> Path | None:
    if not articles:
        return None
    y, m = map(int, month.split('-'))
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f'市委大事记（{y}年{m}月）.docx'
    if path.exists():
        existing = Document(path)
        if 'jrjj-v2-generated' not in (existing.core_properties.keywords or ''):
            path = out_dir / f'市委大事记（{y}年{m}月）_自动整理.docx'
        if path.exists():
            backups = out_dir / '备份'
            backups.mkdir(exist_ok=True)
            stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            shutil.copy2(path, backups / f'{path.stem}_{stamp}.docx')
    doc = Document()
    set_page_layout(doc)
    doc.core_properties.keywords = 'jrjj-v2-generated'
    doc.core_properties.title = f'市委大事记（{y}年{m}月）'
    header = doc.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hrun = header.add_run('市 委 大 事 记')
    hrun.font.name = '方正小标宋简体'
    hrun._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '方正小标宋简体')
    hrun.font.size = Pt(22)
    header.add_run('\n')
    date_run = header.add_run(f'（{y}年{m}月）')
    date_run.font.name = '楷体_GB2312'
    date_run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '楷体_GB2312')
    date_run.font.size = Pt(18)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pnum = footer.add_run()
    pnum.font.name = '仿宋_GB2312'
    pnum.font.size = Pt(12)
    field_start = OxmlElement('w:fldChar')
    field_start.set(qn('w:fldCharType'), 'begin')
    field_instr = OxmlElement('w:instrText')
    field_instr.set(qn('xml:space'), 'preserve')
    field_instr.text = ' PAGE '
    field_end = OxmlElement('w:fldChar')
    field_end.set(qn('w:fldCharType'), 'end')
    pnum._r.append(field_start)
    pnum._r.append(field_instr)
    pnum._r.append(field_end)
    for article in articles:
        doc.add_paragraph('')
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = Pt(28)
        p.paragraph_format.keep_with_next = True
        add_run(p, article['title'], '黑体')
        ppr = p._p.get_or_add_pPr()
        lvl = OxmlElement('w:outlineLvl')
        lvl.set(qn('w:val'), '2')
        ppr.append(lvl)
        for para in article['body']:
            if not para:
                continue
            body = doc.add_paragraph()
            body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            body.paragraph_format.first_line_indent = Cm(1.12)
            body.paragraph_format.line_spacing = Pt(28)
            add_run(body, para, '仿宋_GB2312')
    temp = path.with_suffix('.tmp.docx')
    doc.save(temp)
    temp.replace(path)
    return path


def backup_spreadsheet(path: Path) -> None:
    if not path.exists():
        return
    backups = path.parent / '备份'
    backups.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime('%Y%m%d_%H%M%S_%f')
    backup = backups / f'{path.stem}_运行前_{stamp}{path.suffix}'
    shutil.copy2(path, backup)
    LOG.info('已备份现有Excel：%s', backup)


def previous_month(today: date | None = None) -> str:
    today = today or date.today()
    return (today.replace(day=1) - timedelta(days=1)).strftime('%Y-%m')


def args_parser():
    parser = argparse.ArgumentParser(description='九江市政府 + 九江新闻网：增量抓取、自动匹配、按月生成大事记')
    parser.add_argument('--month', help='输出年月，格式 YYYY-MM；默认上个月')
    parser.add_argument('--year', type=int, help='抓取年份，默认输出月份所在年份')
    parser.add_argument('--mode', choices=['all', 'sync', 'export'], default='all', help='all一键完整流程；sync只更新Excel和正文；export跳过列表抓取，只补正文并导出Word')
    parser.add_argument('--content-scope', choices=['all', 'month'], default='all', help='抓取新文章正文的范围：all或仅指定月份（旧Excel历史文章按需抓取）')
    parser.add_argument('--browser', action='store_true', help='启用Playwright浏览器作为HTML/正文失败时的可选回退')
    parser.add_argument('--config', default='jrjj_config.json', help='配置文件路径')
    parser.add_argument('--excel', default='jiujiang_news.xlsx', help='新闻列表Excel文件')
    parser.add_argument('--database', default='jiujiang_cache.sqlite3', help='新闻正文缓存数据库')
    parser.add_argument('--output-dir', default='.', help='Word输出目录')
    parser.add_argument('--debug-dir', default='debug_html', help='页面识别失败时保存HTML的目录')
    parser.add_argument('--max-pages', type=int, help='覆盖配置中的最多分页数')
    parser.add_argument('--full-scan', action='store_true', help='即使连续遇到旧记录，也扫完当年全部分页')
    parser.add_argument('--only-month', action='store_true', help='仅采集所选月份的新闻列表，遇到更早日期即停止；不会把其他月份写入Excel')
    return parser


def main(argv=None) -> int:
    args = args_parser().parse_args(argv)
    logging.basicConfig(level=logging.INFO, format='%(asctime)s %(levelname)s %(message)s', datefmt='%H:%M:%S')
    month = args.month or previous_month()
    if not re.fullmatch(r'20\d{2}-(0[1-9]|1[0-2])', month):
        raise SystemExit('月份格式错误，请使用 --month 2026-09')
    year = args.year or int(month[:4])
    cfg = load_config(Path(args.config))
    if args.max_pages:
        cfg['max_pages'] = args.max_pages
    if args.full_scan:
        cfg['stop_after_known_pages'] = cfg['max_pages'] + 1
    LOG.info('目标月份：%s；抓取年份：%d；执行模式：%s', month, year, args.mode)
    backup_spreadsheet(Path(args.excel))
    book = open_book(Path(args.excel))
    with connect_db(Path(args.database)) as con:
        existing = existing_rows(book, con)
        session = build_session(cfg['user_agent'])
        new = {source: [] for source in SOURCES}
        if args.mode != 'export':
            new = crawl(session, con, existing, cfg, year, Path(args.debug_dir), args.browser,
                        month if args.only_month else None)
            write_new_rows(book, new)
            # Save the new listing first; failed article requests must not lose the work.
            book.save(args.excel)
            fetch_article_details(session, con, cfg, year, month, new, book, args.browser, args.content_scope)
        if args.mode == 'export':
            # Manual review: skip listing refresh; only fetch needed article bodies.
            fetch_article_details(session, con, cfg, year, month, new, book, args.browser, 'month')
        # Always recompute matches from the current configuration; ignore old metadata.
        names = get_names(cfg)
        selected = refresh_book(book, con, names, cfg)
        book.save(args.excel)
        LOG.info('Excel已更新：%s；自动检测名单 %d 人', args.excel, len(names))
        if args.mode != 'sync':
            to_word, failed = deduplicate_for_month(selected, month)
            output = build_word(month, to_word, Path(args.output_dir))
            if output:
                LOG.info('Word已生成：%s，%d 篇（按日期由早到晚、跨站去重）', output, len(to_word))
            else:
                LOG.warning('目标月份没有成功提取且符合收录条件的正文，未创建空Word')
            if failed:
                LOG.warning('有 %d 篇已选文章正文未抓取成功，本次未导出；下次会重试', len(failed))
        LOG.info('完成。请在Excel中使用E列“收录/排除”，或保持旧版黄色标记进行人工修订。')
    return 0


def configure_cli_encoding() -> None:
    """Make Chinese terminal output safe in Windows CI and redirected consoles.

    Some Windows environments default to cp1252. Setting the encoding before
    argparse handles --help prevents UnicodeEncodeError without requiring the
    caller to manually change the system code page.
    """
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, 'reconfigure'):
            try:
                stream.reconfigure(encoding='utf-8', errors='backslashreplace')
            except (OSError, ValueError):
                pass


if __name__ == '__main__':
    configure_cli_encoding()
    try:
        sys.exit(main())
    except KeyboardInterrupt:
        LOG.error('用户中断；已落盘的缓存、Excel仍被保留')
        sys.exit(130)
    except PermissionError as exc:
        LOG.error('文件被占用，请先关闭 Excel/Word 再运行：%s', exc)
        sys.exit(2)